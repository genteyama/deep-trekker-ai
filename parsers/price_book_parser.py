from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import BinaryIO, Optional, Sequence, Union

from openpyxl import load_workbook

from models import (
    PriceBookImportResult,
    PriceBookIssue,
    PriceBookSheetSummary,
    SKU,
    SKUMasterCandidate,
    SKUSourceOccurrence,
    SkuDuplicateClass,
    SkuSourceStatus,
    ValidationSeverity,
)

HEADER_SCAN_ROWS = 30
NO_DEALER_DISCOUNT_MARKERS = ("NO DEALER DISCOUNT",)
REQUIRED_FIELDS = ("sku", "description", "msrp", "dealer_price")
SKIP_SHEET_KEYS = {"config"}
OBSOLETE_SHEET_KEYS = {"revolution-obsolete", "revolutionobsolete"}
DEFAULT_COLUMN_ALIASES = {
    "sku": ("part number", "part #", "part#", "part no", "part no.", "sku"),
    "description": ("description", "desc", "item description"),
    "msrp": ("msrp", "msrp usd", "msrp (usd)", "list price"),
    "dealer_price": (
        "dealer price",
        "dealer price usd",
        "dealer price (usd)",
        "dealer",
        "dt40",
        "pt30",
    ),
    "notes": ("notes", "notes:", "note", "comments", "comment"),
}
KNOWN_SHEET_HINTS = {
    "revolution": ("ROV", "REVOLUTION"),
    "revolution-dc": ("ROV", "REVOLUTION"),
    "revolutiondc": ("ROV", "REVOLUTION"),
    "revolution-obsolete": ("ROV", "REVOLUTION"),
    "revolutionobsolete": ("ROV", "REVOLUTION"),
    "pivot": ("ROV", "PIVOT"),
    "photon": ("ROV", "PHOTON"),
    "spectra": ("ROV", "SPECTRA"),
    "dtg3": ("ROV", "DTG3"),
    "mag": ("Utility Crawler", "MAG"),
    "vac": ("Utility Crawler", "VAC"),
    "vac&mag": ("Utility Crawler", "VAC & MAG"),
    "vacmag": ("Utility Crawler", "VAC & MAG"),
    "pipetrekker": ("PipeTrekker", "PipeTrekker"),
    "a-150": ("PipeTrekker", "A-150"),
    "a150": ("PipeTrekker", "A-150"),
    "a-200": ("PipeTrekker", "A-200"),
    "a200": ("PipeTrekker", "A-200"),
    "pt-van": ("PipeTrekker", "PT-VAN"),
    "ptvan": ("PipeTrekker", "PT-VAN"),
}
CODE_MISSING_COLUMNS = "MISSING_REQUIRED_COLUMNS"
CODE_EMPTY_SKU = "EMPTY_SKU"
CODE_DUPLICATE_SKU = "DUPLICATE_SKU"
CODE_INVALID_MSRP = "INVALID_MSRP"
CODE_INVALID_DEALER_PRICE = "INVALID_DEALER_PRICE"
CODE_DEALER_GT_MSRP = "DEALER_PRICE_GT_MSRP"
CODE_MSRP_EQUALS_DEALER = "MSRP_EQUALS_DEALER_PRICE"
CODE_NO_DEALER_DISCOUNT = "NO_DEALER_DISCOUNT"
CODE_PART_NUMBER_INVALID = "PART_NUMBER_INVALID"
CODE_DUPLICATE_SAME = "DUPLICATE_SAME"
CODE_DUPLICATE_CONTEXT_DIFFERENT = "DUPLICATE_CONTEXT_DIFFERENT"
CODE_PRICE_CONFLICT = "PRICE_CONFLICT"


def parse_price_book(
    source: Union[str, Path, BinaryIO],
    *,
    source_price_book: Optional[str] = None,
    version: Optional[str] = None,
    column_aliases: Optional[dict] = None,
    imported_at: Optional[datetime] = None,
    data_only: bool = True,
) -> PriceBookImportResult:
    aliases = _merge_aliases(column_aliases)
    book_name = source_price_book or _default_book_name(source)
    timestamp = imported_at or datetime.now(timezone.utc)
    workbook = load_workbook(source, data_only=data_only)
    result = PriceBookImportResult(
        source_price_book=book_name,
        version=version,
        imported_at=timestamp,
    )

    try:
        for sheet in workbook.worksheets:
            _parse_sheet(sheet, aliases, book_name, version, timestamp, result)
    finally:
        workbook.close()
    _build_candidates(result)
    return result


def normalize_sku(value) -> Optional[str]:
    sku, invalid = classify_sku_cell(value)
    if invalid:
        return None
    return sku


def classify_sku_cell(value) -> tuple[Optional[str], Optional[str]]:
    if value is None:
        return None, None
    if isinstance(value, bool):
        return None, None
    if isinstance(value, datetime) or type(value) is date:
        return None, CODE_PART_NUMBER_INVALID
    if isinstance(value, float) and value.is_integer():
        return str(int(value)), None
    if isinstance(value, int):
        return str(value), None
    text = str(value).strip()
    return (text or None), None


def parse_money(value) -> tuple[Optional[float], bool]:
    if value is None or value == "":
        return None, False
    if isinstance(value, bool):
        return None, True
    if isinstance(value, (int, float)):
        return round(float(value), 2), False
    text = str(value).replace("$", "").replace(",", "").replace("USD", "").strip()
    if text == "":
        return None, False
    try:
        return round(float(text), 2), False
    except ValueError:
        return None, True


def calculate_dealer_rate(msrp_usd: Optional[float], dealer_price_usd: Optional[float]) -> Optional[float]:
    if msrp_usd is None or dealer_price_usd is None or msrp_usd == 0:
        return None
    return round(dealer_price_usd / msrp_usd, 4)


def sheet_key(sheet_name: str) -> str:
    return "".join(str(sheet_name).lower().split()).replace("_", "")


def sheet_product_hint(sheet_name: str) -> tuple[Optional[str], Optional[str]]:
    return KNOWN_SHEET_HINTS.get(sheet_key(sheet_name), (None, None))


def sheet_source_status(sheet_name: str) -> SkuSourceStatus:
    if sheet_key(sheet_name) in OBSOLETE_SHEET_KEYS:
        return SkuSourceStatus.OBSOLETE
    return SkuSourceStatus.ACTIVE


def notes_has_no_dealer_discount(notes: Optional[str]) -> bool:
    content = notes or ""
    return any(marker in content.upper() for marker in NO_DEALER_DISCOUNT_MARKERS)


def _parse_sheet(sheet, aliases, book_name, version, timestamp, result) -> None:
    family, model = sheet_product_hint(sheet.title)
    status = sheet_source_status(sheet.title)
    summary = PriceBookSheetSummary(
        sheet_name=sheet.title,
        product_family=family,
        model=model,
        source_status=status,
    )
    if sheet_key(sheet.title) in SKIP_SHEET_KEYS:
        summary.skipped = True
        summary.skip_reason = "CONFIG_SHEET"
        result.sheets.append(summary)
        return

    header = _find_header(sheet, aliases)
    if header is None:
        _add_issue(
            result,
            ValidationSeverity.BLOCKER,
            CODE_MISSING_COLUMNS,
            "Required columns were not found.",
            source_sheet=sheet.title,
            details="sku, description, msrp, dealer_price",
        )
        result.sheets.append(summary)
        return

    summary.header_found = True
    missing = [field for field in REQUIRED_FIELDS if field not in header["columns"]]
    if missing:
        _add_issue(
            result,
            ValidationSeverity.BLOCKER,
            CODE_MISSING_COLUMNS,
            "Required columns were not found.",
            source_sheet=sheet.title,
            details=", ".join(missing),
        )
        result.sheets.append(summary)
        return

    for row_number, values in header["rows"]:
        item, occurrence, issues = _row_to_sku(
            values,
            header["columns"],
            sheet.title,
            family,
            model,
            book_name,
            version,
            timestamp,
            row_number,
            status,
        )
        for issue in issues:
            _add_issue_model(result, issue)
        if item is not None and occurrence is not None:
            result.items.append(item)
            result.occurrences.append(occurrence)
            summary.item_count += 1
    result.sheets.append(summary)


def _find_header(sheet, aliases) -> Optional[dict]:
    rows = [(index, list(row)) for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    for header_index, (_, values) in enumerate(rows[:HEADER_SCAN_ROWS]):
        columns = _map_columns(values, aliases)
        if "sku" not in columns:
            continue
        return {"columns": columns, "rows": rows[header_index + 1 :]}
    return None


def _map_columns(values: Sequence, aliases: dict) -> dict:
    mapped = {}
    for index, value in enumerate(values):
        normalized = _normalize_header(value)
        if not normalized:
            continue
        for field, names in aliases.items():
            if field in mapped:
                continue
            if normalized in names:
                mapped[field] = index
                break
    return mapped


def _row_to_sku(
    values,
    columns,
    sheet_name,
    family,
    model,
    book_name,
    version,
    timestamp,
    row_number,
    status,
) -> tuple[Optional[SKU], Optional[SKUSourceOccurrence], list[PriceBookIssue]]:
    issues = []
    raw_sku = _cell(values, columns["sku"])
    sku, sku_invalid = classify_sku_cell(raw_sku)
    description = _optional_text(_cell(values, columns.get("description")))
    notes = _optional_text(_cell(values, columns.get("notes")))
    msrp, msrp_invalid = parse_money(_cell(values, columns.get("msrp")))
    dealer_price, dealer_invalid = parse_money(_cell(values, columns.get("dealer_price")))
    if sku is None and description is None and msrp is None and dealer_price is None and notes is None:
        return None, None, issues
    if _is_non_product_row(sku, description, msrp, dealer_price):
        return None, None, issues
    if sku_invalid == CODE_PART_NUMBER_INVALID:
        issues.append(
            _issue(
                ValidationSeverity.WARNING,
                CODE_PART_NUMBER_INVALID,
                "Part Number is not a usable SKU string.",
                source_sheet=sheet_name,
                details=f"row={row_number}",
            )
        )
        return None, None, issues
    if sku is None:
        if description or msrp is not None or dealer_price is not None:
            issues.append(
                _issue(
                    ValidationSeverity.BLOCKER,
                    CODE_EMPTY_SKU,
                    "SKU is empty.",
                    source_sheet=sheet_name,
                    details=f"row={row_number}",
                )
            )
        return None, None, issues
    if msrp_invalid:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKER,
                CODE_INVALID_MSRP,
                "MSRP is not a number.",
                sku=sku,
                source_sheet=sheet_name,
                details=f"row={row_number}",
            )
        )
    if dealer_invalid:
        issues.append(
            _issue(
                ValidationSeverity.BLOCKER,
                CODE_INVALID_DEALER_PRICE,
                "Dealer Price is not a number.",
                sku=sku,
                source_sheet=sheet_name,
                details=f"row={row_number}",
            )
        )
    dealer_equals_msrp = msrp is not None and dealer_price is not None and dealer_price == msrp
    no_dealer_discount_note = notes_has_no_dealer_discount(notes)
    if msrp is not None and dealer_price is not None and dealer_price > msrp:
        issues.append(
            _issue(
                ValidationSeverity.WARNING,
                CODE_DEALER_GT_MSRP,
                "Dealer Price is greater than MSRP.",
                sku=sku,
                source_sheet=sheet_name,
            )
        )
    if dealer_equals_msrp:
        issues.append(
            _issue(
                ValidationSeverity.INFO,
                CODE_MSRP_EQUALS_DEALER,
                "Dealer Price equals MSRP.",
                sku=sku,
                source_sheet=sheet_name,
            )
        )
    if no_dealer_discount_note:
        issues.append(
            _issue(
                ValidationSeverity.INFO,
                CODE_NO_DEALER_DISCOUNT,
                "Notes include NO DEALER DISCOUNT.",
                sku=sku,
                source_sheet=sheet_name,
            )
        )
    item = SKU(
        sku=sku,
        product_family=family,
        model=model,
        description=description,
        msrp_usd=msrp,
        dealer_price_usd=dealer_price,
        dealer_rate=calculate_dealer_rate(msrp, dealer_price),
        dealer_equals_msrp=dealer_equals_msrp,
        no_dealer_discount_note=no_dealer_discount_note,
        notes=notes,
        source_price_book=book_name,
        source_sheet=sheet_name,
        source_row=row_number,
        source_status=status,
        price_book_version=version,
        last_synced_at=timestamp,
    )
    occurrence = SKUSourceOccurrence(
        sku=sku,
        source_price_book=book_name,
        source_sheet=sheet_name,
        source_row=row_number,
        description=description,
        msrp_usd=msrp,
        dealer_price_usd=dealer_price,
        notes=notes,
        source_status=status,
        dealer_equals_msrp=dealer_equals_msrp,
        no_dealer_discount_note=no_dealer_discount_note,
    )
    return item, occurrence, issues


def _is_non_product_row(sku, description, msrp, dealer_price) -> bool:
    if sku:
        return False
    desc = (description or "").strip().upper()
    if desc.startswith("PRICING HAS BEEN UPDATED"):
        return True
    if description is None and msrp == 1 and dealer_price in (0.6, 0.7):
        return True
    if msrp is None and dealer_price is None and any(
        marker in desc for marker in ("INCLUDE", "THIS PAGE IS FOR", "ARE OBSOLETE")
    ):
        return True
    return False


def _build_candidates(result: PriceBookImportResult) -> None:
    grouped: dict[str, list[SKUSourceOccurrence]] = defaultdict(list)
    for occurrence in result.occurrences:
        grouped[occurrence.sku].append(occurrence)

    for sku, occurrences in grouped.items():
        duplicate_class = _classify_duplicates(occurrences)
        if duplicate_class == SkuDuplicateClass.PRICE_CONFLICT:
            _add_issue(
                result,
                ValidationSeverity.WARNING,
                CODE_PRICE_CONFLICT,
                "Same SKU has different manufacturer prices.",
                sku=sku,
                details="do_not_auto_select",
            )
        elif duplicate_class == SkuDuplicateClass.DUPLICATE_CONTEXT_DIFFERENT:
            _add_issue(
                result,
                ValidationSeverity.INFO,
                CODE_DUPLICATE_CONTEXT_DIFFERENT,
                "Same SKU and price appear with different notes or description.",
                sku=sku,
            )
        elif duplicate_class == SkuDuplicateClass.DUPLICATE_SAME:
            _add_issue(
                result,
                ValidationSeverity.INFO,
                CODE_DUPLICATE_SAME,
                "Same SKU is listed more than once with identical content.",
                sku=sku,
            )

        active = [item for item in occurrences if item.source_status != SkuSourceStatus.OBSOLETE]
        value_source = active or occurrences
        source_status = SkuSourceStatus.ACTIVE if active else SkuSourceStatus.OBSOLETE
        prices = {(_money(item.msrp_usd), _money(item.dealer_price_usd)) for item in value_source}
        price_conflict = len(prices) > 1
        descriptions = {_text(item.description) for item in value_source}
        msrp = dealer = description = dealer_rate = None
        dealer_equals = False
        if not price_conflict:
            first = value_source[0]
            msrp = first.msrp_usd
            dealer = first.dealer_price_usd
            dealer_rate = calculate_dealer_rate(msrp, dealer)
            dealer_equals = all(item.dealer_equals_msrp for item in value_source)
            if len(descriptions) == 1:
                description = first.description
        result.candidates.append(
            SKUMasterCandidate(
                sku=sku,
                description=description,
                msrp_usd=msrp,
                dealer_price_usd=dealer,
                dealer_rate=dealer_rate,
                dealer_equals_msrp=dealer_equals,
                no_dealer_discount_note=any(item.no_dealer_discount_note for item in occurrences),
                source_status=source_status,
                occurrence_count=len(occurrences),
                duplicate_class=duplicate_class,
                occurrences=list(occurrences),
            )
        )


def _classify_duplicates(occurrences: Sequence[SKUSourceOccurrence]) -> Optional[SkuDuplicateClass]:
    if len(occurrences) < 2:
        return None
    prices = {(_money(item.msrp_usd), _money(item.dealer_price_usd)) for item in occurrences}
    if len(prices) > 1:
        return SkuDuplicateClass.PRICE_CONFLICT
    descriptions = {_text(item.description) for item in occurrences}
    notes = {_text(item.notes) for item in occurrences}
    if len(descriptions) > 1 or len(notes) > 1:
        return SkuDuplicateClass.DUPLICATE_CONTEXT_DIFFERENT
    return SkuDuplicateClass.DUPLICATE_SAME


def _cell(values, index: Optional[int]):
    if index is None or index < 0 or index >= len(values):
        return None
    return values[index]


def _optional_text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().lower().split()).rstrip(":").strip()


def _merge_aliases(column_aliases: Optional[dict]) -> dict:
    merged = {field: tuple(names) for field, names in DEFAULT_COLUMN_ALIASES.items()}
    if not column_aliases:
        return merged
    for field, names in column_aliases.items():
        extra = tuple(_normalize_header(name) for name in names)
        merged[field] = extra + merged.get(field, ())
    return merged


def _default_book_name(source) -> Optional[str]:
    name = getattr(source, "name", None)
    if name:
        return Path(name).stem
    if isinstance(source, (str, Path)):
        return Path(source).stem
    return None


def _money(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    return round(float(value), 2)


def _text(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    text = value.strip()
    return text or None


def _issue(
    severity: ValidationSeverity,
    code: str,
    message: str,
    sku: Optional[str] = None,
    source_sheet: Optional[str] = None,
    details: Optional[str] = None,
) -> PriceBookIssue:
    return PriceBookIssue(
        severity=severity,
        code=code,
        message=message,
        sku=sku,
        source_sheet=source_sheet,
        details=details,
    )


def _add_issue(result: PriceBookImportResult, severity, code, message, **kwargs) -> None:
    _add_issue_model(result, _issue(severity, code, message, **kwargs))


def _add_issue_model(result: PriceBookImportResult, issue: PriceBookIssue) -> None:
    if issue.severity == ValidationSeverity.BLOCKER:
        result.errors.append(issue)
    elif issue.severity == ValidationSeverity.WARNING:
        result.warnings.append(issue)
    else:
        result.infos.append(issue)

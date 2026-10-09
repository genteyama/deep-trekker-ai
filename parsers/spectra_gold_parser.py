from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
import math
from pathlib import Path
from typing import BinaryIO, Optional, Union

from openpyxl import load_workbook

from models import (
    PriceBookImportResult,
    PriceBookIssue,
    PriceBookSheetSummary,
    SKU,
    SKUMasterCandidate,
    SKUSourceOccurrence,
    SkuDuplicateClass,
    SkuSourceStatus,
    ValidationSeverity,
)
from parsers.price_book_parser import normalize_sku

SPECTRA_SHEET = "spectra"
REQUIRED_HEADERS = ("part number", "description", "spectra", "gold")
HEADER_SCAN_ROWS = 50
UNAVAILABLE_PRICE_MARKERS = {
    "CONTACT PRODUCT TO QUOTE",
    "CONTACT PRODUCT FOR QUOTE",
    "CONTACT FOR QUOTE",
    "TBD",
    "N/A",
    "NA",
    "-",
}

CODE_SHEET_MISSING = "SPECTRA_SHEET_MISSING"
CODE_COLUMNS_MISSING = "SPECTRA_COLUMNS_MISSING"
CODE_IDENTITY_UNCONFIRMED = "SPECTRA_IDENTITY_UNCONFIRMED"
CODE_NO_SKUS = "SPECTRA_NO_SKUS"
CODE_PRICE_UNAVAILABLE = "SPECTRA_PRICE_UNAVAILABLE"
CODE_PRICE_INVALID = "SPECTRA_PRICE_INVALID"
CODE_DUPLICATE_CONFLICT = "SPECTRA_DUPLICATE_CONFLICT"


def parse_spectra_gold(
    source: Union[str, Path, BinaryIO],
    *,
    source_price_book: str = "SPECTRA_GOLD",
    version: Optional[str] = None,
) -> PriceBookImportResult:
    """Parse only the SPECTRA sheet. Other Gold-tier product sheets are intentionally excluded."""
    timestamp = datetime.now(timezone.utc)
    result = PriceBookImportResult(
        source_price_book=source_price_book,
        version=version,
        imported_at=timestamp,
    )
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = next((item for item in workbook.worksheets if _header(item.title) == SPECTRA_SHEET), None)
        if sheet is None:
            result.errors.append(_issue(CODE_SHEET_MISSING, "SPECTRA sheet was not found."))
            return result
        summary = PriceBookSheetSummary(
            sheet_name=sheet.title,
            product_family="SPECTRA",
            model="SPECTRA",
            source_status=SkuSourceStatus.ACTIVE,
            dealer_price_label="gold",
        )
        result.sheets.append(summary)
        rows = [(index, list(row)) for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
        header = _find_header(rows)
        if header is None:
            result.errors.append(
                _issue(
                    CODE_COLUMNS_MISSING,
                    "Required SPECTRA columns were not found.",
                    source_sheet=sheet.title,
                    details=", ".join(REQUIRED_HEADERS),
                )
            )
            return result
        summary.header_found = True
        if not _has_source_identity(rows):
            result.errors.append(
                _issue(
                    CODE_IDENTITY_UNCONFIRMED,
                    "SPECTRA MSRP and Gold dealer source markers were not confirmed.",
                    source_sheet=sheet.title,
                )
            )
            return result
        columns, header_offset = header
        for row_number, values in rows[header_offset:]:
            _parse_row(
                result,
                values,
                columns,
                row_number=row_number,
                sheet_name=sheet.title,
                source_price_book=source_price_book,
                version=version,
                timestamp=timestamp,
            )
        summary.item_count = len(result.items)
        if not result.items:
            result.errors.append(
                _issue(CODE_NO_SKUS, "No SPECTRA SKUs were found.", source_sheet=sheet.title)
            )
        _build_candidates(result)
        return result
    finally:
        workbook.close()


def _parse_row(
    result: PriceBookImportResult,
    values: list,
    columns: dict[str, int],
    *,
    row_number: int,
    sheet_name: str,
    source_price_book: str,
    version: Optional[str],
    timestamp: datetime,
) -> None:
    sku = normalize_sku(_cell(values, columns["part number"]))
    if not sku:
        return
    description = _text(_cell(values, columns["description"]))
    notes = _text(_cell(values, columns.get("notes")))
    raw_msrp, msrp_explicit, msrp_invalid = _price(_cell(values, columns["spectra"]))
    raw_dealer, dealer_explicit, dealer_invalid = _price(_cell(values, columns["gold"]))
    complete = msrp_explicit and dealer_explicit and not msrp_invalid and not dealer_invalid
    msrp = raw_msrp if complete else None
    dealer = raw_dealer if complete else None
    if not complete:
        result.warnings.append(
            PriceBookIssue(
                severity=ValidationSeverity.WARNING,
                code=CODE_PRICE_INVALID if msrp_invalid or dealer_invalid else CODE_PRICE_UNAVAILABLE,
                message="SPECTRA price is incomplete; manual review is required.",
                sku=sku,
                source_sheet=sheet_name,
                details=f"row={row_number}",
            )
        )
    dealer_rate = None
    if msrp is not None and dealer is not None and msrp != 0:
        dealer_rate = round(dealer / msrp, 4)
    item = SKU(
        sku=sku,
        product_family="SPECTRA",
        model="SPECTRA",
        description=description,
        msrp_usd=msrp,
        dealer_price_usd=dealer,
        dealer_rate=dealer_rate,
        dealer_equals_msrp=msrp is not None and dealer is not None and msrp == dealer,
        notes=notes,
        source_price_book=source_price_book,
        source_sheet=sheet_name,
        source_row=row_number,
        source_status=SkuSourceStatus.ACTIVE,
        price_book_version=version,
        last_synced_at=timestamp,
    )
    occurrence = SKUSourceOccurrence(
        sku=sku,
        source_price_book=source_price_book,
        source_sheet=sheet_name,
        source_row=row_number,
        description=description,
        # Raw partial numeric values are retained for audit, but never promoted to a candidate.
        msrp_usd=raw_msrp,
        dealer_price_usd=raw_dealer,
        notes=notes,
        source_status=SkuSourceStatus.ACTIVE,
        dealer_equals_msrp=(
            msrp_explicit
            and dealer_explicit
            and raw_msrp is not None
            and raw_dealer is not None
            and raw_msrp == raw_dealer
        ),
    )
    result.items.append(item)
    result.occurrences.append(occurrence)


def _build_candidates(result: PriceBookImportResult) -> None:
    grouped: dict[str, list[SKUSourceOccurrence]] = defaultdict(list)
    safe_items = {item.sku: item for item in result.items}
    for occurrence in result.occurrences:
        grouped[occurrence.sku].append(occurrence)
    for sku, occurrences in grouped.items():
        raw_prices = {(item.msrp_usd, item.dealer_price_usd) for item in occurrences}
        duplicate_class = None
        if len(occurrences) > 1:
            if len(raw_prices) > 1:
                duplicate_class = SkuDuplicateClass.PRICE_CONFLICT
                result.errors.append(
                    _issue(
                        CODE_DUPLICATE_CONFLICT,
                        "The same SPECTRA SKU has conflicting prices.",
                        sku=sku,
                        source_sheet=occurrences[0].source_sheet,
                    )
                )
            elif len({(_text(item.description), _text(item.notes)) for item in occurrences}) > 1:
                duplicate_class = SkuDuplicateClass.DUPLICATE_CONTEXT_DIFFERENT
            else:
                duplicate_class = SkuDuplicateClass.DUPLICATE_SAME
        safe = safe_items[sku]
        if duplicate_class == SkuDuplicateClass.PRICE_CONFLICT:
            msrp = dealer = dealer_rate = None
        else:
            msrp = safe.msrp_usd
            dealer = safe.dealer_price_usd
            dealer_rate = safe.dealer_rate
        result.candidates.append(
            SKUMasterCandidate(
                sku=sku,
                description=safe.description,
                msrp_usd=msrp,
                dealer_price_usd=dealer,
                dealer_rate=dealer_rate,
                dealer_equals_msrp=msrp is not None and dealer is not None and msrp == dealer,
                source_status=SkuSourceStatus.ACTIVE,
                occurrence_count=len(occurrences),
                duplicate_class=duplicate_class,
                occurrences=occurrences,
            )
        )


def _find_header(rows: list[tuple[int, list]]) -> Optional[tuple[dict[str, int], int]]:
    for offset, (_, values) in enumerate(rows[:HEADER_SCAN_ROWS], start=1):
        mapped = {}
        for index, value in enumerate(values):
            name = _header(value)
            if name in REQUIRED_HEADERS or name == "notes":
                mapped.setdefault(name, index)
        if all(name in mapped for name in REQUIRED_HEADERS):
            return mapped, offset
    return None


def _has_source_identity(rows: list[tuple[int, list]]) -> bool:
    values = [_marker(value) for _, row in rows for value in row if value is not None]
    has_msrp = any(
        "manufacturer s suggested retail price" in value and "usd" in value
        for value in values
    )
    has_gold = any(
        "spectra dealer" in value and "30" in value and "discount" in value
        for value in values
    )
    return has_msrp and has_gold


def _price(value) -> tuple[Optional[float], bool, bool]:
    if value is None or value == "":
        return None, False, False
    if isinstance(value, bool):
        return None, False, True
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number):
            return None, False, True
        return round(number, 2), True, False
    text = str(value).strip()
    if text.upper() in UNAVAILABLE_PRICE_MARKERS:
        return None, False, False
    normalized = text.replace("$", "").replace(",", "").replace("USD", "").strip()
    try:
        number = float(normalized)
    except ValueError:
        return None, False, True
    if not math.isfinite(number):
        return None, False, True
    return round(number, 2), True, False


def _issue(
    code: str,
    message: str,
    *,
    sku: Optional[str] = None,
    source_sheet: Optional[str] = None,
    details: Optional[str] = None,
) -> PriceBookIssue:
    return PriceBookIssue(
        severity=ValidationSeverity.BLOCKER,
        code=code,
        message=message,
        sku=sku,
        source_sheet=source_sheet,
        details=details,
    )


def _cell(values: list, index: Optional[int]):
    if index is None or index >= len(values):
        return None
    return values[index]


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _header(value) -> str:
    return " ".join(str(value or "").strip().casefold().rstrip(":").split())


def _marker(value) -> str:
    text = str(value or "").replace("’", "'").replace("'", " ")
    return " ".join("".join(character if character.isalnum() else " " for character in text.casefold()).split())

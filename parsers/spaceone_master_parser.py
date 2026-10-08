import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import BinaryIO, Optional, Union

from openpyxl import load_workbook

from models import (
    CellReference,
    SpaceOneMasterImportResult,
    SpaceOneMasterItem,
    SpaceOneValues,
)
from parsers.price_book_parser import CODE_PART_NUMBER_INVALID, parse_money

HEADER_SCAN_ROWS = 20
SHIPPING_MARKERS = (
    "LARGE BOX",
    "SMALL BOX",
    "PIPE TREKKER SHIPPING",
    "PIPETREKKER SHIPPING",
    "輸送費",
)
HEADERISH_SKU = ("part number", "dtマスター", "価格表より")
# Sheets named "_SRC_..." mirror manufacturer data for lookups inside the SO Master. They are never item sources.
HELPER_SHEET_PREFIX = "_SRC_"
IMPORT_RE = re.compile(
    r'IMPORTRANGE\(\s*"+(?P<url>[^"]+)"+\s*,\s*"+(?P<range>[^"]+)"+',
    re.IGNORECASE,
)
BOOK_ID_RE = re.compile(r"/d/([A-Za-z0-9_-]+)")
KNOWN_BOOKS = {
    "1xVJqlhF-sqMKJ3bnZ5hYjP2lTiYMl_5VaODFBL5CIx4": "DT40",
    "1VoWnPU8KRnRM7ddNGQJIEN2I_zAborsBfBmOxsNMXl4": "PT30",
}


def parse_spaceone_master(
    source: Union[str, Path, BinaryIO],
    *,
    source_name: Optional[str] = None,
    imported_at: Optional[datetime] = None,
) -> SpaceOneMasterImportResult:
    name = source_name or _default_name(source)
    timestamp = imported_at or datetime.now(timezone.utc)
    formula_wb = load_workbook(source, data_only=False)
    if hasattr(source, "seek"):
        source.seek(0)
    value_wb = load_workbook(source, data_only=True)
    result = SpaceOneMasterImportResult(source_name=name, imported_at=timestamp)
    counter = 0
    try:
        for formula_sheet, value_sheet in zip(formula_wb.worksheets, value_wb.worksheets):
            if value_sheet.title.startswith(HELPER_SHEET_PREFIX):
                continue
            header_row = _find_header_row(value_sheet)
            if header_row is None:
                continue
            max_row = max(formula_sheet.max_row or 0, value_sheet.max_row or 0)
            for row_number in range(header_row + 1, max_row + 1):
                item, counter = _row_to_item(
                    formula_sheet,
                    value_sheet,
                    row_number,
                    header_row,
                    max_row,
                    counter,
                )
                if item is not None:
                    result.items.append(item)
    finally:
        formula_wb.close()
        value_wb.close()
    return result


def classify_spaceone_sku(value) -> tuple[Optional[str], Optional[str], str, Optional[str]]:
    if value is None or value == "":
        return None, None, "empty", None
    if isinstance(value, bool):
        return None, CODE_PART_NUMBER_INVALID, "bool", str(value)
    if isinstance(value, datetime) or type(value) is date:
        return None, CODE_PART_NUMBER_INVALID, "datetime", str(value)
    if isinstance(value, float):
        raw = str(value)
        if value.is_integer():
            return str(int(value)), None, "float", raw
        return None, CODE_PART_NUMBER_INVALID, "float", raw
    if isinstance(value, int):
        return str(value), None, "int", str(value)
    text = str(value).strip()
    if not text:
        return None, None, "str", None
    normalized = "".join(text.split())
    return normalized or None, None, "str", text


def parse_cell_reference(formula: Optional[str]) -> Optional[CellReference]:
    if not formula or not isinstance(formula, str):
        return None
    match = IMPORT_RE.search(formula.replace("'", ""))
    if match is None:
        return None
    url = match.group("url")
    range_text = match.group("range").strip()
    book_id = None
    book_match = BOOK_ID_RE.search(url)
    if book_match:
        book_id = book_match.group(1)
    sheet = cell = None
    if "!" in range_text:
        sheet, cell = range_text.rsplit("!", 1)
        sheet = sheet.strip()
        cell = cell.strip()
    return CellReference(
        workbook=KNOWN_BOOKS.get(book_id, book_id),
        workbook_id=book_id,
        sheet=sheet,
        cell=cell,
        formula=formula,
    )


def _find_header_row(sheet) -> Optional[int]:
    for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1):
        blob = " ".join(str(value).lower() for value in values if value is not None)
        if "part number" in blob and "価格" in blob:
            return row_number
        if "part number" in blob:
            return row_number
    return None


def _row_to_item(formula_sheet, value_sheet, row_number, header_row, max_row, counter):
    sku_value = value_sheet.cell(row_number, 2).value
    name = _text(value_sheet.cell(row_number, 3).value)
    description = _text(value_sheet.cell(row_number, 4).value)
    category = _text(value_sheet.cell(row_number, 1).value)
    normalized, invalid, cell_type, raw = classify_spaceone_sku(sku_value)
    msrp_formula = formula_sheet.cell(row_number, 5).value
    dealer_formula = formula_sheet.cell(row_number, 6).value
    msrp, _ = parse_money(value_sheet.cell(row_number, 5).value)
    dealer, _ = parse_money(value_sheet.cell(row_number, 6).value)
    if msrp is None:
        msrp = _iferror_fallback(msrp_formula)
    if dealer is None:
        dealer = _iferror_fallback(dealer_formula)
    sales_price = _sales_price(value_sheet, row_number, max_row)
    shipping = _is_legacy_shipping(normalized, name, description)
    headerish = _is_headerish(raw or normalized)
    if headerish:
        return None, counter
    if normalized is None and not invalid and not shipping and name is None and msrp is None and dealer is None:
        return None, counter
    if normalized is None and not invalid and not shipping:
        return None, counter
    formula = msrp_formula
    if not (isinstance(formula, str) and "IMPORTRANGE" in formula.upper()):
        formula = dealer_formula
    sales_formula, sales_formula_row, jpy_msrp_formula = _sales_formula(
        formula_sheet, value_sheet, row_number, max_row
    )
    detected_rate, rate_cell = _sheet_exchange_rate(formula_sheet, value_sheet)
    counter += 1
    item = SpaceOneMasterItem(
        spaceone_item_id=f"so-{counter:03d}",
        source_sheet=value_sheet.title,
        source_row=row_number,
        spaceone_sku=raw if cell_type == "str" else (normalized or raw),
        normalized_sku=None if invalid else normalized,
        sku_cell_type=cell_type,
        sku_raw=raw,
        part_number_invalid=invalid == CODE_PART_NUMBER_INVALID,
        is_legacy_shipping=shipping,
        old_reference=parse_cell_reference(formula if isinstance(formula, str) else None),
        sales_price_formula=_formula_text(sales_formula),
        sales_price_formula_row=sales_formula_row,
        jpy_msrp_formula=_formula_text(jpy_msrp_formula) if isinstance(jpy_msrp_formula, str) else None,
        detected_exchange_rate=detected_rate,
        exchange_rate_source_cell=rate_cell,
        values=SpaceOneValues(
            name_ja=name,
            description=description,
            manufacturer_msrp_usd=msrp,
            manufacturer_dealer_price_usd=dealer,
            sales_price=sales_price,
            notes=description,
            category=category,
        ),
    )
    return item, counter


def _sales_formula(formula_sheet, value_sheet, row_number, max_row):
    current = formula_sheet.cell(row_number, 18).value
    next_row = row_number + 1
    nxt = None
    if next_row <= max_row and value_sheet.cell(next_row, 2).value in (None, ""):
        nxt = formula_sheet.cell(next_row, 18).value
        if isinstance(nxt, str) and nxt.startswith("="):
            return nxt, next_row, formula_sheet.cell(next_row, 5).value
    if isinstance(current, str) and current.startswith("="):
        return current, row_number, formula_sheet.cell(row_number, 5).value
    if _is_sales_formula_or_fixed(nxt):
        return nxt, next_row, formula_sheet.cell(next_row, 5).value
    if _is_sales_formula_or_fixed(current):
        return current, row_number, formula_sheet.cell(row_number, 5).value
    return None, None, None


def _is_sales_formula_or_fixed(value) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return isinstance(value, str) and value.startswith("=")


def _formula_text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return str(value)
    return None


def _sheet_exchange_rate(formula_sheet, value_sheet):
    label = value_sheet["E2"].value or formula_sheet["E2"].value
    raw = value_sheet["F2"].value
    if raw is None:
        raw = formula_sheet["F2"].value
    if isinstance(raw, bool) or not isinstance(raw, (int, float)):
        return None, None
    cell = f"{formula_sheet.title}!F2"
    if label and "ドル" not in str(label) and "USD" not in str(label).upper():
        return float(raw), cell
    return float(raw), cell


def _sales_price(value_sheet, row_number, max_row) -> Optional[float]:
    current, _ = parse_money(value_sheet.cell(row_number, 18).value)
    if current is not None:
        return current
    next_row = row_number + 1
    if next_row > max_row:
        return None
    if value_sheet.cell(next_row, 2).value not in (None, ""):
        return None
    next_price, _ = parse_money(value_sheet.cell(next_row, 18).value)
    return next_price


def _is_legacy_shipping(sku, name, description) -> bool:
    blob = " ".join(part for part in (sku, name, description) if part).upper()
    return any(marker in blob for marker in SHIPPING_MARKERS)


def _is_headerish(text: Optional[str]) -> bool:
    content = (text or "").lower()
    return any(marker in content for marker in HEADERISH_SKU)


def _iferror_fallback(formula) -> Optional[float]:
    if not isinstance(formula, str):
        return None
    match = re.search(r",\s*(-?[0-9]+(?:\.[0-9]+)?)\s*\)\s*$", formula)
    if match is None:
        return None
    value, invalid = parse_money(match.group(1))
    if invalid:
        return None
    return value


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _default_name(source) -> Optional[str]:
    name = getattr(source, "name", None)
    if name:
        return Path(name).stem
    if isinstance(source, (str, Path)):
        return Path(source).stem
    return None

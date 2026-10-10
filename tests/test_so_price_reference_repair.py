import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

import agents.so_price_reference_repair as repair_module
from agents.so_price_reference_audit import ReferenceAuditStatus, audit_price_references
from agents.so_price_reference_repair import (
    RepairPreviewError,
    SourceSnapshot,
    create_repair_preview,
)
from parsers.spaceone_master_parser import parse_cell_reference
from repositories.sqlite_price_master_repository import SqlitePriceMasterRepository
from tests.test_so_price_reference_audit import _books


def _bytes(source):
    return source.getvalue()


def _snapshots(so, dt40, pt30):
    return {
        "SO_MASTER": SourceSnapshot("SO-TEST", hashlib.sha256(_bytes(so)).hexdigest()),
        "DT40": SourceSnapshot("DT40-TEST", hashlib.sha256(_bytes(dt40)).hexdigest()),
        "PT30": SourceSnapshot("PT30-TEST", hashlib.sha256(_bytes(pt30)).hexdigest()),
    }


def _preview(tmp_path):
    so, dt40, pt30 = _books()
    result = create_repair_preview(
        so,
        dt40,
        pt30,
        output_dir=tmp_path,
        snapshots=_snapshots(so, dt40, pt30),
        expected_fields=3,
        expected_rows=2,
    )
    return result, so, dt40, pt30


def test_preview_plans_only_existing_exact_references_and_preserves_formula(tmp_path):
    result, so, dt40, pt30 = _preview(tmp_path)

    assert len(result.plan) == 3
    assert {(item.so_sku, item.field) for item in result.plan} == {
        ("SKU-B", "MSRP"),
        ("SKU-B", "DEALER"),
        ("SKU-OK", "DEALER"),
    }
    assert all(item.status == "READY" for item in result.plan)
    assert all(item.current_workbook == item.expected_workbook == "DT40" for item in result.plan)
    assert all("SKU-NO-REF" != item.so_sku for item in result.plan)
    assert all("9685" != item.so_sku for item in result.plan)

    msrp = next(item for item in result.plan if item.so_sku == "SKU-B" and item.field == "MSRP")
    assert msrp.target_so_cell == "D3"
    assert msrp.current_cell == "C2"
    assert msrp.expected_cell == "C3"
    assert msrp.current_referenced_sku == "SKU-A"
    assert msrp.current_formula.startswith("=IFERROR(IMPORTRANGE(")
    assert f"/d/" in msrp.replacement_formula
    assert msrp.current_formula.split('","')[0] == msrp.replacement_formula.split('","')[0]
    assert msrp.current_fallback_value == 100
    assert msrp.fallback_matches_official is False
    assert msrp.replacement_formula.endswith(",100)")
    parsed = parse_cell_reference(msrp.replacement_formula)
    assert (parsed.workbook, parsed.sheet, parsed.cell) == ("DT40", "PHOTON", "C3")

    original = load_workbook(BytesIO(_bytes(so)), data_only=False)
    original_9685 = original["PHOTON"]["D4"].value
    original.close()
    preview = load_workbook(result.output_paths["preview"], data_only=False)
    try:
        assert preview["PHOTON"]["D3"].value == msrp.replacement_formula
        assert preview["PHOTON"]["D10"].value == 700  # NO_REFERENCE is unchanged.
        assert preview["PHOTON"]["D4"].value == original_9685  # 9685 is unchanged.
    finally:
        preview.close()

    after = audit_price_references(result.output_paths["preview"], dt40, pt30)
    after_by_key = {(row.normalized_sku, row.field): row for row in after.rows}
    assert after_by_key[("SKU-B", "MSRP")].status == ReferenceAuditStatus.CORRECT
    assert after_by_key[("SKU-NO-REF", "MSRP")].status == ReferenceAuditStatus.NO_REFERENCE
    assert after_by_key[("DUP-SKU", "MSRP")].status == ReferenceAuditStatus.AMBIGUOUS_OCCURRENCE
    assert after_by_key[("9685", "MSRP")].status == ReferenceAuditStatus.SKU_NOT_FOUND


def test_preview_logical_diff_contains_only_planned_formula_cells(tmp_path):
    result, *_ = _preview(tmp_path)

    targets = {(item.so_sheet, item.target_so_cell) for item in result.plan}
    changes = {(item["sheet"], item["cell"]) for item in result.logical_diff}
    assert changes == targets
    assert len(changes) == 3
    assert all(item["before"].startswith("=") and item["after"].startswith("=") for item in result.logical_diff)


def test_preview_is_read_only_and_never_opens_registry(tmp_path, monkeypatch):
    so, dt40, pt30 = _books()
    before = [hashlib.sha256(_bytes(source)).hexdigest() for source in (so, dt40, pt30)]

    def fail_registry(*_args, **_kwargs):
        raise AssertionError("registry write attempted")

    monkeypatch.setattr(SqlitePriceMasterRepository, "__init__", fail_registry)
    create_repair_preview(
        so,
        dt40,
        pt30,
        output_dir=tmp_path,
        snapshots=_snapshots(so, dt40, pt30),
        expected_fields=3,
        expected_rows=2,
    )

    after = [hashlib.sha256(_bytes(source)).hexdigest() for source in (so, dt40, pt30)]
    assert after == before


def test_stale_source_sha_aborts_before_preview(tmp_path):
    so, dt40, pt30 = _books()
    snapshots = _snapshots(so, dt40, pt30)
    snapshots["SO_MASTER"] = SourceSnapshot("SO-TEST", "0" * 64)

    with pytest.raises(RepairPreviewError, match="SOURCE_CHANGED"):
        create_repair_preview(
            so,
            dt40,
            pt30,
            output_dir=tmp_path,
            snapshots=snapshots,
            expected_fields=3,
            expected_rows=2,
        )
    assert not (tmp_path / "SO_MASTER_DQ5B_PREVIEW.xlsx").exists()


def test_changed_repair_count_aborts_without_partial_preview(tmp_path):
    so, dt40, pt30 = _books()

    with pytest.raises(RepairPreviewError, match="STALE_AUDIT"):
        create_repair_preview(
            so,
            dt40,
            pt30,
            output_dir=tmp_path,
            snapshots=_snapshots(so, dt40, pt30),
            expected_fields=4,
            expected_rows=2,
        )
    assert not (tmp_path / "SO_MASTER_DQ5B_PREVIEW.xlsx").exists()


def test_changed_current_formula_aborts_whole_batch(tmp_path, monkeypatch):
    so, dt40, pt30 = _books()
    stale_audit = audit_price_references(so, dt40, pt30)
    workbook = load_workbook(BytesIO(_bytes(so)), data_only=False)
    workbook["PHOTON"]["D3"] = workbook["PHOTON"]["D3"].value.replace("C2", "C4")
    changed = BytesIO()
    workbook.save(changed)
    workbook.close()
    changed.seek(0)
    snapshots = _snapshots(changed, dt40, pt30)
    monkeypatch.setattr(repair_module, "audit_price_references", lambda *_args, **_kwargs: stale_audit)

    with pytest.raises(RepairPreviewError, match="current formula changed"):
        create_repair_preview(
            changed,
            dt40,
            pt30,
            output_dir=tmp_path,
            snapshots=snapshots,
            expected_fields=3,
            expected_rows=2,
        )
    assert not (tmp_path / "SO_MASTER_DQ5B_PREVIEW.xlsx").exists()


def test_workbook_switch_aborts_batch(tmp_path, monkeypatch):
    so, dt40, pt30 = _books()
    audit = audit_price_references(so, dt40, pt30)
    target = next(row for row in audit.rows if row.status == ReferenceAuditStatus.EXACT_REPAIRABLE)
    target.expected_workbook = "PT30"
    monkeypatch.setattr(repair_module, "audit_price_references", lambda *_args, **_kwargs: audit)

    with pytest.raises(RepairPreviewError, match="WORKBOOK_SWITCH_REQUIRES_REVIEW"):
        create_repair_preview(
            so,
            dt40,
            pt30,
            output_dir=tmp_path,
            snapshots=_snapshots(so, dt40, pt30),
            expected_fields=3,
            expected_rows=2,
        )
    assert not (tmp_path / "SO_MASTER_DQ5B_PREVIEW.xlsx").exists()

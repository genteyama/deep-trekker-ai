from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from agents.activity_catalog import (
    ActivityFilter,
    ActivityRow,
    ActivitySelection,
    ActivitySort,
    apply_activity_filter,
    build_quote_preview,
    case_row,
    classify_case,
    classify_quote,
    link_quote_to_case,
    load_activity_rows,
    quotes_for_case,
    related_source_label,
    reset_activity_filter,
    reset_activity_sort,
    sort_activity_rows,
    timeline_event_label,
)
from agents.activity_log import (
    append_activity_event,
    record_case_lifecycle,
    record_case_saved,
    record_quote_approved,
    record_quote_lifecycle,
    record_quote_document_generated,
    record_quote_saved,
    record_quote_submitted,
)
from agents.activity_report import (
    activity_report_filename,
    customer_csv_filename,
    customer_records_from_rows,
    render_activity_report_docx,
    render_activity_report_pdf,
    render_customer_csv,
)
from agents.quote_pdf import extract_pdf_text
from agents.technical_case_agent import run_technical_case_analysis
from agents.technical_case_persistence import build_record
from agents.work_lifecycle import (
    archive_record,
    derive_quote,
    derive_technical_case,
    restore_quote,
    restore_record,
    soft_delete_quote,
    soft_delete_record,
)
from llm.mock_provider import MockTechnicalCaseProvider
from agents.quote_export import export_spaceone_quote_excel, export_spaceone_quote_pdf
from models import (
    ActivityCategory,
    ActivityEventType,
    CaseLineageType,
    CustomerRecord,
    ExportFileType,
    ExportPurpose,
    QuoteDraftStatus,
    TechnicalCaseRecord,
    TechnicalCaseStatus,
)
from repositories.sqlite_activity_repository import SqliteActivityRepository
from repositories.sqlite_customer_repository import SqliteCustomerRepository
from repositories.sqlite_quote_repository import SqliteQuoteRepository
from repositories.sqlite_technical_case_repository import SqliteTechnicalCaseRepository
from tests.test_quote_approval import _approve, _ready_photon
from tests.test_quote_builder import _mag_draft, _photon_draft
from ui.quote_persistence import maybe_autosave, save_draft_now
from ui.quote_steps import SESSION_QUOTE_STEP
from ui.technical_case_persistence import persist_technical_case


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"


def _row(**kwargs) -> ActivityRow:
    defaults = dict(
        row_id="case:C1",
        entity_kind="case",
        entity_id="C1",
        category=ActivityCategory.TECHNICAL_CASE.value,
        date="2026-10-01T00:00:00+00:00",
        last_activity_at="2026-10-01T00:00:00+00:00",
        customer="A社",
        end_user="現場A",
        product="MAG",
        title="案件A",
        status=TechnicalCaseStatus.ANALYZED.value,
        updated_at="2026-10-01T00:00:00+00:00",
    )
    defaults.update(kwargs)
    return ActivityRow(**defaults)


def _saved_case(repo, *, title="タンク肉厚測定", customer="IHI検査計測", end_user="End User"):
    run = run_technical_case_analysis(title, customer, end_user, "MAGで肉厚測定したい。", provider=MockTechnicalCaseProvider())
    record = build_record(run)
    repo.save_case(record)
    return repo.get_case(record.case_id)


def test_activity_list_chronological(tmp_path):
    older = _row(row_id="1", last_activity_at="2026-10-01T10:00:00+00:00", date="2026-09-01T00:00:00+00:00", title="古い")
    newer = _row(row_id="2", last_activity_at="2026-10-04T12:00:00+00:00", date="2026-09-02T00:00:00+00:00", title="新しい")
    ordered = sort_activity_rows([older, newer], reset_activity_sort())
    assert [row.title for row in ordered] == ["新しい", "古い"]
    oldest_first = sort_activity_rows([older, newer], ActivitySort(field="last_activity_at", descending=False))
    assert [row.title for row in oldest_first] == ["古い", "新しい"]


def test_filter_customer():
    rows = [_row(customer="IHI検査計測"), _row(row_id="2", customer="他社")]
    filtered = apply_activity_filter(rows, ActivityFilter(customers=("IHI検査計測",)))
    assert [row.customer for row in filtered] == ["IHI検査計測"]


def test_filter_product():
    rows = [_row(product="MAG"), _row(row_id="2", product="PHOTON")]
    filtered = apply_activity_filter(rows, ActivityFilter(products=("MAG",)))
    assert [row.product for row in filtered] == ["MAG"]


def test_filter_category():
    rows = [
        _row(category=ActivityCategory.INQUIRY.value),
        _row(row_id="2", category=ActivityCategory.QUOTE.value),
    ]
    filtered = apply_activity_filter(rows, ActivityFilter(categories=(ActivityCategory.QUOTE.value,)))
    assert [row.category for row in filtered] == [ActivityCategory.QUOTE.value]


def test_filter_status():
    rows = [_row(status="DRAFT"), _row(row_id="2", status="COMPLETED")]
    filtered = apply_activity_filter(rows, ActivityFilter(statuses=("COMPLETED",)))
    assert [row.status for row in filtered] == ["COMPLETED"]


def test_filter_date():
    rows = [
        _row(date="2026-09-01T00:00:00+00:00", updated_at="2026-10-08T00:00:00+00:00"),
        _row(row_id="2", date="2026-10-04T00:00:00+00:00", updated_at="2026-10-04T00:00:00+00:00"),
    ]
    filtered = apply_activity_filter(rows, ActivityFilter(date_from="2026-10-01", date_to="2026-10-31"))
    assert [row.row_id for row in filtered] == ["2"]


def test_reset_filters():
    filt = ActivityFilter(customers=("A社",), keyword="MAG", include_archived=True)
    reset = reset_activity_filter()
    assert reset == ActivityFilter()
    assert reset_activity_sort().field == "last_activity_at"
    assert reset_activity_sort().descending is True
    rows = [_row(), _row(row_id="2", customer="B社")]
    assert apply_activity_filter(rows, reset) == rows


def test_select_one_many_all_filtered_and_deselect():
    rows = [_row(row_id="a"), _row(row_id="b"), _row(row_id="c", customer="B社")]
    selection = ActivitySelection()
    selection.select_one("a")
    assert selection.selected == {"a"}
    selection.select_many(["b"])
    assert selection.selected == {"a", "b"}
    filtered = apply_activity_filter(rows, ActivityFilter(customers=("A社",)))
    selection.select_all(filtered)
    assert selection.selected == {"a", "b"}
    selection.deselect_all()
    assert selection.selected == set()


def test_word_export(tmp_path):
    rows = [_row(customer="IHI検査計測", title="IHI タンク肉厚測定")]
    path = tmp_path / activity_report_filename(".docx", now="20261004")
    render_activity_report_docx(rows, path, filt=ActivityFilter(customers=("IHI検査計測",)))
    assert path.name == "DeepTrekker_Activity_Report_20261004.docx"
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Deep Trekker / PipeTrekker 対応履歴レポート" in text
    assert "選択件数: 1" in text
    assert "IHI検査計測" in text
    headers = [cell.text for cell in document.tables[0].rows[0].cells]
    assert headers == ["No.", "日付", "顧客", "End User", "区分", "機種", "案件名", "対応ステータス"]


def test_pdf_export(tmp_path):
    rows = [_row(customer="IHI検査計測", title="IHI タンク肉厚測定")]
    path = tmp_path / activity_report_filename(".pdf", now="20261004")
    render_activity_report_pdf(rows, path)
    assert path.name == "DeepTrekker_Activity_Report_20261004.pdf"
    text = extract_pdf_text(path)
    assert "対応履歴レポート" in text
    assert "IHI検査計測" in text
    assert "選択件数" in text


def test_customer_csv_export_and_japanese(tmp_path):
    rows = [
        _row(customer="株式会社ＩＨＩ", end_user="現場A"),
        _row(row_id="2", customer="株式会社ＩＨＩ", end_user="現場A"),
        _row(row_id="3", customer="他社"),
    ]
    records = customer_records_from_rows(rows)
    assert [item.customer_name for item in records] == ["他社", "株式会社ＩＨＩ"]
    text, path = render_customer_csv(records, tmp_path / customer_csv_filename(now="20261004"))
    assert path.name == "DeepTrekker_Customers_20261004.csv"
    assert text.startswith("\ufeff")
    assert "会社名,部署,担当者名,メール,電話,住所,End User,備考" in text
    assert "株式会社ＩＨＩ" in text
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")


def test_soft_deleted_excluded(tmp_path):
    db = tmp_path / "ledger.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    record = _saved_case(cases)
    cases.save_case(soft_delete_record(record))
    rows = apply_activity_filter(load_activity_rows(cases, quotes, events))
    assert rows == []
    history = cases.list_recent_cases(limit=20, view="history")
    assert history == []
    trash = cases.list_recent_cases(limit=20, view="trash")
    assert trash[0].case_id == record.case_id


def test_archived_filtering(tmp_path):
    db = tmp_path / "ledger.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    record = _saved_case(cases)
    cases.save_case(archive_record(record))
    rows = load_activity_rows(cases, quotes, events)
    assert apply_activity_filter(rows) == []
    shown = apply_activity_filter(rows, ActivityFilter(include_archived=True))
    assert shown[0].entity_id == record.case_id
    assert shown[0].archived is True


def test_derived_relationship_shown(tmp_path):
    db = tmp_path / "ledger.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    parent = _saved_case(cases)
    child = derive_technical_case(parent, relation_type="FOLLOW_UP")
    cases.save_case(child)
    rows = load_activity_rows(cases, quotes, events)
    derived = next(row for row in rows if row.entity_id == child.case_id)
    assert derived.category == ActivityCategory.DERIVED_CASE.value
    assert related_source_label(derived, rows) == parent.case_title


def test_technical_case_quote_relationship():
    case = TechnicalCaseRecord(case_id="CASE-IHI", customer_name="IHI検査計測", case_title="IHI タンク肉厚測定")
    first = link_quote_to_case(_mag_draft(), case.case_id)
    second = link_quote_to_case(_photon_draft(), case.case_id)
    related = quotes_for_case([first, second, _photon_draft()], case.case_id)
    assert {item.configuration_name for item in related} == {"MAG", "PHOTON"}
    assert first.case_id == "CASE-IHI"
    assert second.case_id == "CASE-IHI"


def test_activity_event_append_only(tmp_path):
    repo = SqliteActivityRepository(tmp_path / "events.sqlite3")
    first = append_activity_event(repo, event_type=ActivityEventType.CASE_CREATED.value, entity_kind="case", entity_id="C1", payload={"api_key": "secret"})
    second = append_activity_event(repo, event_type=ActivityEventType.CASE_UPDATED.value, entity_kind="case", entity_id="C1")
    with pytest.raises(RuntimeError, match="append-only"):
        repo.update(first)
    events = repo.list_events(entity_kind="case", entity_id="C1")
    assert [item.event_id for item in events] == [second.event_id, first.event_id]
    assert events[-1].event_type == ActivityEventType.CASE_CREATED.value
    assert "api_key" not in first.payload
    count = repo._connection.execute("SELECT COUNT(*) AS n FROM activity_events").fetchone()["n"]
    assert count == 2


def test_no_api_secrets_exported(tmp_path):
    rows = [_row(customer="A社", title="thinking leak check")]
    repo = SqliteCustomerRepository(tmp_path / "cust.sqlite3")
    repo.save(
        CustomerRecord(
            customer_id="CUS-1",
            customer_name="A社",
            notes="api_key=AIzaSyDummyKeyValue1234567890 thinking=secret",
        )
    )
    records = customer_records_from_rows(rows, repo)
    text, _ = render_customer_csv(records)
    word = tmp_path / "secret.docx"
    pdf = tmp_path / "secret.pdf"
    render_activity_report_docx(rows, word)
    render_activity_report_pdf(rows, pdf)
    pdf_text = extract_pdf_text(pdf)
    word_text = "\n".join(paragraph.text for paragraph in Document(str(word)).paragraphs)
    assert "AIzaSyDummyKeyValue1234567890" not in text
    assert "AIza" not in text
    assert "AIza" not in word_text
    assert "AIza" not in pdf_text
    assert "sk-ant-" not in pdf_text
    assert "reasoning" not in word_text


def test_existing_quote_and_technical_case_persistence_unchanged(tmp_path):
    db = tmp_path / "persist.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    record = _saved_case(cases)
    before = record.model_dump()
    reload_case = cases.get_case(record.case_id)
    assert reload_case.model_dump() == before
    draft = _photon_draft()
    quotes.save_draft(draft)
    loaded = quotes.get_draft(draft.quote_draft_id, draft.quote_version)
    assert loaded.draft.quote_draft_id == draft.quote_draft_id
    assert loaded.draft.customer == draft.customer
    assert loaded.draft.total_jpy == draft.total_jpy
    assert loaded.draft.configuration_lines == draft.configuration_lines


def test_last_activity_uses_event_not_created_at(tmp_path, monkeypatch):
    # save_case() stamps updated_at with the current time; pin it so the event stays newest on any run date.
    monkeypatch.setattr(
        "repositories.sqlite_technical_case_repository.now_iso",
        lambda: "2026-09-02T00:00:00+00:00",
    )
    db = tmp_path / "ledger.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    record = TechnicalCaseRecord(
        case_id="CASE-DATE",
        customer_name="IHI検査計測",
        case_title="日付確認",
        created_at="2026-09-01T00:00:00+00:00",
        updated_at="2026-09-02T00:00:00+00:00",
        status=TechnicalCaseStatus.ANALYZED.value,
        inquiry_success=True,
    )
    cases.save_case(record)
    append_activity_event(
        events,
        event_type=ActivityEventType.MANUFACTURER_RESPONSE_RECEIVED.value,
        entity_kind="case",
        entity_id=record.case_id,
        occurred_at="2026-10-04T09:00:00+00:00",
    )
    rows = sort_activity_rows(load_activity_rows(cases, quotes, events))
    assert rows[0].entity_id == record.case_id
    assert rows[0].updated_at == "2026-09-02T00:00:00+00:00"
    assert rows[0].last_activity_at.startswith("2026-10-04")
    assert rows[0].date.startswith("2026-09-01")


def test_quote_preview_uses_snapshot_not_current_master():
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    preview = build_quote_preview(snapshot)
    original_total = preview["total"]
    draft.total_jpy = 1
    draft.customer_lines[0].unit_price_jpy = 1
    assert build_quote_preview(snapshot)["total"] == original_total
    assert preview["quote_number"] == snapshot.official_quote_number or snapshot.quote_number_candidate
    assert preview["status"] == QuoteDraftStatus.APPROVED.value


def test_classify_inquiry_and_revision():
    inquiry = TechnicalCaseRecord(case_id="INQ", status=TechnicalCaseStatus.DRAFT.value, inquiry_success=False)
    analyzed = TechnicalCaseRecord(case_id="AN", status=TechnicalCaseStatus.ANALYZED.value, inquiry_success=True)
    assert classify_case(inquiry) == ActivityCategory.INQUIRY.value
    assert classify_case(analyzed) == ActivityCategory.TECHNICAL_CASE.value
    revision = derive_quote(_photon_draft(), relation_type="REVISION")
    assert classify_quote(revision) == ActivityCategory.REVISION.value


def test_persist_hooks_record_major_events_only(tmp_path):
    db = tmp_path / "hooks.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    events = SqliteActivityRepository(db)
    run = run_technical_case_analysis("案件", "顧客", "EU", "MAGで確認したい。", provider=MockTechnicalCaseProvider())
    session = {"technical_case_run": run}
    first = persist_technical_case(session, run, repository=cases)
    second = persist_technical_case(session, run, repository=cases)
    logged = events.list_events(entity_kind="case", entity_id=first.case_id)
    assert [item.event_type for item in logged] == [ActivityEventType.CASE_CREATED.value]
    assert second.case_id == first.case_id
    record_case_saved(events, first, first.model_copy(update={"status": TechnicalCaseStatus.COMPLETED.value}), completed=True)
    assert [item.event_type for item in events.list_events(entity_kind="case", entity_id=first.case_id)][0] == ActivityEventType.CASE_UPDATED.value


def test_quote_create_and_approve_events(tmp_path):
    db = tmp_path / "hooks.sqlite3"
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    draft = _ready_photon()
    record_quote_saved(events, None, draft)
    quotes.save_draft(draft)
    _approval, snapshot = _approve(draft)
    quotes.save_snapshot(snapshot)
    record_quote_approved(events, draft, snapshot)
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=draft.quote_draft_id)]
    assert ActivityEventType.QUOTE_CREATED.value in types
    assert ActivityEventType.QUOTE_APPROVED.value in types


def test_activity_ledger_page_opens():
    at = AppTest.from_file(str(APP_PATH)).run()
    at.button(key="open_activity_ledger").click().run()
    visible = " ".join(
        [getattr(item, "value", "") for item in at.text]
        + [getattr(item, "value", "") for item in at.markdown]
        + [getattr(item, "value", "") for item in at.caption]
        + [getattr(item, "value", "") for item in getattr(at, "title", [])]
    )
    assert "時系列で確認" in visible
    assert at.button(key="activity_reset_filters").label == "フィルターをリセット"
    assert at.button(key="activity_reset_filters")
    assert at.button(key="activity_select_all")
    assert at.button(key="activity_deselect_all")
    assert at.button(key="activity_export_word")
    assert at.button(key="activity_export_pdf")


def test_quote_approved_and_submitted_are_separate_events(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    approved = record_quote_approved(events, draft, snapshot)
    submitted = record_quote_submitted(events, draft, submitted_at="2026-10-04T12:00:00+00:00")
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=draft.quote_draft_id)]
    assert approved.event_type == ActivityEventType.QUOTE_APPROVED.value
    assert submitted.event_type == ActivityEventType.QUOTE_SUBMITTED.value
    assert approved.event_type != submitted.event_type
    assert ActivityEventType.QUOTE_APPROVED.value in types
    assert ActivityEventType.QUOTE_SUBMITTED.value in types
    assert submitted.payload["quote_draft_id"] == draft.quote_draft_id
    assert submitted.payload["quote_version"] == draft.quote_version
    assert submitted.payload["customer_name"] == draft.customer
    assert submitted.payload["submitted_at"] == "2026-10-04T12:00:00+00:00"
    assert "api_key" not in submitted.payload
    assert "thinking" not in submitted.payload


def test_quote_submitted_is_recorded_only_once(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _photon_draft()
    first = record_quote_submitted(events, draft)
    second = record_quote_submitted(events, draft)
    third = record_quote_submitted(events, draft)
    logged = [item for item in events.list_events(entity_kind="quote", entity_id=draft.quote_draft_id) if item.event_type == ActivityEventType.QUOTE_SUBMITTED.value]
    assert first is not None
    assert second is None
    assert third is None
    assert len(logged) == 1


def test_autosave_does_not_duplicate_submitted_event(tmp_path):
    db = tmp_path / "hooks.sqlite3"
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    draft = _photon_draft()
    session = {SESSION_QUOTE_STEP: 2}
    save_draft_now(quotes, draft, session)
    record_quote_submitted(events, draft)
    save_draft_now(quotes, draft, session)
    maybe_autosave(quotes, draft, session)
    maybe_autosave(quotes, draft, session)
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=draft.quote_draft_id)]
    assert types.count(ActivityEventType.QUOTE_SUBMITTED.value) == 1
    assert types.count(ActivityEventType.QUOTE_CREATED.value) == 1
    assert ActivityEventType.QUOTE_APPROVED.value not in types


def test_derived_case_event(tmp_path):
    db = tmp_path / "hooks.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    events = SqliteActivityRepository(db)
    parent = _saved_case(cases)
    child = derive_technical_case(parent, relation_type=CaseLineageType.FOLLOW_UP.value)
    cases.save_case(child)
    first = record_case_lifecycle(events, child, "derive", previous=parent)
    second = record_case_lifecycle(events, child, "derive", previous=parent)
    logged = events.list_events(entity_kind="case", entity_id=child.case_id)
    assert first is not None
    assert second is None
    assert [item.event_type for item in logged] == [ActivityEventType.CASE_DERIVED.value]
    assert first.payload["parent_case_id"] == parent.case_id
    assert first.payload["new_case_id"] == child.case_id
    assert first.payload["relation_type"] == CaseLineageType.FOLLOW_UP.value
    assert ActivityEventType.CASE_CREATED.value not in [item.event_type for item in logged]


def test_move_to_trash_event(tmp_path):
    db = tmp_path / "hooks.sqlite3"
    cases = SqliteTechnicalCaseRepository(db)
    quotes = SqliteQuoteRepository(db)
    events = SqliteActivityRepository(db)
    record = _saved_case(cases)
    trashed = soft_delete_record(record)
    first = record_case_lifecycle(events, trashed, "trash", previous=record)
    second = record_case_lifecycle(events, trashed, "trash", previous=trashed)
    assert first.event_type == ActivityEventType.MOVED_TO_TRASH.value
    assert second is None
    draft = _photon_draft()
    quotes.save_draft(draft)
    deleted = soft_delete_quote(draft)
    quote_first = record_quote_lifecycle(events, deleted, "trash", previous=draft)
    quote_again = record_quote_lifecycle(events, deleted, "trash", previous=deleted)
    assert quote_first.event_type == ActivityEventType.MOVED_TO_TRASH.value
    assert quote_again is None


def test_restore_from_trash_event(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    record = TechnicalCaseRecord(case_id="CASE-TRASH", customer_name="A社", case_title="ゴミ箱確認")
    trashed = soft_delete_record(record)
    record_case_lifecycle(events, trashed, "trash", previous=record)
    restored = restore_record(trashed)
    first = record_case_lifecycle(events, restored, "restore", previous=trashed)
    second = record_case_lifecycle(events, restored, "restore", previous=restored)
    types = [item.event_type for item in events.list_events(entity_kind="case", entity_id=record.case_id)]
    assert first.event_type == ActivityEventType.RESTORED_FROM_TRASH.value
    assert second is None
    assert ActivityEventType.RESTORED.value not in types
    draft = _photon_draft()
    deleted = soft_delete_quote(draft)
    record_quote_lifecycle(events, deleted, "trash", previous=draft)
    quote_restored = restore_quote(deleted)
    quote_event = record_quote_lifecycle(events, quote_restored, "restore", previous=deleted)
    assert quote_event.event_type == ActivityEventType.RESTORED_FROM_TRASH.value


def test_timeline_japanese_labels():
    assert timeline_event_label(ActivityEventType.QUOTE_APPROVED.value) == "見積承認"
    assert timeline_event_label(ActivityEventType.QUOTE_SUBMITTED.value) == "見積提出"
    assert timeline_event_label(ActivityEventType.QUOTE_DOCUMENT_GENERATED.value) == "正式見積書作成"
    assert timeline_event_label(ActivityEventType.CASE_DERIVED.value) == "派生案件作成"
    assert timeline_event_label(ActivityEventType.MOVED_TO_TRASH.value) == "ゴミ箱へ移動"
    assert timeline_event_label(ActivityEventType.RESTORED_FROM_TRASH.value) == "ゴミ箱から復元"
    assert timeline_event_label(ActivityEventType.QUOTE_APPROVED.value) != timeline_event_label(ActivityEventType.QUOTE_SUBMITTED.value)
    assert timeline_event_label(ActivityEventType.QUOTE_DOCUMENT_GENERATED.value) != timeline_event_label(ActivityEventType.QUOTE_SUBMITTED.value)


def test_formal_pdf_generation_is_not_quote_submitted(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    export_spaceone_quote_pdf(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        generated_by="弦",
        purpose=ExportPurpose.FORMAL,
    )
    recorded = record_quote_document_generated(
        events,
        snapshot,
        document_type=ExportFileType.SPACEONE_QUOTE_PDF.value,
    )
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=snapshot.quote_draft_id)]
    assert recorded.event_type == ActivityEventType.QUOTE_DOCUMENT_GENERATED.value
    assert ActivityEventType.QUOTE_SUBMITTED.value not in types
    assert "record_quote_submitted" not in Path("ui/quote_control.py").read_text(encoding="utf-8")


def test_formal_excel_generation_is_not_quote_submitted(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        generated_by="弦",
        purpose=ExportPurpose.FORMAL,
    )
    recorded = record_quote_document_generated(
        events,
        snapshot,
        document_type=ExportFileType.SPACEONE_QUOTE_XLSX.value,
    )
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=snapshot.quote_draft_id)]
    assert recorded.event_type == ActivityEventType.QUOTE_DOCUMENT_GENERATED.value
    assert recorded.payload["document_type"] == ExportFileType.SPACEONE_QUOTE_XLSX.value
    assert ActivityEventType.QUOTE_SUBMITTED.value not in types
    assert "path" not in recorded.payload
    assert "api_key" not in recorded.payload


def test_formal_document_generation_is_quote_document_generated(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    pdf = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    excel = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_XLSX.value)
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=snapshot.quote_draft_id)]
    assert types.count(ActivityEventType.QUOTE_DOCUMENT_GENERATED.value) == 2
    assert {pdf.payload["document_type"], excel.payload["document_type"]} == {
        ExportFileType.SPACEONE_QUOTE_PDF.value,
        ExportFileType.SPACEONE_QUOTE_XLSX.value,
    }
    assert ActivityEventType.QUOTE_SUBMITTED.value not in types


def test_explicit_submit_only_creates_quote_submitted(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    record_quote_approved(events, draft, snapshot)
    record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    submitted = record_quote_submitted(events, draft)
    types = [item.event_type for item in events.list_events(entity_kind="quote", entity_id=draft.quote_draft_id)]
    assert submitted.event_type == ActivityEventType.QUOTE_SUBMITTED.value
    assert types.count(ActivityEventType.QUOTE_SUBMITTED.value) == 1
    assert ActivityEventType.QUOTE_DOCUMENT_GENERATED.value in types
    assert ActivityEventType.QUOTE_APPROVED.value in types


def test_formal_document_redownload_does_not_duplicate_event(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    first = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    second = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    excel = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_XLSX.value)
    events_for = events.list_events(entity_kind="quote", entity_id=snapshot.quote_draft_id)
    pdf_events = [item for item in events_for if item.payload.get("document_type") == ExportFileType.SPACEONE_QUOTE_PDF.value]
    assert first is not None
    assert second is None
    assert excel is not None
    assert len(pdf_events) == 1


def test_new_quote_version_document_is_separate_event(tmp_path):
    events = SqliteActivityRepository(tmp_path / "events.sqlite3")
    draft = _ready_photon()
    _approval, snapshot = _approve(draft)
    first = record_quote_document_generated(events, snapshot, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    revision = snapshot.model_copy(update={"quote_version": snapshot.quote_version + 1})
    second = record_quote_document_generated(events, revision, document_type=ExportFileType.SPACEONE_QUOTE_PDF.value)
    assert first is not None
    assert second is not None
    assert first.entity_version != second.entity_version
    assert first.payload["quote_version"] == snapshot.quote_version
    assert second.payload["quote_version"] == snapshot.quote_version + 1
    pdf_events = [
        item
        for item in events.list_events(entity_kind="quote", entity_id=snapshot.quote_draft_id)
        if item.event_type == ActivityEventType.QUOTE_DOCUMENT_GENERATED.value
    ]
    assert len(pdf_events) == 2

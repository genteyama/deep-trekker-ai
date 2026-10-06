import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from agents.price_master import get_active_master, import_price_master, validate_price_master
from models import (
    PriceMasterImport,
    PriceMasterImportOutcome,
    PriceMasterImportStatus,
    PriceMasterType,
    PriceMasterValidationStatus,
)
from repositories.sqlite import now_iso
from repositories.sqlite_price_master_repository import SqlitePriceMasterRepository
from tests.price_book_fixtures import (
    build_workbook,
    official_dt40_style_book,
    official_ihi_sku_snapshot_book,
    official_pt30_style_book,
    pricing_policy_master_book,
    quote_calc_formula_book,
    spaceone_master_book,
    valid_price_book,
)
from tests.test_price_master_ui import _open_quote_page, _texts, uploads  # noqa: F401  (fixture)
from ui.price_master import outcome_message

LOCALE = json.loads((Path(__file__).resolve().parents[1] / "locales" / "ja.json").read_text(encoding="utf-8"))
LABELS = LOCALE["pages"]["quote_control"]["price_masters"]


def _bytes(factory):
    return factory().getvalue()


def _mixed_dt40_pt30_book():
    return build_workbook(
        {
            "PHOTON": [["Part Number", "Description", "MSRP", "DT40", "Notes:"], ["9680-BASE", "PHOTON", 17391, 10434.6, None]],
            "A-200": [["Part Number", "Description", "MSRP", "PT30", "Notes:"], ["A200-BASE", "A-200", 40000, 28000, None]],
        }
    )


def _dt40_plus_generic_sheet_book():
    return build_workbook(
        {
            "PHOTON": [["Part Number", "Description", "MSRP", "DT40", "Notes:"], ["9680-BASE", "PHOTON", 17391, 10434.6, None]],
            "MISC": [["Part Number", "Description", "MSRP", "Dealer Price", "Notes"], ["X-1", "Unknown", 100, 60, None]],
        }
    )


def _quote_calc_variant(*, drop_insurance=False, drop_markup=False, drop_tax=False):
    workbook = load_workbook(quote_calc_formula_book())
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                formula = isinstance(value, str) and value.startswith("=H")
                if drop_insurance and cell.column_letter == "K" and formula:
                    cell.value = None
                if drop_tax and cell.column_letter == "I" and formula:
                    cell.value = None
                if drop_markup and cell.column_letter in ("F", "G") and (
                    isinstance(value, (int, float)) or (isinstance(value, str) and "1.2" in value)
                ):
                    cell.value = None
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


# A / B -----------------------------------------------------------------------------------------


@pytest.mark.parametrize("factory", [official_ihi_sku_snapshot_book, official_dt40_style_book])
def test_dt40_slot_accepts_a_dt40_book(factory):
    outcome = import_price_master(PriceMasterType.DT40, "anything.xlsx", _bytes(factory))
    assert outcome.status == PriceMasterImportStatus.ACTIVATED


def test_pt30_slot_accepts_a_pt30_book():
    outcome = import_price_master(PriceMasterType.PT30, "anything.xlsx", _bytes(official_pt30_style_book))
    assert outcome.status == PriceMasterImportStatus.ACTIVATED


# C / D -----------------------------------------------------------------------------------------


def test_dt40_slot_rejects_a_pt30_book_and_keeps_the_current_dt40():
    current = import_price_master(PriceMasterType.DT40, "dt40.xlsx", _bytes(official_ihi_sku_snapshot_book)).record

    outcome = import_price_master(PriceMasterType.DT40, "DT40.xlsx", _bytes(official_pt30_style_book))

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "MASTER_TYPE_UNCONFIRMED"
    assert get_active_master(PriceMasterType.DT40).record.import_id == current.import_id
    assert len(SqlitePriceMasterRepository().list_imports(PriceMasterType.DT40)) == 1


def test_pt30_slot_rejects_a_dt40_book_and_keeps_the_current_pt30():
    current = import_price_master(PriceMasterType.PT30, "pt30.xlsx", _bytes(official_pt30_style_book)).record

    outcome = import_price_master(PriceMasterType.PT30, "PT30.xlsx", _bytes(official_dt40_style_book))

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "MASTER_TYPE_UNCONFIRMED"
    assert get_active_master(PriceMasterType.PT30).record.import_id == current.import_id


# E ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize("factory", [valid_price_book, _mixed_dt40_pt30_book, _dt40_plus_generic_sheet_book])
@pytest.mark.parametrize("master_type", [PriceMasterType.DT40, PriceMasterType.PT30])
def test_manufacturer_book_that_cannot_be_identified_is_rejected(master_type, factory):
    outcome = import_price_master(master_type, f"{master_type.value}.xlsx", _bytes(factory))

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "MASTER_TYPE_UNCONFIRMED"
    assert get_active_master(master_type) is None


def test_identity_comes_from_the_workbook_not_the_file_name():
    assert import_price_master(PriceMasterType.DT40, "PT30 価格表.xlsx", _bytes(official_dt40_style_book)).status == (
        PriceMasterImportStatus.ACTIVATED
    )
    assert import_price_master(PriceMasterType.PT30, "DT40 価格表.xlsx", _bytes(official_pt30_style_book)).status == (
        PriceMasterImportStatus.ACTIVATED
    )


# SO_MASTER wrong type --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "factory",
    [official_ihi_sku_snapshot_book, official_dt40_style_book, official_pt30_style_book, quote_calc_formula_book, valid_price_book],
)
def test_so_master_slot_rejects_other_files_and_keeps_the_current_master(factory):
    current = import_price_master(PriceMasterType.SO_MASTER, "so.xlsx", _bytes(pricing_policy_master_book)).record

    outcome = import_price_master(PriceMasterType.SO_MASTER, "SO_MASTER.xlsx", _bytes(factory))

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert get_active_master(PriceMasterType.SO_MASTER).record.import_id == current.import_id


def test_so_master_slot_accepts_spaceone_masters():
    for factory in (pricing_policy_master_book, spaceone_master_book):
        assert import_price_master(PriceMasterType.SO_MASTER, "so.xlsx", _bytes(factory)).status in {
            PriceMasterImportStatus.ACTIVATED,
            PriceMasterImportStatus.ALREADY_ACTIVE,
        }


# QUOTE_CALC ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "variant, missing",
    [
        ({"drop_insurance": True}, "insurance_rate"),
        ({"drop_markup": True}, "shipping_markup"),
        ({"drop_tax": True}, "import_tax_rate"),
    ],
)
def test_incomplete_quote_calc_is_rejected_and_keeps_the_current_policy(variant, missing):
    current = import_price_master(PriceMasterType.QUOTE_CALC, "calc.xlsx", _bytes(quote_calc_formula_book)).record

    outcome = import_price_master(PriceMasterType.QUOTE_CALC, "calc-v2.xlsx", _quote_calc_variant(**variant).getvalue())

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "POLICY_INCOMPLETE"
    assert get_active_master(PriceMasterType.QUOTE_CALC).record.import_id == current.import_id


def test_quote_calc_requires_all_three_values_as_real_numbers(tmp_path, monkeypatch):
    from models import LandedCostPolicyCandidate, QuoteCalcAudit
    from agents import price_master

    path = tmp_path / "calc.xlsx"
    path.write_bytes(_bytes(quote_calc_formula_book))
    for bad in (None, float("nan"), float("inf"), True):
        policy = LandedCostPolicyCandidate(
            landed_cost_policy_candidate_id="x", import_tax_rate=0.1, insurance_rate=0.03, shipping_markup_multiplier=1.2
        ).model_copy(update={"insurance_rate": bad})
        monkeypatch.setattr(price_master, "extract_quote_calc_audit", lambda _path, policy=policy: QuoteCalcAudit(policy_candidate=policy))
        with pytest.raises(price_master.PriceMasterValidationError) as raised:
            validate_price_master(PriceMasterType.QUOTE_CALC, path)
        assert raised.value.code == "POLICY_INCOMPLETE"


def test_complete_quote_calc_is_accepted():
    outcome = import_price_master(PriceMasterType.QUOTE_CALC, "calc.xlsx", _bytes(quote_calc_formula_book))

    assert outcome.status == PriceMasterImportStatus.ACTIVATED
    assert outcome.record.validation_summary == {"import_tax_rate": 0.1, "insurance_rate": 0.03, "shipping_markup": 1.2}


# SHA dedupe and reactivation -------------------------------------------------------------------


def test_reactivation_of_an_old_import_is_revalidated_against_current_rules():
    current = import_price_master(PriceMasterType.DT40, "dt40.xlsx", _bytes(official_ihi_sku_snapshot_book)).record
    repo = SqlitePriceMasterRepository()
    generic = _bytes(valid_price_book)
    import hashlib

    legacy = PriceMasterImport(
        import_id="DT40-legacy-generic",
        master_type=PriceMasterType.DT40,
        original_filename="old.xlsx",
        stored_path="DT40/DT40-legacy-generic.xlsx",
        sha256=hashlib.sha256(generic).hexdigest(),
        imported_at=now_iso(),
        validation_status=PriceMasterValidationStatus.VALID,
    )
    (repo.storage_root / "DT40" / "DT40-legacy-generic.xlsx").write_bytes(generic)
    repo.add(legacy)

    outcome = import_price_master(PriceMasterType.DT40, "old.xlsx", generic)

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "MASTER_TYPE_UNCONFIRMED"
    assert get_active_master(PriceMasterType.DT40).record.import_id == current.import_id


# UI --------------------------------------------------------------------------------------------


def test_wrong_type_message_is_plain_japanese():
    outcome = PriceMasterImportOutcome(
        status=PriceMasterImportStatus.REJECTED, master_type=PriceMasterType.DT40, reason_code="MASTER_TYPE_UNCONFIRMED"
    )
    level, text = outcome_message(LABELS, outcome)

    assert level == "error"
    assert text.startswith("DT40（Deep Trekker 価格表）として確認できませんでした。選択したファイルの種類をご確認ください")
    assert "現在の価格マスターは変更していません" in text
    incomplete = outcome.model_copy(update={"master_type": PriceMasterType.QUOTE_CALC, "reason_code": "POLICY_INCOMPLETE"})
    assert "輸入税・保険・Shipping倍率" in outcome_message(LABELS, incomplete)[1]


def test_pt30_file_in_the_dt40_slot_is_refused_in_the_ui(uploads):  # noqa: F811
    from tests.test_price_master_ui import _Upload

    current = import_price_master(PriceMasterType.DT40, "dt40.xlsx", _bytes(official_ihi_sku_snapshot_book)).record
    uploads["price_master_upload_DT40"] = _Upload("DT40 最新.xlsx", _bytes(official_pt30_style_book))
    at = _open_quote_page()
    at.button(key="price_master_import_DT40").click().run()

    assert not at.exception
    errors = " ".join(item.value for item in at.error)
    assert "DT40（Deep Trekker 価格表）として確認できませんでした" in errors
    assert "Traceback" not in _texts(at) and "Error" not in errors
    assert get_active_master(PriceMasterType.DT40).record.import_id == current.import_id


# QUOTE_CALC arithmetic validity ----------------------------------------------------------------


def _policy_with(**values):
    from models import LandedCostPolicyCandidate

    base = {"import_tax_rate": 0.1, "insurance_rate": 0.03, "shipping_markup_multiplier": 1.2}
    base.update(values)
    return LandedCostPolicyCandidate(landed_cost_policy_candidate_id="x", **base)


def _distinct_quote_calc_bytes():
    # An extra sheet guarantees different bytes from the active master, so dedupe never short-circuits.
    workbook = load_workbook(quote_calc_formula_book())
    workbook.create_sheet("memo")["A1"] = "different upload"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _use_policy(monkeypatch, policy):
    from agents import price_master
    from models import QuoteCalcAudit

    monkeypatch.setattr(price_master, "extract_quote_calc_audit", lambda _path: QuoteCalcAudit(policy_candidate=policy))


@pytest.mark.parametrize(
    "values",
    [
        {"import_tax_rate": -0.1},
        {"insurance_rate": -0.03},
        {"shipping_markup_multiplier": 0.0},
        {"shipping_markup_multiplier": -1.2},
    ],
)
def test_quote_calc_with_unusable_values_is_rejected_and_keeps_the_current_policy(monkeypatch, values):
    current = import_price_master(PriceMasterType.QUOTE_CALC, "calc.xlsx", _bytes(quote_calc_formula_book)).record

    _use_policy(monkeypatch, _policy_with(**values))
    outcome = import_price_master(PriceMasterType.QUOTE_CALC, "calc-bad.xlsx", _distinct_quote_calc_bytes())

    assert outcome.status == PriceMasterImportStatus.REJECTED
    assert outcome.reason_code == "VALIDATION_FAILED"
    assert outcome.record is None
    assert get_active_master(PriceMasterType.QUOTE_CALC).record.import_id == current.import_id
    level, text = outcome_message(LABELS, outcome)
    assert level == "error" and "現在の価格マスターは変更していません" in text and "Traceback" not in text


@pytest.mark.parametrize("bad", [None, float("nan"), float("inf"), True, "0.1"])
@pytest.mark.parametrize("field", ["import_tax_rate", "insurance_rate", "shipping_markup_multiplier"])
def test_quote_calc_non_numeric_values_are_rejected(tmp_path, monkeypatch, field, bad):
    from agents import price_master
    from models import QuoteCalcAudit

    policy = _policy_with().model_copy(update={field: bad})
    monkeypatch.setattr(price_master, "extract_quote_calc_audit", lambda _path: QuoteCalcAudit(policy_candidate=policy))
    path = tmp_path / "calc.xlsx"
    path.write_bytes(_bytes(quote_calc_formula_book))

    with pytest.raises(price_master.PriceMasterValidationError):
        validate_price_master(PriceMasterType.QUOTE_CALC, path)


def test_zero_tax_and_zero_insurance_are_accepted(monkeypatch):
    _use_policy(monkeypatch, _policy_with(import_tax_rate=0.0, insurance_rate=0.0, shipping_markup_multiplier=0.8))

    outcome = import_price_master(PriceMasterType.QUOTE_CALC, "calc-zero.xlsx", _bytes(quote_calc_formula_book))

    assert outcome.status == PriceMasterImportStatus.ACTIVATED
    assert outcome.record.validation_summary == {"import_tax_rate": 0.0, "insurance_rate": 0.0, "shipping_markup": 0.8}

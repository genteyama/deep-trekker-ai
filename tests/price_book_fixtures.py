from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures" / "excel"


def build_workbook(sheets: dict) -> BytesIO:
    workbook = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            worksheet = workbook.active
            worksheet.title = sheet_name
            first = False
        else:
            worksheet = workbook.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def valid_price_book() -> BytesIO:
    return build_workbook(
        {
            "MAG": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["9701-MAG-4K", "MAG Utility Crawler 4K", 25000, 15000, "Bare steel surface"],
                ["9735", "MAG accessory", 800, 480, None],
            ],
            "PHOTON": [
                ["SKU", "Description", "MSRP", "Dealer Price", "Notes"],
                ["7851-PHOTON", "PHOTON ROV", 8000, 4800, None],
            ],
            "PipeTrekker": [
                ["Part #", "Description", "MSRP", "Dealer Price", "Notes"],
                ["9680-BASE", "PipeTrekker Base", 40000, 28000, None],
            ],
            "Cygnus": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["5608", "Cygnus Thickness Gauge", 5000, 5000, "NO DEALER DISCOUNT"],
            ],
            "Bundle Pack": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["8459", "Unknown bundle item", 100, 90, None],
            ],
        }
    )


def shifted_columns_price_book() -> BytesIO:
    return build_workbook(
        {
            "MAG": [
                [None, None, None, None, None],
                [None, None, None, None, None],
                [None, None, "Ignore", "Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                [None, None, "skip", "9701-MAG-4K", "MAG Utility Crawler 4K", 25000, 15000, "Bare steel surface"],
                [None, None, "skip", "9735", "MAG accessory", 800, 480, None],
            ]
        }
    )


def validation_price_book() -> BytesIO:
    return build_workbook(
        {
            "MAG": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                [None, "Missing SKU crawler", 1000, 600, None],
                ["9701-MAG-4K", "MAG first", 25000, 15000, None],
                ["9701-MAG-4K", "MAG duplicate", 25000, 15000, None],
                ["BAD-MSRP", "Invalid MSRP", "TBD", 100, None],
                ["BAD-DEALER", "Invalid dealer", 200, "ASK", None],
                ["HIGH-DEALER", "Dealer above MSRP", 100, 150, None],
                ["5608", "Cygnus Thickness Gauge", 5000, 5000, "NO DEALER DISCOUNT"],
            ]
        }
    )


def missing_columns_price_book() -> BytesIO:
    return build_workbook(
        {
            "MAG": [
                ["Item Name", "List Amount", "Comment"],
                ["MAG Utility Crawler", 25000, "No SKU column"],
            ]
        }
    )


def official_dt40_style_book() -> BytesIO:
    return build_workbook(
        {
            "CONFIG": [
                ["Dealer Retail:", "D_MSRP"],
                ["Discount Code:", "E_DT-40"],
            ],
            "PHOTON": [
                ["Product:", "PHOTON"],
                [None, None, None, " PHOTON MSRP"],
                [None, "Part Number", "Description", "MSRP", "DT40", "Notes:"],
                [None, None, "PRICING HAS BEEN UPDATED FOR APRIL 1ST, 2026", "MSRP USD", "Deep Trekker Dealer 40% Discount"],
                [None, None, None, 1, 0.6],
                [None, "9680-BASE", "PHOTON BASE PACKAGE", 17391, 10434.6, "Base package"],
                [None, "2535", "GAME PAD", 105, 105, "*** NO DEALER DISCOUNT - Requires Bridge Box"],
                [None, "9757-2", "BRIDGE BOX, ROV", 3280, 1968, None],
                [None, "00123", "LEADING ZERO SAMPLE", 10, 6, None],
            ],
            "PIVOT": [
                [None, "Part Number", "Description", "MSRP", "DT40", "Notes:"],
                [None, "2535", "GAME PAD", 105, 105, "*** NO DEALER DISCOUNT - Requires Bridge Box"],
                [None, "8998", "CONTROLLER CABLE SUPPORT", 381, 228.6, "PIVOT note"],
            ],
            "REVOLUTION-OBSOLETE": [
                [None, "Part Number", "Description", "MSRP", "DT40", "Notes:"],
                [None, "7511-SC-BASE", "Obsolete package", 1000, 600, "obsolete"],
                [None, "8998", "CONTROLLER CABLE SUPPORT", 381, 228.6, "OBSOLETE note"],
            ],
        }
    )


def official_pt30_style_book() -> BytesIO:
    return build_workbook(
        {
            "CONFIG": [
                ["Dealer Retail:", "D_MSRP"],
                ["Discount Code:", "E_PT-30"],
            ],
            "A-150": [
                [None, "Part Number", "Description", "MSRP", "PT30", "Notes:"],
                [None, "11000S", "A-150 S", 47500, 33250, "Base system"],
            ],
            "A-200": [
                [None, "Part Number", "Description", "MSRP", "PT30", "Notes:"],
                [None, "10800S", "A-200S", 52500, 36750, None],
            ],
        }
    )


def duplicate_same_price_book() -> BytesIO:
    return build_workbook(
        {
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["2535", "GAME PAD", 105, 105, "NO DEALER DISCOUNT"],
            ],
            "PIVOT": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["2535", "GAME PAD", 105, 105, "NO DEALER DISCOUNT"],
            ],
        }
    )


def price_conflict_book() -> BytesIO:
    return build_workbook(
        {
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["11490", "ALL TERRAIN WHEELS", 1945, 1361.5, None],
            ],
            "A-200": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["11490", "ALL TERRAIN WHEELS", 77000, 68701.1, None],
            ],
        }
    )


def datetime_sku_book() -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PHOTON"
    worksheet.append(["Part Number", "Description", "MSRP", "Dealer Price", "Notes"])
    worksheet.append([datetime(9757, 2, 1), "Broken date SKU", 3280, 1968, None])
    worksheet.append(["2500-1", "BRIDGE CONSOLE ONLY, NO DPK", 26250, 15750, None])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def previous_master_book() -> BytesIO:
    return build_workbook(
        {
            "MAG": [
                ["Part Number", "Description", "MSRP", "Dealer Price", "Notes"],
                ["9701-MAG-4K", "MAG Utility Crawler old name", 24000, 14400, "Old note"],
                ["9735", "MAG accessory", 800, 480, None],
                ["RETIRED-1", "Removed later", 10, 6, None],
            ],
            "PHOTON": [
                ["SKU", "Description", "MSRP", "Dealer Price", "Notes"],
                ["7851-PHOTON", "PHOTON ROV", 8000, 4800, None],
            ],
        }
    )


def supplier_quote_validation_books() -> tuple[BytesIO, BytesIO]:
    dt40 = build_workbook(
        {
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                ["DEALER-1", "Dealer priced item", 1000, 600, None],
                ["MSRP-1", "Quoted at list", 2000, 1200, None],
                ["MISMATCH-1", "Neither price", 3000, 1800, None],
                ["5608", "CYGNUS", 10448, 10448, "***NO DEALER DISCOUNT"],
                ["NDD-NUM-ONLY", "Dealer equals MSRP without note", 500, 500, None],
                ["SPECIAL-1", "Special candidate", 800, 480, None],
            ],
            "REVOLUTION-OBSOLETE": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                ["7511-SC-BASE", "Obsolete package", 1000, 600, "obsolete"],
            ],
        }
    )
    pt30 = build_workbook(
        {
            "A-200": [
                ["Part Number", "Description", "MSRP", "PT30", "Notes:"],
                ["CONFLICT-1", "Conflict A", 100, 60, "note A"],
            ],
            "A-150": [
                ["Part Number", "Description", "MSRP", "PT30", "Notes:"],
                ["CONFLICT-1", "Conflict B", 200, 140, "note B"],
            ],
        }
    )
    return dt40, pt30


def official_ihi_sku_snapshot_book() -> BytesIO:
    return build_workbook(
        {
            "VAC & MAG": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                [
                    "9701-MAG-4K",
                    "MAG CRAWLER PACKAGE 4K",
                    35437,
                    21262.2,
                    "75m Tether, Reel, 7\" LCD Controller, Carry Case, Sensors, MAG Wheels, Forward Facing 4K Body Camera, LED Lights",
                ],
                ["9735", "ELEVATING PAN TILT CAMERA KIT + LEDS - UTILITY CRAWLERS", 12075, 7245, None],
                ["5608", "CYGNUS THICKNESS GAUGE", 10448, 10448, "***NO DEALER DISCOUNT"],
            ],
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                [
                    "9680-BASE",
                    "PHOTON BASE PACKAGE",
                    17391,
                    10434.6,
                    "7\" Controller, 4K Camera, Lights, 1 Battery, Sensor Pod, Case, 150m Tether on reel, 1 Year Warranty",
                ],
                ["7851-PHOTON", "CYGNUS THICKNESS GAUGE INTEGRATION KIT (ONLY) - PHOTON", 1746, 1047.6, "USES RS485 COMS"],
                ["8459", "SPARE BATTERY - PHOTON", 787, 472.2, "ONE (1) BATTERY PACK"],
            ],
        }
    )


def manufacturer_books_for_reconciliation() -> tuple[BytesIO, BytesIO]:
    dt40 = build_workbook(
        {
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                ["9680-BASE", "PHOTON BASE PACKAGE", 17391, 10434.6, "Base"],
                ["8459", "POWER PACK", 787, 472.2, None],
                ["7851-PHOTON", "CYGNUS INTEGRATION", 1746, 1047.6, None],
                ["11490", "ALL TERRAIN WHEELS", 1945, 1361.5, None],
                ["5608", "CYGNUS", 10448, 10448, "NO DEALER DISCOUNT"],
            ],
            "REVOLUTION-OBSOLETE": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                ["7511-SC-BASE", "Obsolete package", 1000, 600, "obsolete"],
            ],
        }
    )
    pt30 = build_workbook(
        {
            "A-200": [
                ["Part Number", "Description", "MSRP", "PT30", "Notes:"],
                ["10800S", "A-200S", 52500, 36750, None],
                ["CONFLICT-1", "Conflict A", 100, 60, "note A"],
            ],
            "A-150": [
                ["Part Number", "Description", "MSRP", "PT30", "Notes:"],
                ["CONFLICT-1", "Conflict B", 200, 140, "note B"],
            ],
        }
    )
    return dt40, pt30


def spaceone_master_book() -> BytesIO:
    workbook = Workbook()
    photon = workbook.active
    photon.title = "PHOTON"
    photon.append(
        [None, "Part Number", "製品名", "内容", "価格", None, None, None, None, None, None, None, None, None, None, None, None, "スペースワン設定価格"]
    )
    photon.append([None, "DTマスター\n価格表より", None, None, "定価", "卸値"])
    rows = [
        (3, "PHOTON", "9680-BASE", "PHOTON 基本構成", "日本語説明", 17391, 10434.6, 3540000),
        (5, None, "9680-EXPEET", "EXPERT誤記", None, 26575, 15948, None),
        (7, None, 8459.0, "パワーパック", None, 787, 472.2, 160000),
        (9, None, datetime(9757, 2, 1), "BRIDGE BOX", None, 3280, 1968, None),
        (11, None, "7851-PHOTON", "Cygnus取付", None, None, None, 370000),
        (13, None, "11490", "全地形ホイール", None, 77000, 68701.1, None),
        (15, None, "5608", "Cygnus 1", None, 10448, 10448, 2220000),
        (17, None, "5608", "Cygnus 2 別用途", None, 10448, 10448, 2220000),
        (19, None, "MISSING-SKU", "存在しない", None, 10, 6, None),
        (21, None, "7511-SC-BASE", "廃番候補", None, 1000, 600, None),
        (23, None, "CONFLICT-1", "価格Conflict", None, 100, 60, None),
        (25, None, None, "キャビブラスター輸送費", "パレット", 9980, 9980, None),
    ]
    for row_number, category, sku, name, content, msrp, dealer, sales in rows:
        photon.cell(row_number, 1, category)
        photon.cell(row_number, 2, sku)
        photon.cell(row_number, 3, name)
        photon.cell(row_number, 4, content)
        photon.cell(row_number, 5, msrp)
        photon.cell(row_number, 6, dealer)
        if sales is not None:
            photon.cell(row_number + 1, 18, sales)
    photon["E3"] = (
        '=IFERROR(IMPORTRANGE("https://docs.google.com/spreadsheets/d/'
        '1xVJqlhF-sqMKJ3bnZ5hYjP2lTiYMl_5VaODFBL5CIx4/edit","PHOTON!D10"),17391)'
    )
    photon["E13"] = (
        '=IFERROR(IMPORTRANGE("https://docs.google.com/spreadsheets/d/'
        '1xVJqlhF-sqMKJ3bnZ5hYjP2lTiYMl_5VaODFBL5CIx4/edit","PHOTON!D40"),77000)'
    )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def write_fixture_files(directory: Path = FIXTURE_DIR) -> dict:
    directory.mkdir(parents=True, exist_ok=True)
    files = {
        "valid_dt40.xlsx": valid_price_book(),
        "shifted_columns.xlsx": shifted_columns_price_book(),
        "validation.xlsx": validation_price_book(),
        "missing_columns.xlsx": missing_columns_price_book(),
        "previous_master.xlsx": previous_master_book(),
    }
    paths = {}
    for name, buffer in files.items():
        path = directory / name
        path.write_bytes(buffer.getvalue())
        paths[name] = path
    return paths

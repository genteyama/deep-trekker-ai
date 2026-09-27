from datetime import date, datetime

from openpyxl import load_workbook
import csv
import pytest

from agents.quote_builder import apply_final_price, apply_presentation_mode, apply_shipping_final_price
from agents.quote_dates import format_quote_date, parse_quote_date
from agents.quote_export import (
    CUSTOMER_FORBIDDEN_TERMS,
    EXCEL_DATE_NUMBER_FORMAT,
    QuoteExportError,
    approved_remark_texts,
    build_export_bundle,
    export_ihi_photon_files,
    export_internal_calc_excel,
    export_moneyforward_csv,
    export_moneyforward_tsv,
    export_spaceone_quote_excel,
    scan_customer_workbook_leaks,
    unique_output_path,
)
from models import CustomerPresentationMode, ExportFileType, ExportPurpose, FinalPriceStatus
from tests.test_quote_approval import _approve, _confirm_remarks, _mag_draft, _ready_photon
from tests.test_quote_builder import _mag_draft as _build_mag_draft


def _photon_snapshot():
    return _approve(_ready_photon())[1]


def _bundled_mag_snapshot():
    draft = _mag_draft(tax_rate=0.1)
    parent = next(line for line in draft.configuration_lines if line.manufacturer_sku == "2604")
    child = next(line for line in draft.configuration_lines if line.manufacturer_sku == "2601")
    apply_presentation_mode(
        draft,
        child.line_id,
        CustomerPresentationMode.BUNDLED_WITH_PARENT,
        bundled_into_line_id=parent.line_id,
    )
    for line in draft.configuration_lines:
        if line.customer_presentation_status == CustomerPresentationMode.SEPARATE_LINE:
            apply_final_price(draft, line.line_id, FinalPriceStatus.USE_STANDARD_CANDIDATE)
    apply_shipping_final_price(draft, 1000000)
    _confirm_remarks(draft)
    return _approve(draft)[1]


def test_formal_export_rejects_non_snapshot_source(tmp_path):
    draft = _mag_draft(tax_rate=0.1)

    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        export_moneyforward_tsv(draft, output_dir=tmp_path)
    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        export_spaceone_quote_excel(draft, output_dir=tmp_path, official_quote_number="8195")


def test_moneyforward_tsv_and_csv_match_approved_totals(tmp_path):
    snapshot = _photon_snapshot()
    bundle, tsv_path = export_moneyforward_tsv(snapshot, output_dir=tmp_path, generated_by="弦")
    bundle, csv_path = export_moneyforward_csv(snapshot, output_dir=tmp_path, bundle=bundle, generated_by="弦")
    tsv_text = tsv_path.read_text(encoding="utf-8")
    csv_bytes = csv_path.read_bytes()
    csv_text = csv_bytes.decode("utf-8-sig")
    tsv_amounts = [int(line.split("\t")[4]) for line in tsv_text.splitlines()[1:] if line]
    csv_amounts = [int(row[4]) for row in csv.reader(csv_text.splitlines()) if row and row[0] != "品目"]

    assert tsv_path.name == "IHI_QUOTE_001_v1_moneyforward.tsv"
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    assert tsv_text.splitlines()[0] == "品目\t品目詳細\t単価\t数量\t金額\t備考"
    assert csv_text.splitlines()[0] == "品目,品目詳細,単価,数量,金額,備考"
    assert sum(tsv_amounts) == 7160000
    assert sum(csv_amounts) == 7160000
    assert bundle.moneyforward_payload.tax_jpy == 716000
    assert bundle.moneyforward_payload.total_jpy == 7876000
    assert all(item.total == 7876000 for item in bundle.file_manifest)
    assert all(item.subtotal_ex_tax == 7160000 for item in bundle.file_manifest)
    assert "保険" not in tsv_text.splitlines()[0]
    assert ExportFileType.MONEYFORWARD_TSV in [item.file_type for item in bundle.file_manifest]
    assert ExportFileType.MONEYFORWARD_CSV in [item.file_type for item in bundle.file_manifest]


def test_internal_excel_keeps_bundled_parts_and_metadata(tmp_path):
    snapshot = _bundled_mag_snapshot()
    bundle, path = export_internal_calc_excel(snapshot, output_dir=tmp_path, generated_by="弦")
    workbook = load_workbook(path)
    metadata = {row[0]: row[1] for row in workbook["Metadata"].iter_rows(values_only=True)}
    parts = [row[0] for row in workbook["Calculation"].iter_rows(min_row=2, values_only=True) if row[0]]

    assert path.name.endswith("_internal_calc.xlsx")
    assert "2601" in parts
    assert metadata["Approved Snapshot ID"] == snapshot.approved_quote_snapshot_id
    assert metadata["Quote Version"] == snapshot.quote_version
    assert metadata["Case ID"] == "IHI_QUOTE_001"
    assert metadata["Customer Sales Total"] == snapshot.total_jpy
    assert metadata["Gross Margin"] == snapshot.gross_margin_rate
    assert metadata["Exchange Rate"] == snapshot.exchange_rate
    assert metadata["Approved By"] == "弦"
    assert bundle.internal_transfer_payload.customer_total_jpy == snapshot.total_jpy


def test_spaceone_quote_excel_hides_internal_cost_and_uses_snapshot_text(tmp_path):
    snapshot = _photon_snapshot()
    with pytest.raises(QuoteExportError, match="official_quote_number"):
        export_spaceone_quote_excel(snapshot, output_dir=tmp_path, purpose=ExportPurpose.FORMAL)
    bundle, path = export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        generated_by="弦",
        purpose=ExportPurpose.FORMAL,
    )
    workbook = load_workbook(path)
    sheet = workbook["見積書"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    text = " ".join(str(item) for item in values if item is not None)

    assert path.name == "8195_spaceone_quote.xlsx"
    assert "見積書" in values
    assert "株式会社IHI検査計測 御中" in values
    assert snapshot.title in values
    assert "8195" in values
    assert "国際輸送費" in text
    assert "カナダ→日本" in text
    assert "大型梱包×1" in text
    assert 7160000 in values
    assert 716000 in values
    assert 7876000 in values
    assert 870000 in values
    for remark in bundle.spaceone_quote_payload.remarks:
        assert remark in values
    assert scan_customer_workbook_leaks(path) == set()
    for term in CUSTOMER_FORBIDDEN_TERMS:
        assert term not in text
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_area.endswith("$A$1:$D$26") or "A1:D" in sheet.print_area.replace("$", "")
    assert bundle.file_manifest[-1].pdf_layout_prepared is True
    assert bundle.file_manifest[-1].purpose == ExportPurpose.FORMAL


def test_export_does_not_overwrite_and_ignores_later_master_changes(tmp_path):
    snapshot = _photon_snapshot()
    _, first = export_moneyforward_tsv(snapshot, output_dir=tmp_path)
    _, second = export_moneyforward_tsv(snapshot, output_dir=tmp_path)
    original_total = snapshot.total_jpy
    snapshot.configuration_snapshot[0].final_sales_price_jpy = 1
    _, excel = export_internal_calc_excel(snapshot, output_dir=tmp_path)
    workbook = load_workbook(excel)
    metadata = {row[0]: row[1] for row in workbook["Metadata"].iter_rows(values_only=True)}

    assert first != second
    assert first.exists() and second.exists()
    assert metadata["Customer Sales Total"] == original_total
    assert snapshot.total_jpy == original_total


def test_ihi_photon_three_outputs_share_approved_totals(tmp_path):
    snapshot = _photon_snapshot()
    bundle, paths = export_ihi_photon_files(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        generated_by="弦",
    )
    tsv_amounts = [
        int(line.split("\t")[4])
        for line in paths["moneyforward_tsv"].read_text(encoding="utf-8").splitlines()[1:]
        if line
    ]
    spaceone = load_workbook(paths["spaceone_quote"])["見積書"]
    spaceone_values = [cell.value for row in spaceone.iter_rows() for cell in row]
    internal_meta = {
        row[0]: row[1] for row in load_workbook(paths["internal_calc"])["Metadata"].iter_rows(values_only=True)
    }

    assert sum(tsv_amounts) == 7160000
    assert 7160000 in spaceone_values
    assert 716000 in spaceone_values
    assert 7876000 in spaceone_values
    assert internal_meta["Customer Sales Total"] == 7876000
    assert internal_meta["Customer Subtotal"] == 7160000
    assert internal_meta["Customer Tax"] == 716000
    assert {item.total for item in bundle.file_manifest} == {7876000}
    assert {item.subtotal_ex_tax for item in bundle.file_manifest} == {7160000}
    assert bundle.approved_quote_snapshot_id == snapshot.approved_quote_snapshot_id
    assert all(item.sha256 for item in bundle.file_manifest)


def test_mag_review_draft_cannot_export_formally(tmp_path):
    draft = _build_mag_draft(tax_rate=0.1)

    assert draft.status.value == "REVIEW_REQUIRED"
    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        build_export_bundle(draft)
    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        export_internal_calc_excel(draft, output_dir=tmp_path)
    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        export_spaceone_quote_excel(draft, output_dir=tmp_path, official_quote_number="8194")


def test_internal_excel_uses_snapshot_standard_and_shipping_summary(tmp_path):
    snapshot = _photon_snapshot()
    original_standards = {
        line.manufacturer_sku: line.standard_sales_price_candidate_jpy
        for line in snapshot.configuration_snapshot
    }
    bundle, path = export_internal_calc_excel(snapshot, output_dir=tmp_path, generated_by="弦")
    workbook = load_workbook(path)
    calc_rows = {
        row[0]: row for row in workbook["Calculation"].iter_rows(min_row=2, values_only=True) if row[0]
    }
    shipping_rows = list(workbook["Shipping"].iter_rows(min_row=2, values_only=True))
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True)}
    component_finals = [
        row[7]
        for row in shipping_rows
        if row[0] in {"LARGE_BOX", "SMALL_BOX"}
    ]
    aggregate = next(row for row in shipping_rows if row[0] == "INTERNATIONAL_SHIPPING")

    assert original_standards["9680-BASE"] == 3547764
    assert calc_rows["9680-BASE"][10] == 3547764
    assert calc_rows["8459"][10] == 160548
    assert calc_rows["5608"][10] == 2309008
    assert calc_rows["7851-PHOTON"][10] == 371025
    assert calc_rows["9680-BASE"][11] == 3540000
    assert {row[0] for row in shipping_rows} >= {"LARGE_BOX", "SMALL_BOX", "INTERNATIONAL_SHIPPING"}
    assert all(value is None for value in component_finals)
    assert aggregate[7] == 870000
    assert aggregate[5] == bundle.internal_transfer_payload.shipping_cost_jpy
    assert aggregate[8] == round(870000 - bundle.internal_transfer_payload.shipping_cost_jpy, 4)
    assert summary["Product Sales"] == 6290000
    assert summary["Shipping Sales"] == 870000
    assert summary["Total Sales ex Tax"] == 7160000
    assert summary["Product Landed Cost"] == bundle.internal_transfer_payload.product_landed_cost_jpy
    assert summary["Shipping Cost"] == bundle.internal_transfer_payload.shipping_cost_jpy
    assert summary["Total Landed Cost"] == snapshot.total_landed_cost_jpy
    assert summary["Gross Profit"] == snapshot.gross_profit_jpy
    assert summary["Gross Margin"] == snapshot.gross_margin_rate
    assert workbook["Calculation"].column_dimensions["B"].width >= 40
    assert snapshot.total_jpy == 7876000


def test_moneyforward_uses_customer_description_and_does_not_copy_item_name(tmp_path):
    snapshot = _photon_snapshot()
    snapshot.customer_lines_snapshot[1].description = None
    _, path = export_moneyforward_tsv(snapshot, output_dir=tmp_path)
    rows = [line.split("\t") for line in path.read_text(encoding="utf-8").splitlines()[1:] if line]
    by_name = {row[0]: row[1] for row in rows}

    assert "7インチLCDコントローラー" in by_name["DeepTrekker PHOTON BASE Package"]
    assert "PHOTON用予備バッテリー" not in by_name["POWER PACK ASY, PHOTON"]
    assert by_name["POWER PACK ASY, PHOTON"] == ""
    assert by_name["THICKNESS GUAGE - CYGNUS"] == "超音波肉厚計測機－本体セット"
    assert by_name["THICKNESS GAUGE - CYGNUS INTEGRATION KIT"] == "肉厚計測機組込みインテグレーションキット（PHOTON用）"
    assert "カナダ→日本" in by_name["国際輸送費"]
    assert by_name["POWER PACK ASY, PHOTON"] != "POWER PACK ASY, PHOTON"


def test_formal_spaceone_requires_dates_and_uses_issuer_snapshot(tmp_path):
    snapshot = _photon_snapshot()
    missing_until = snapshot.model_copy(update={"valid_until": None})
    missing_issue = snapshot.model_copy(update={"issue_date": None})
    missing_issuer = snapshot.model_copy(update={"issuer_snapshot": None})
    current_company = {"company_name": "Current Company Master"}

    with pytest.raises(QuoteExportError, match="valid_until"):
        export_spaceone_quote_excel(
            missing_until, output_dir=tmp_path, official_quote_number="8195", purpose=ExportPurpose.FORMAL
        )
    with pytest.raises(QuoteExportError, match="issue_date"):
        export_spaceone_quote_excel(
            missing_issue, output_dir=tmp_path, official_quote_number="8195", purpose=ExportPurpose.FORMAL
        )
    with pytest.raises(QuoteExportError, match="issuer_snapshot"):
        export_spaceone_quote_excel(
            missing_issuer, output_dir=tmp_path, official_quote_number="8195", purpose=ExportPurpose.FORMAL
        )

    bundle, path = export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        purpose=ExportPurpose.FORMAL,
    )
    sheet = load_workbook(path)["見積書"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    text = " ".join(str(item) for item in values if item is not None)
    base_row = next(row for row in sheet.iter_rows() if row[0].value and "PHOTON BASE Package" in str(row[0].value))

    assert snapshot.issue_date == "2026-09-26"
    assert snapshot.valid_until == "2026-10-31"
    assert isinstance(sheet["B9"].value, (date, datetime))
    assert isinstance(sheet["B10"].value, (date, datetime))
    assert parse_quote_date(sheet["B9"].value) == date(2026, 9, 26)
    assert parse_quote_date(sheet["B10"].value) == date(2026, 10, 31)
    assert sheet["B9"].number_format.lower() == EXCEL_DATE_NUMBER_FORMAT
    assert sheet["B10"].number_format.lower() == EXCEL_DATE_NUMBER_FORMAT
    assert format_quote_date(sheet["B9"].value) == "2026/09/26"
    assert format_quote_date(sheet["B10"].value) == "2026/10/31"
    assert "2026-09-26" not in values
    assert "2026-10-31" not in values
    assert snapshot.issuer_snapshot.company_name in values
    assert snapshot.issuer_snapshot.address in values
    assert current_company["company_name"] not in values
    assert "7インチLCDコントローラー" in text
    assert "＜納期目安＞発注より3ヵ月程度（2026年9月26日現在）" in values
    assert "※納期はご発注確定後改めてお知らせいたします。" in values
    assert base_row[0].alignment.wrap_text is True
    assert sheet.row_dimensions[base_row[0].row].height >= 28
    assert scan_customer_workbook_leaks(path) == set()
    assert bundle.spaceone_quote_payload.total == 7876000
    assert snapshot.total_jpy == 7876000


def test_spaceone_excel_writes_snapshot_remarks_and_dynamic_print_area(tmp_path):
    snapshot = _photon_snapshot()
    remarks = approved_remark_texts(snapshot)
    _, path = export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        purpose=ExportPurpose.FORMAL,
    )
    sheet = load_workbook(path)["見積書"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    remark_rows = [row[0].row for row in sheet.iter_rows() if row[0].value in remarks]
    heading = next(row[0].row for row in sheet.iter_rows() if row[0].value == "備考")
    print_area = sheet.print_area.replace("$", "")

    assert remarks == [
        "＜納期目安＞発注より3ヵ月程度（2026年9月26日現在）",
        "※１：導入講習は実施場所により別途交通費諸経費がかかる場合があります。",
        "※本見積書の価格は、為替レートの変動により変更される可能性があります。",
        "※センサー等の組み込み工賃、輸送時保険料は代金に含まれます。",
        "※納期はご発注確定後改めてお知らせいたします。",
    ]
    assert all(item in values for item in remarks)
    assert all(row > heading for row in remark_rows)
    assert max(remark_rows) > heading
    assert print_area.endswith(f"D{max(remark_rows)}")
    assert print_area != "A1:D25"
    assert sheet.page_setup.fitToWidth == 1
    assert snapshot.total_jpy == 7876000
    assert 7876000 in values
    assert scan_customer_workbook_leaks(path) == set()


def test_empty_snapshot_remarks_do_not_add_historical_text(tmp_path):
    snapshot = _photon_snapshot().model_copy(update={"remarks": []})
    _, path = export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        purpose=ExportPurpose.FORMAL,
    )
    values = [cell.value for row in load_workbook(path)["見積書"].iter_rows() for cell in row]
    historical = "＜納期目安＞発注より3ヵ月程度（2026年9月26日現在）"

    assert "備考" in values
    assert historical not in values
    assert "輸送時保険料は代金に含む。" not in values
    assert 7876000 in values


def test_spaceone_excel_dates_are_excel_date_cells_not_iso_strings(tmp_path):
    snapshot = _photon_snapshot()
    original_issue = snapshot.issue_date
    original_valid = snapshot.valid_until
    original_total = snapshot.total_jpy
    _, path = export_spaceone_quote_excel(
        snapshot,
        output_dir=tmp_path,
        official_quote_number="8195",
        purpose=ExportPurpose.FORMAL,
    )
    sheet = load_workbook(path)["見積書"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]

    assert original_issue == "2026-09-26"
    assert original_valid == "2026-10-31"
    assert snapshot.issue_date == original_issue
    assert snapshot.valid_until == original_valid
    assert isinstance(sheet["B9"].value, (date, datetime))
    assert isinstance(sheet["B10"].value, (date, datetime))
    assert parse_quote_date(sheet["B9"].value) == date(2026, 9, 26)
    assert parse_quote_date(sheet["B10"].value) == date(2026, 10, 31)
    assert sheet["B9"].number_format == "yyyy/mm/dd"
    assert sheet["B10"].number_format == "yyyy/mm/dd"
    assert format_quote_date(sheet["B9"].value) == "2026/09/26"
    assert format_quote_date(sheet["B10"].value) == "2026/10/31"
    assert "2026-09-26" not in values
    assert "2026-10-31" not in values
    assert 7876000 in values
    assert snapshot.total_jpy == 7876000
    assert original_total == 7876000


def test_unique_output_path_does_not_overwrite(tmp_path):
    first = unique_output_path(tmp_path, "8195_moneyforward.csv")
    first.write_text("one", encoding="utf-8")
    second = unique_output_path(tmp_path, "8195_moneyforward.csv")
    second.write_text("two", encoding="utf-8")

    assert first != second
    assert first.read_text(encoding="utf-8") == "one"
    assert second.read_text(encoding="utf-8") == "two"

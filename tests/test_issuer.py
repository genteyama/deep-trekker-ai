import json

from openpyxl import load_workbook
import pytest

import agents.issuer as issuer_module
from agents.formal_quote_document import build_formal_quote_document
from agents.issuer import ISSUER_CONFIG_PATH, IssuerConfigError, load_issuer_snapshot
from agents.quote_approval import QuoteApprovalError, QuoteApprovalStore, create_revision_draft
from agents.quote_builder import (
    apply_final_price,
    apply_issuer_snapshot,
    apply_presentation_mode,
    apply_shipping_final_price,
)
from agents.quote_export import QuoteExportError, export_spaceone_quote_excel, export_spaceone_quote_pdf
from agents.quote_pdf import extract_pdf_text
from models import CustomerPresentationMode, IssuerSnapshot, ExportPurpose, FinalPriceStatus, QuoteDraftStatus
from tests.test_quote_approval import _approve, _confirm_remarks, _mag_draft, _ready_photon

OFFICIAL_ISSUER = {
    "company_name": "株式会社スペースワン",
    "address": "〒963-8833 福島県郡山市香久池1-17-3",
    "office_address": "東京営業所：〒110-0005 東京都台東区上野1-20-1-5F",
    "telephone": "TEL: 024-954-9930",
    "source_reference": "config/issuer.json",
}
HISTORICAL_SOURCE = "IHI_QUOTE_001 Golden Historical Issuer Snapshot"


def _write_config(tmp_path, payload, name="issuer.json"):
    path = tmp_path / name
    path.write_text(payload if isinstance(payload, str) else json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def _use_config(monkeypatch, path):
    monkeypatch.setattr(issuer_module, "ISSUER_CONFIG_PATH", path)


def _ready_mag():
    draft = _mag_draft(tax_rate=0.1)
    parent = next(line for line in draft.configuration_lines if line.manufacturer_sku == "2604")
    child = next(line for line in draft.configuration_lines if line.manufacturer_sku == "2601")
    apply_presentation_mode(
        draft,
        child.line_id,
        CustomerPresentationMode.BUNDLED_WITH_PARENT,
        bundled_into_line_id=parent.line_id,
    )
    for line in draft.configuration_lines:
        if line.customer_presentation_status == CustomerPresentationMode.SEPARATE_LINE:
            apply_final_price(draft, line.line_id, FinalPriceStatus.USE_STANDARD_CANDIDATE)
    apply_shipping_final_price(draft, 1000000)
    _confirm_remarks(draft)
    return draft


def _ready_photon_without_issuer():
    draft = _ready_photon()
    draft.issuer_snapshot = None
    return draft


def _issuer_dict(snapshot):
    return snapshot.issuer_snapshot.model_dump()


def test_official_issuer_config_loads_all_fields():
    issuer = load_issuer_snapshot()

    assert ISSUER_CONFIG_PATH.name == "issuer.json"
    assert issuer.company_name == OFFICIAL_ISSUER["company_name"]
    assert issuer.address == OFFICIAL_ISSUER["address"]
    assert issuer.office_address == OFFICIAL_ISSUER["office_address"]
    assert issuer.telephone == OFFICIAL_ISSUER["telephone"]
    assert issuer.source_reference == OFFICIAL_ISSUER["source_reference"]


def test_missing_issuer_config_fails_closed(tmp_path):
    with pytest.raises(IssuerConfigError, match="not found"):
        load_issuer_snapshot(tmp_path / "missing.json")


def test_invalid_issuer_json_fails_closed(tmp_path):
    with pytest.raises(IssuerConfigError, match="unreadable"):
        load_issuer_snapshot(_write_config(tmp_path, "{not json"))
    with pytest.raises(IssuerConfigError, match="JSON object"):
        load_issuer_snapshot(_write_config(tmp_path, "[]", name="list.json"))


@pytest.mark.parametrize("field", list(OFFICIAL_ISSUER))
@pytest.mark.parametrize("value", ["", "   ", None])
def test_blank_required_issuer_field_fails_closed(tmp_path, field, value):
    payload = dict(OFFICIAL_ISSUER, **{field: value})
    with pytest.raises(IssuerConfigError, match=field):
        load_issuer_snapshot(_write_config(tmp_path, payload))


def test_absent_required_issuer_field_fails_closed(tmp_path):
    payload = {key: value for key, value in OFFICIAL_ISSUER.items() if key != "telephone"}
    with pytest.raises(IssuerConfigError, match="telephone"):
        load_issuer_snapshot(_write_config(tmp_path, payload))


def test_mag_approval_without_draft_issuer_freezes_official_issuer():
    draft = _ready_mag()
    assert draft.issuer_snapshot is None

    _, snapshot = _approve(draft)

    assert _issuer_dict(snapshot) == OFFICIAL_ISSUER


def test_photon_approval_without_draft_issuer_freezes_official_issuer():
    _, snapshot = _approve(_ready_photon_without_issuer())

    assert _issuer_dict(snapshot) == OFFICIAL_ISSUER


def test_historical_draft_issuer_is_kept():
    _, snapshot = _approve(_ready_photon())

    assert snapshot.issuer_snapshot.source_reference == HISTORICAL_SOURCE
    assert snapshot.issuer_snapshot.address == "福島県郡山市香久池1-17-3"


@pytest.mark.parametrize("broken", ["missing", "invalid", "blank"])
def test_broken_issuer_config_blocks_approval_without_mutation(tmp_path, monkeypatch, broken):
    if broken == "missing":
        path = tmp_path / "missing.json"
    elif broken == "invalid":
        path = _write_config(tmp_path, "{")
    else:
        path = _write_config(tmp_path, dict(OFFICIAL_ISSUER, company_name=""))
    _use_config(monkeypatch, path)
    draft = _ready_mag()
    status_before = draft.status

    with pytest.raises(QuoteApprovalError, match="Issuer snapshot unavailable"):
        _approve(draft)
    assert draft.status == status_before
    assert draft.status != QuoteDraftStatus.APPROVED


def test_mag_formal_pdf_and_excel_succeed_with_official_issuer(tmp_path):
    _, snapshot = _approve(_ready_mag())

    _, pdf_path = export_spaceone_quote_pdf(snapshot, output_dir=tmp_path, official_quote_number="9001")
    _, xlsx_path = export_spaceone_quote_excel(
        snapshot, output_dir=tmp_path, official_quote_number="9001", purpose=ExportPurpose.FORMAL
    )
    pdf_text = "".join(extract_pdf_text(pdf_path).split())
    sheet = load_workbook(xlsx_path)["見積書"]
    values = [cell.value for row in sheet.iter_rows() for cell in row]

    for key in ("company_name", "address", "office_address", "telephone"):
        assert OFFICIAL_ISSUER[key].replace(" ", "") in pdf_text
        assert OFFICIAL_ISSUER[key] in values


def test_mag_formal_document_keeps_rounded_arithmetic():
    _, snapshot = _approve(_ready_mag())
    document = build_formal_quote_document(snapshot, official_quote_number="9001")

    assert document.issuer.company_name == OFFICIAL_ISSUER["company_name"]
    assert sum(line.amount for line in document.customer_lines) == document.subtotal
    assert document.subtotal + document.tax_amount == document.total
    assert all(line.unit_price % 1000 == 0 for line in document.customer_lines)


def test_formal_export_still_rejects_draft_and_snapshot_without_issuer(tmp_path):
    draft = _ready_mag()
    with pytest.raises(QuoteExportError, match="ApprovedQuoteSnapshot"):
        export_spaceone_quote_pdf(draft, output_dir=tmp_path, official_quote_number="9001")

    _, snapshot = _approve(draft)
    missing = snapshot.model_copy(update={"issuer_snapshot": None})
    with pytest.raises(QuoteExportError, match="issuer_snapshot"):
        export_spaceone_quote_pdf(missing, output_dir=tmp_path, official_quote_number="9001")
    with pytest.raises(QuoteExportError, match="issuer_snapshot"):
        export_spaceone_quote_excel(
            missing, output_dir=tmp_path, official_quote_number="9001", purpose=ExportPurpose.FORMAL
        )


def test_issuer_config_change_after_approval_does_not_touch_existing_quote(tmp_path, monkeypatch):
    _use_config(monkeypatch, _write_config(tmp_path, OFFICIAL_ISSUER, name="v1.json"))
    _, snapshot_a = _approve(_ready_mag())
    frozen = snapshot_a.model_dump(mode="json")
    _, pdf_before = export_spaceone_quote_pdf(snapshot_a, output_dir=tmp_path / "before", official_quote_number="9001")
    _, xlsx_before = export_spaceone_quote_excel(
        snapshot_a, output_dir=tmp_path / "before", official_quote_number="9001", purpose=ExportPurpose.FORMAL
    )

    changed = dict(
        OFFICIAL_ISSUER,
        company_name="株式会社スペースワン新社名",
        telephone="TEL: 000-000-0000",
        source_reference="config/issuer.json@v2",
    )
    _use_config(monkeypatch, _write_config(tmp_path, changed, name="v2.json"))

    _, pdf_after = export_spaceone_quote_pdf(snapshot_a, output_dir=tmp_path / "after", official_quote_number="9001")
    _, xlsx_after = export_spaceone_quote_excel(
        snapshot_a, output_dir=tmp_path / "after", official_quote_number="9001", purpose=ExportPurpose.FORMAL
    )
    xlsx_values = lambda path: [cell.value for row in load_workbook(path)["見積書"].iter_rows() for cell in row]

    assert snapshot_a.model_dump(mode="json") == frozen
    assert _issuer_dict(snapshot_a) == OFFICIAL_ISSUER
    assert extract_pdf_text(pdf_after) == extract_pdf_text(pdf_before)
    assert xlsx_values(xlsx_after) == xlsx_values(xlsx_before)
    assert "新社名" not in extract_pdf_text(pdf_after)

    _, snapshot_b = _approve(_ready_photon_without_issuer())
    assert _issuer_dict(snapshot_b) == changed


ISSUER_B = dict(
    OFFICIAL_ISSUER,
    company_name="株式会社スペースワン新社名",
    telephone="TEL: 000-000-0000",
    source_reference="config/issuer.json@B",
)


def _formal_exports(snapshot, output_dir):
    _, pdf_path = export_spaceone_quote_pdf(snapshot, output_dir=output_dir, official_quote_number="9001")
    _, xlsx_path = export_spaceone_quote_excel(
        snapshot, output_dir=output_dir, official_quote_number="9001", purpose=ExportPurpose.FORMAL
    )
    pdf_text = "".join(extract_pdf_text(pdf_path).split())
    xlsx_values = [cell.value for row in load_workbook(xlsx_path)["見積書"].iter_rows() for cell in row]
    return pdf_text, xlsx_values


def test_revision_draft_does_not_inherit_approved_issuer():
    _, snapshot_v1 = _approve(_ready_mag())
    revision = create_revision_draft(snapshot_v1)

    assert _issuer_dict(snapshot_v1) == OFFICIAL_ISSUER
    assert revision.issuer_snapshot is None


def test_golden_revision_draft_does_not_inherit_historical_issuer():
    _, snapshot_v1 = _approve(_ready_photon())
    revision = create_revision_draft(snapshot_v1)

    assert snapshot_v1.issuer_snapshot.source_reference == HISTORICAL_SOURCE
    assert revision.issuer_snapshot is None


def test_revision_keeps_non_issuer_inheritance():
    _, snapshot_v1 = _approve(_ready_mag())
    revision = create_revision_draft(snapshot_v1)

    assert revision.customer == snapshot_v1.customer
    assert revision.configuration_name == snapshot_v1.configuration_name
    assert revision.exchange_rate == snapshot_v1.exchange_rate
    assert revision.pricing_context.exchange_rate_source == snapshot_v1.exchange_rate_source
    assert revision.pricing_context.market_reference_rate == snapshot_v1.market_reference_rate
    assert revision.pricing_context.market_reference_date == snapshot_v1.market_reference_date
    assert revision.subtotal_ex_tax_jpy == snapshot_v1.subtotal_ex_tax_jpy
    assert revision.total_jpy == snapshot_v1.total_jpy


def test_revision_approval_with_unchanged_config_freezes_same_issuer(tmp_path, monkeypatch):
    _use_config(monkeypatch, _write_config(tmp_path, OFFICIAL_ISSUER))
    _, snapshot_v1 = _approve(_ready_mag())
    _, snapshot_v2 = _approve(create_revision_draft(snapshot_v1))

    assert _issuer_dict(snapshot_v2) == OFFICIAL_ISSUER


def test_revision_approval_uses_current_issuer_config_and_keeps_v1(tmp_path, monkeypatch):
    store = QuoteApprovalStore()
    _use_config(monkeypatch, _write_config(tmp_path, OFFICIAL_ISSUER, name="a.json"))
    _, snapshot_v1 = _approve(_ready_mag(), store=store)
    frozen_v1 = snapshot_v1.model_dump(mode="json")
    revision = create_revision_draft(snapshot_v1, store=store)

    _use_config(monkeypatch, _write_config(tmp_path, ISSUER_B, name="b.json"))
    _, snapshot_v2 = _approve(revision, store=store)

    assert _issuer_dict(snapshot_v2) == ISSUER_B
    assert _issuer_dict(snapshot_v1) == OFFICIAL_ISSUER
    assert _issuer_dict(store.snapshots[snapshot_v1.approved_quote_snapshot_id]) == OFFICIAL_ISSUER
    assert snapshot_v1.model_dump(mode="json") == frozen_v1

    v1_pdf, v1_xlsx = _formal_exports(snapshot_v1, tmp_path / "v1")
    v2_pdf, v2_xlsx = _formal_exports(snapshot_v2, tmp_path / "v2")
    assert OFFICIAL_ISSUER["company_name"] in v1_xlsx
    assert ISSUER_B["company_name"] not in v1_xlsx
    assert ISSUER_B["company_name"] not in v1_pdf
    assert ISSUER_B["telephone"] not in v1_xlsx
    assert ISSUER_B["company_name"] in v2_xlsx
    assert ISSUER_B["company_name"] in v2_pdf
    assert ISSUER_B["telephone"] in v2_xlsx
    assert OFFICIAL_ISSUER["telephone"] not in v2_xlsx


def test_explicit_issuer_on_revision_draft_is_still_respected(tmp_path, monkeypatch):
    _, snapshot_v1 = _approve(_ready_mag())
    revision = create_revision_draft(snapshot_v1)
    explicit = IssuerSnapshot.model_validate(dict(OFFICIAL_ISSUER, source_reference="explicit"))
    apply_issuer_snapshot(revision, explicit)

    _use_config(monkeypatch, _write_config(tmp_path, ISSUER_B))
    _, snapshot_v2 = _approve(revision)

    assert snapshot_v2.issuer_snapshot.source_reference == "explicit"

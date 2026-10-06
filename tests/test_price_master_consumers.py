from pathlib import Path

from streamlit.testing.v1 import AppTest

from agents.landed_cost import calculate_landed_cost_scenario
from agents.master_reconciliation import collect_manufacturer_candidates
from agents.price_master import (
    SALES_MASTER_MISSING_WARNING,
    active_price_books,
    attach_price_master_provenance,
    get_active_master,
    import_price_master,
    load_active_landed_policy,
)
from agents.quote_builder import build_quote_draft
from agents.supplier_quote_validation import load_official_manufacturer_price_books
from models import InsuranceMode, PriceMasterType, PriceSourceType, QuoteDraftStatus
from parsers.quote_calc_parser import locate_quote_calc_workbook
from repositories.sqlite_quote_repository import SqliteQuoteRepository
from tests.price_book_fixtures import (
    build_workbook,
    official_ihi_sku_snapshot_book,
    official_pt30_style_book,
    pricing_policy_manufacturer_book,
    pricing_policy_master_book,
    quote_calc_formula_book,
)
from tests.test_price_source_safety import _policy
from tests.test_quote_approval import _approve, _ready_photon

APP_PATH = Path(__file__).resolve().parents[1] / "app.py"
ROOT = Path(__file__).resolve().parents[1]
PHOTON_SKUS = ("9680-BASE", "8459", "5608", "7851-PHOTON")


def _photon_dt40_book(base_dealer: float):
    return build_workbook(
        {
            "PHOTON": [
                ["Part Number", "Description", "MSRP", "DT40", "Notes:"],
                ["9680-BASE", "PHOTON BASE PACKAGE", 17391, base_dealer, None],
                ["7851-PHOTON", "CYGNUS INTEGRATION KIT - PHOTON", 1746, 1047.6, None],
                ["8459", "SPARE BATTERY - PHOTON", 787, 472.2, None],
                ["5608", "CYGNUS THICKNESS GAUGE", 10448, 10448, "***NO DEALER DISCOUNT"],
            ]
        }
    )


def _quote_calc_with_memo_sheet():
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(quote_calc_formula_book())
    workbook.create_sheet("memo")["A1"] = "v2"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _import(master_type, factory, name):
    outcome = import_price_master(master_type, name, factory().getvalue())
    assert outcome.record is not None and outcome.record.active
    return outcome.record


def _create_ihi_photon_draft():
    at = AppTest.from_file(str(APP_PATH)).run()
    at.button(key="open_quote_control").click().run()
    at.button(key="ihi_photon_draft").click().run()
    assert not at.exception
    return at, at.session_state["quote_draft"]


def _line(draft, sku="9680-BASE"):
    return next(line for line in draft.configuration_lines if line.manufacturer_sku == sku)


def test_without_masters_the_quote_stays_development_reference_and_says_so():
    _, draft = _create_ihi_photon_draft()

    assert all(
        line.manufacturer_price_snapshot.price_source_type == PriceSourceType.DEVELOPMENT_REFERENCE
        for line in draft.configuration_lines
    )
    assert draft.status == QuoteDraftStatus.REVIEW_REQUIRED
    assert all(any(SALES_MASTER_MISSING_WARNING in item for item in line.warnings) for line in draft.configuration_lines)
    assert not any("Price master" in item for item in draft.source_references)


def test_active_dt40_is_used_by_a_new_quote_with_provenance():
    record = _import(PriceMasterType.DT40, official_ihi_sku_snapshot_book, "Deep Trekker DT40 2026-10.xlsx")

    _, draft = _create_ihi_photon_draft()

    for line in draft.configuration_lines:
        snapshot = line.manufacturer_price_snapshot
        assert snapshot.price_source_type == PriceSourceType.OFFICIAL_PRICE_BOOK
        assert snapshot.price_book == "DT40"
        assert snapshot.price_master_import_id == record.import_id
        assert snapshot.price_master_sha256 == record.sha256
        assert snapshot.price_master_filename == "Deep Trekker DT40 2026-10.xlsx"
    assert all(item.price_master_import_id == record.import_id for item in draft.pricing_context.manufacturer_price_snapshots)
    assert any(record.import_id in item for item in draft.source_references)
    assert _line(draft).dealer_price_usd == 10434.6


def test_active_pt30_is_used_for_pipetrekker_prices_with_provenance():
    record = _import(PriceMasterType.PT30, official_pt30_style_book, "PT30.xlsx")
    books = load_official_manufacturer_price_books()
    assert [book.source_price_book for book in books] == ["PT30"]
    candidate = next(item for item in collect_manufacturer_candidates(*books) if item.dealer_price_usd is not None)

    scenario, _ = calculate_landed_cost_scenario(
        scenario_id="pt30",
        case_id="PT30",
        name="PT30",
        exchange_rate=160,
        policy=_policy(),
        product_inputs=[{"sku": candidate.sku, "quantity": 1}],
        shipping_lines=[],
        insurance_mode=InsuranceMode.PERCENTAGE,
        domestic_shipping_jpy=0,
        price_book_candidates=collect_manufacturer_candidates(*books),
    )
    draft = build_quote_draft(
        quote_draft_id="pt30",
        case_id="PT30",
        customer="x",
        title="PT30",
        configuration_name="PT30",
        landed_scenario=scenario,
        sales_candidates=[],
        presentation={"product_presentations": [{"sku": candidate.sku, "presentation_mode": "SEPARATE_LINE"}]},
    )
    attach_price_master_provenance(draft, [item.record for item in active_price_books()])

    snapshot = draft.configuration_lines[0].manufacturer_price_snapshot
    assert snapshot.price_source_type == PriceSourceType.OFFICIAL_PRICE_BOOK
    assert snapshot.price_book == "PT30"
    assert snapshot.price_master_import_id == record.import_id


def test_active_so_master_supplies_standard_sales_candidates():
    _import(PriceMasterType.SO_MASTER, pricing_policy_master_book, "SpaceOne master.xlsx")
    _import(PriceMasterType.DT40, pricing_policy_manufacturer_book, "dt40.xlsx")

    _, draft = _create_ihi_photon_draft()

    assert _line(draft).standard_sales_price_candidate_jpy == 3339072.0
    assert not any(SALES_MASTER_MISSING_WARNING in item for line in draft.configuration_lines for item in line.warnings)
    assert any("SO_MASTER" in item for item in draft.source_references)


def test_active_quote_calc_supplies_the_landed_policy_and_switches_with_a_new_import():
    first = _import(PriceMasterType.QUOTE_CALC, quote_calc_formula_book, "見積試算.xlsx")

    _, draft = _create_ihi_photon_draft()
    policy = draft.pricing_context.landed_cost_policy_snapshot
    assert policy.import_tax_rate == 0.1
    assert policy.insurance_rate == 0.03
    assert policy.landed_cost_policy_candidate_id == f"lcp-quote-calc:{first.import_id}"
    assert first.import_id in (policy.notes or "")

    second = _import(PriceMasterType.QUOTE_CALC, _quote_calc_with_memo_sheet, "見積試算 v2.xlsx")
    assert second.import_id != first.import_id
    assert load_active_landed_policy().landed_cost_policy_candidate_id == f"lcp-quote-calc:{second.import_id}"
    at, newer = _create_ihi_photon_draft()
    assert newer.pricing_context.landed_cost_policy_snapshot.landed_cost_policy_candidate_id.endswith(second.import_id)


def test_tmp_and_downloads_are_never_searched_by_business_code():
    for folder in ("agents", "parsers", "ui", "models", "repositories", "llm"):
        for path in (ROOT / folder).rglob("*.py"):
            text = path.read_text(encoding="utf-8")
            assert "/tmp/" not in text, path
            assert "Downloads" not in text, path
            assert "dt_price_investigation" not in text, path


def test_a_quote_calc_file_in_downloads_is_not_used(tmp_path, monkeypatch):
    downloads = tmp_path / "home" / "Downloads"
    downloads.mkdir(parents=True)
    (downloads / "DT_PT見積もり試算シート.xlsx").write_bytes(quote_calc_formula_book().getvalue())
    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path / "home"))

    assert locate_quote_calc_workbook() is None
    assert load_active_landed_policy() is None
    _, draft = _create_ihi_photon_draft()
    assert draft.pricing_context.landed_cost_policy_snapshot.import_tax_rate is None


def test_new_dt40_master_does_not_change_existing_drafts_and_applies_only_to_new_ones():
    first = _import(PriceMasterType.DT40, lambda: _photon_dt40_book(10434.6), "dt40-sep.xlsx")
    at, existing = _create_ihi_photon_draft()
    before = existing.model_dump(mode="json")
    stored_before = SqliteQuoteRepository().get_draft(existing.quote_draft_id).draft.model_dump(mode="json")

    second = _import(PriceMasterType.DT40, lambda: _photon_dt40_book(11000.0), "dt40-oct.xlsx")
    assert get_active_master(PriceMasterType.DT40).record.import_id == second.import_id

    stored_after = SqliteQuoteRepository().get_draft(existing.quote_draft_id).draft
    assert stored_after.model_dump(mode="json") == stored_before
    assert existing.model_dump(mode="json") == before
    assert _line(stored_after).dealer_price_usd == 10434.6
    assert _line(stored_after).manufacturer_price_snapshot.price_master_import_id == first.import_id

    at.run()
    assert _line(at.session_state["quote_draft"]).dealer_price_usd == 10434.6
    at.button(key="ihi_photon_draft").click().run()
    new = at.session_state["quote_draft"]
    assert _line(new).dealer_price_usd == 11000.0
    assert _line(new).manufacturer_price_snapshot.price_master_import_id == second.import_id


def test_new_master_does_not_change_an_approved_snapshot():
    _, snapshot = _approve(_ready_photon())
    repo = SqliteQuoteRepository()
    repo.save_snapshot(snapshot)
    before = repo.get_snapshot(snapshot.approved_quote_snapshot_id).model_dump(mode="json")

    _import(PriceMasterType.DT40, lambda: _photon_dt40_book(99999.0), "dt40-new.xlsx")
    _import(PriceMasterType.QUOTE_CALC, quote_calc_formula_book, "calc.xlsx")
    _import(PriceMasterType.SO_MASTER, pricing_policy_master_book, "so.xlsx")

    after = SqliteQuoteRepository().get_snapshot(snapshot.approved_quote_snapshot_id)
    assert after.model_dump(mode="json") == before
    assert after.total_jpy == 7876000

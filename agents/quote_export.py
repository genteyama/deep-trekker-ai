from datetime import date, datetime, timezone
from pathlib import Path
from typing import Optional
import csv
import hashlib
import math
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from agents.quote_approval import generate_quote_outputs
from agents.quote_dates import parse_quote_date
from models import (
    ApprovedQuoteExportBundle,
    ApprovedQuoteSnapshot,
    ExportBundleStatus,
    ExportFileManifest,
    ExportFileType,
    ExportPurpose,
    MoneyForwardQuotePayload,
    QuoteDraft,
    QuoteDraftStatus,
    QuoteOutputBundle,
)

OUTPUT_DIR = Path(__file__).resolve().parents[1] / "outputs"
MONEYFORWARD_COLUMNS = ("品目", "品目詳細", "単価", "数量", "金額", "備考")
INTERNAL_COLUMNS = (
    "Part Number",
    "品名",
    "数量",
    "DT単価USD",
    "DT金額USD",
    "卸値円",
    "輸入税",
    "保険",
    "国内送料",
    "仕入小計",
    "SpaceOne標準売価",
    "売価調整後単価",
    "売価金額",
    "粗利額",
    "粗利率",
)
CUSTOMER_FORBIDDEN_TERMS = (
    "Dealer",
    "Landed Cost",
    "Gross Margin",
    "DT40",
    "PT30",
    "Pricing Policy",
    "Manufacturer Discount",
    "Supplier Quote",
)
_UNSAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")
EXCEL_DATE_NUMBER_FORMAT = "yyyy/mm/dd"


class QuoteExportError(ValueError):
    pass


def default_output_dir() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR


def require_approved_snapshot(source: object) -> ApprovedQuoteSnapshot:
    if isinstance(source, QuoteDraft):
        raise QuoteExportError(
            "Export requires an ApprovedQuoteSnapshot. QuoteDraft is not a file source."
        )
    if not isinstance(source, ApprovedQuoteSnapshot):
        raise QuoteExportError("Export requires an ApprovedQuoteSnapshot.")
    if source.status not in {QuoteDraftStatus.APPROVED, QuoteDraftStatus.SUPERSEDED}:
        raise QuoteExportError(
            f"Export requires an approved snapshot. Current status is {source.status.value}."
        )
    return source


def validate_export_payloads(snapshot: ApprovedQuoteSnapshot, outputs: QuoteOutputBundle) -> list[str]:
    require_approved_snapshot(snapshot)
    if outputs.source_approved_quote_snapshot_id != snapshot.approved_quote_snapshot_id:
        raise QuoteExportError("Export payloads are not from the supplied ApprovedQuoteSnapshot.")
    warnings = []
    _validate_moneyforward_amounts(snapshot, outputs.moneyforward)
    _assert_equal(outputs.moneyforward.subtotal_ex_tax_jpy, snapshot.subtotal_ex_tax_jpy, "MoneyForward subtotal")
    _assert_equal(outputs.moneyforward.tax_jpy, snapshot.tax_jpy, "MoneyForward tax")
    _assert_equal(outputs.moneyforward.total_jpy, snapshot.total_jpy, "MoneyForward total")
    _assert_equal(outputs.spaceone_quote.subtotal, snapshot.subtotal_ex_tax_jpy, "SpaceOne subtotal")
    _assert_equal(outputs.spaceone_quote.tax, snapshot.tax_jpy, "SpaceOne tax")
    _assert_equal(outputs.spaceone_quote.total, snapshot.total_jpy, "SpaceOne total")
    _assert_equal(
        outputs.internal_transfer.customer_subtotal_ex_tax_jpy,
        snapshot.subtotal_ex_tax_jpy,
        "Internal customer subtotal",
    )
    _assert_equal(outputs.internal_transfer.customer_tax_jpy, snapshot.tax_jpy, "Internal customer tax")
    _assert_equal(outputs.internal_transfer.customer_total_jpy, snapshot.total_jpy, "Internal customer total")
    if len(outputs.moneyforward.rows) != len(snapshot.customer_lines_snapshot):
        raise QuoteExportError("MoneyForward row count does not match Approved customer lines.")
    if len(outputs.spaceone_quote.lines) != len(snapshot.customer_lines_snapshot):
        raise QuoteExportError("SpaceOne line count does not match Approved customer lines.")
    if len(outputs.internal_transfer.rows) != len(snapshot.configuration_snapshot):
        raise QuoteExportError("Internal BOM count does not match Approved configuration.")
    if outputs.spaceone_quote.remarks != [item.text for item in snapshot.remarks]:
        raise QuoteExportError("SpaceOne remarks must be the Approved Snapshot final remarks only.")
    leaked = scan_customer_payload_leaks(outputs.spaceone_quote.model_dump()) | scan_customer_payload_leaks(
        outputs.moneyforward.model_dump()
    )
    if leaked:
        raise QuoteExportError(f"Customer export payload contains internal terms: {sorted(leaked)}")
    return warnings


def build_export_bundle(
    snapshot: ApprovedQuoteSnapshot,
    *,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
    generated_at: Optional[datetime] = None,
) -> ApprovedQuoteExportBundle:
    snapshot = require_approved_snapshot(snapshot)
    outputs = generate_quote_outputs(snapshot)
    warnings = validate_export_payloads(snapshot, outputs)
    captured = generated_at or datetime.now(timezone.utc)
    official = _clean_optional(official_quote_number) or snapshot.official_quote_number
    return ApprovedQuoteExportBundle(
        export_bundle_id=f"exp-{snapshot.approved_quote_snapshot_id}-{captured.strftime('%Y%m%d%H%M%S')}",
        approved_quote_snapshot_id=snapshot.approved_quote_snapshot_id,
        case_id=snapshot.case_id,
        quote_version=snapshot.quote_version,
        official_quote_number=official,
        quote_number_candidate=snapshot.quote_number_candidate,
        generated_at=captured,
        generated_by=generated_by,
        moneyforward_payload=outputs.moneyforward,
        internal_transfer_payload=outputs.internal_transfer,
        spaceone_quote_payload=outputs.spaceone_quote,
        file_manifest=[],
        status=ExportBundleStatus.VALIDATED,
        warnings=warnings,
    )


def export_moneyforward_tsv(
    snapshot: ApprovedQuoteSnapshot,
    *,
    output_dir: Optional[Path] = None,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
    purpose: ExportPurpose = ExportPurpose.DEVELOPMENT,
    bundle: Optional[ApprovedQuoteExportBundle] = None,
) -> tuple[ApprovedQuoteExportBundle, Path]:
    return _export_moneyforward(
        snapshot,
        file_type=ExportFileType.MONEYFORWARD_TSV,
        output_dir=output_dir,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
        purpose=purpose,
        bundle=bundle,
    )


def export_moneyforward_csv(
    snapshot: ApprovedQuoteSnapshot,
    *,
    output_dir: Optional[Path] = None,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
    purpose: ExportPurpose = ExportPurpose.DEVELOPMENT,
    bundle: Optional[ApprovedQuoteExportBundle] = None,
) -> tuple[ApprovedQuoteExportBundle, Path]:
    return _export_moneyforward(
        snapshot,
        file_type=ExportFileType.MONEYFORWARD_CSV,
        output_dir=output_dir,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
        purpose=purpose,
        bundle=bundle,
    )


def export_internal_calc_excel(
    snapshot: ApprovedQuoteSnapshot,
    *,
    output_dir: Optional[Path] = None,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
    purpose: ExportPurpose = ExportPurpose.DEVELOPMENT,
    bundle: Optional[ApprovedQuoteExportBundle] = None,
) -> tuple[ApprovedQuoteExportBundle, Path]:
    bundle = bundle or build_export_bundle(
        snapshot,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
    )
    directory = Path(output_dir) if output_dir else default_output_dir()
    directory.mkdir(parents=True, exist_ok=True)
    path = unique_output_path(directory, _filename(bundle, "internal_calc", ".xlsx"))
    workbook = Workbook()
    _write_internal_metadata_sheet(workbook.active, snapshot, bundle)
    calc = workbook.create_sheet("Calculation")
    _write_internal_calc_sheet(calc, bundle)
    shipping = workbook.create_sheet("Shipping")
    _write_internal_shipping_sheet(shipping, bundle)
    summary = workbook.create_sheet("Summary")
    _write_internal_summary_sheet(summary, snapshot, bundle)
    workbook.save(path)
    _append_manifest(
        bundle,
        file_type=ExportFileType.INTERNAL_CALC_XLSX,
        path=path,
        purpose=purpose,
        row_count=len(bundle.internal_transfer_payload.rows),
        subtotal=snapshot.subtotal_ex_tax_jpy,
        tax=snapshot.tax_jpy,
        total=snapshot.total_jpy,
    )
    _validate_internal_excel(path, snapshot, bundle)
    return bundle, path


def export_spaceone_quote_excel(
    snapshot: ApprovedQuoteSnapshot,
    *,
    output_dir: Optional[Path] = None,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
    purpose: ExportPurpose = ExportPurpose.FORMAL,
    bundle: Optional[ApprovedQuoteExportBundle] = None,
) -> tuple[ApprovedQuoteExportBundle, Path]:
    bundle = bundle or build_export_bundle(
        snapshot,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
    )
    if purpose == ExportPurpose.FORMAL:
        if not bundle.official_quote_number:
            raise QuoteExportError("Formal SpaceOne quote Excel requires an official_quote_number.")
        if not snapshot.issue_date:
            raise QuoteExportError("Formal SpaceOne quote Excel requires issue_date.")
        if not snapshot.valid_until:
            raise QuoteExportError("Formal SpaceOne quote Excel requires valid_until.")
        if snapshot.issuer_snapshot is None:
            raise QuoteExportError("Formal SpaceOne quote Excel requires an issuer_snapshot.")
    directory = Path(output_dir) if output_dir else default_output_dir()
    directory.mkdir(parents=True, exist_ok=True)
    path = unique_output_path(directory, _filename(bundle, "spaceone_quote", ".xlsx"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "見積書"
    _write_spaceone_quote_sheet(sheet, snapshot, bundle, purpose)
    prepare_spaceone_quote_pdf_layout(sheet)
    workbook.save(path)
    leaked = scan_customer_workbook_leaks(path)
    if leaked:
        path.unlink(missing_ok=True)
        raise QuoteExportError(f"SpaceOne quote Excel contains internal terms: {sorted(leaked)}")
    _append_manifest(
        bundle,
        file_type=ExportFileType.SPACEONE_QUOTE_XLSX,
        path=path,
        purpose=purpose,
        row_count=len(bundle.spaceone_quote_payload.lines),
        subtotal=snapshot.subtotal_ex_tax_jpy,
        tax=snapshot.tax_jpy,
        total=snapshot.total_jpy,
        pdf_layout_prepared=True,
    )
    _validate_spaceone_excel(path, snapshot, bundle, purpose)
    return bundle, path


def export_ihi_photon_files(
    snapshot: ApprovedQuoteSnapshot,
    *,
    output_dir: Optional[Path] = None,
    official_quote_number: Optional[str] = None,
    generated_by: Optional[str] = None,
) -> tuple[ApprovedQuoteExportBundle, dict[str, Path]]:
    bundle = build_export_bundle(
        snapshot,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
    )
    paths = {}
    bundle, paths["moneyforward_tsv"] = export_moneyforward_tsv(
        snapshot, output_dir=output_dir, bundle=bundle, generated_by=generated_by
    )
    bundle, paths["moneyforward_csv"] = export_moneyforward_csv(
        snapshot, output_dir=output_dir, bundle=bundle, generated_by=generated_by
    )
    bundle, paths["internal_calc"] = export_internal_calc_excel(
        snapshot, output_dir=output_dir, bundle=bundle, generated_by=generated_by
    )
    spaceone_purpose = ExportPurpose.FORMAL if bundle.official_quote_number else ExportPurpose.DEVELOPMENT
    bundle, paths["spaceone_quote"] = export_spaceone_quote_excel(
        snapshot,
        output_dir=output_dir,
        bundle=bundle,
        generated_by=generated_by,
        purpose=spaceone_purpose,
    )
    return bundle, paths


def unique_output_path(directory: Path, filename: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    candidate = directory / filename
    if not candidate.exists():
        return candidate
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    stamped = directory / f"{candidate.stem}_{stamp}{candidate.suffix}"
    if not stamped.exists():
        return stamped
    index = 2
    while True:
        numbered = directory / f"{candidate.stem}_{stamp}_{index}{candidate.suffix}"
        if not numbered.exists():
            return numbered
        index += 1


def prepare_spaceone_quote_pdf_layout(sheet) -> None:
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.horizontalCentered = True
    sheet.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.6)
    last_row = max(_used_row(sheet), 1)
    sheet.print_area = f"A1:D{last_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def approved_remark_texts(snapshot: ApprovedQuoteSnapshot) -> list[str]:
    texts = []
    for item in snapshot.remarks or []:
        text = item.text if hasattr(item, "text") else str(item)
        if text and str(text).strip():
            texts.append(str(text).strip())
    return texts


def scan_customer_workbook_leaks(path: Path) -> set[str]:
    workbook = load_workbook(path, data_only=True)
    leaked = set()
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                leaked.update(_forbidden_terms_in(value))
    return leaked


def scan_customer_payload_leaks(payload: object) -> set[str]:
    leaked = set()
    if isinstance(payload, dict):
        for key, value in payload.items():
            if key in {"source_approved_quote_snapshot_id"}:
                continue
            leaked.update(scan_customer_payload_leaks(value))
    elif isinstance(payload, list):
        for item in payload:
            leaked.update(scan_customer_payload_leaks(item))
    else:
        leaked.update(_forbidden_terms_in(payload))
    return leaked


def _export_moneyforward(
    snapshot: ApprovedQuoteSnapshot,
    *,
    file_type: ExportFileType,
    output_dir: Optional[Path],
    official_quote_number: Optional[str],
    generated_by: Optional[str],
    purpose: ExportPurpose,
    bundle: Optional[ApprovedQuoteExportBundle],
) -> tuple[ApprovedQuoteExportBundle, Path]:
    bundle = bundle or build_export_bundle(
        snapshot,
        official_quote_number=official_quote_number,
        generated_by=generated_by,
    )
    directory = Path(output_dir) if output_dir else default_output_dir()
    directory.mkdir(parents=True, exist_ok=True)
    suffix = ".tsv" if file_type == ExportFileType.MONEYFORWARD_TSV else ".csv"
    path = unique_output_path(directory, _filename(bundle, "moneyforward", suffix))
    rows = bundle.moneyforward_payload.rows
    if file_type == ExportFileType.MONEYFORWARD_TSV:
        path.write_text(_moneyforward_text(rows, "\t"), encoding="utf-8")
    else:
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(MONEYFORWARD_COLUMNS)
            for row in rows:
                writer.writerow(_moneyforward_values(row))
    _validate_moneyforward_file(path, snapshot, bundle, file_type)
    _append_manifest(
        bundle,
        file_type=file_type,
        path=path,
        purpose=purpose,
        row_count=len(rows),
        subtotal=snapshot.subtotal_ex_tax_jpy,
        tax=snapshot.tax_jpy,
        total=snapshot.total_jpy,
    )
    return bundle, path


def _moneyforward_text(rows, delimiter: str) -> str:
    lines = [delimiter.join(MONEYFORWARD_COLUMNS)]
    for row in rows:
        lines.append(delimiter.join(_moneyforward_values(row)))
    return "\n".join(lines) + "\n"


def _moneyforward_values(row) -> list[str]:
    return [
        _cell_text(row.item_name),
        _cell_text(row.item_detail),
        _number_text(row.unit_price_jpy),
        str(row.quantity),
        _number_text(row.amount_jpy),
        _cell_text(row.notes),
    ]


def _validate_moneyforward_amounts(snapshot: ApprovedQuoteSnapshot, payload: MoneyForwardQuotePayload) -> None:
    line_total = 0.0
    for row in payload.rows:
        if row.unit_price_jpy is None or row.amount_jpy is None:
            raise QuoteExportError("MoneyForward row is missing unit price or amount.")
        expected = round(row.unit_price_jpy * row.quantity, 4)
        if abs(expected - row.amount_jpy) >= 1:
            raise QuoteExportError(
                f"MoneyForward amount mismatch: {row.item_name} {row.unit_price_jpy} x {row.quantity} != {row.amount_jpy}"
            )
        line_total += row.amount_jpy
    if abs(line_total - (snapshot.subtotal_ex_tax_jpy or 0)) >= 1:
        raise QuoteExportError("MoneyForward line total does not match Approved subtotal.")
    if abs((payload.tax_jpy or 0) - (snapshot.tax_jpy or 0)) >= 1:
        raise QuoteExportError("MoneyForward tax does not match Approved tax.")
    if abs((payload.total_jpy or 0) - (snapshot.total_jpy or 0)) >= 1:
        raise QuoteExportError("MoneyForward total does not match Approved total.")


def _validate_moneyforward_file(
    path: Path,
    snapshot: ApprovedQuoteSnapshot,
    bundle: ApprovedQuoteExportBundle,
    file_type: ExportFileType,
) -> None:
    raw = path.read_bytes()
    if file_type == ExportFileType.MONEYFORWARD_CSV and not raw.startswith(b"\xef\xbb\xbf"):
        raise QuoteExportError("MoneyForward CSV must be UTF-8 with BOM.")
    text = raw.decode("utf-8-sig")
    delimiter = "\t" if file_type == ExportFileType.MONEYFORWARD_TSV else ","
    parsed = list(csv.reader(text.splitlines(), delimiter=delimiter))
    if not parsed or parsed[0] != list(MONEYFORWARD_COLUMNS):
        raise QuoteExportError("MoneyForward file header is invalid.")
    data = parsed[1:]
    if len(data) != len(bundle.moneyforward_payload.rows):
        raise QuoteExportError("MoneyForward file row count does not match the Approved payload.")
    amounts = [float(row[4]) for row in data if row[4]]
    if abs(sum(amounts) - (snapshot.subtotal_ex_tax_jpy or 0)) >= 1:
        raise QuoteExportError("MoneyForward file amounts do not match Approved subtotal.")


def _write_internal_metadata_sheet(sheet, snapshot: ApprovedQuoteSnapshot, bundle: ApprovedQuoteExportBundle) -> None:
    sheet.title = "Metadata"
    rows = [
        ("Approved Snapshot ID", snapshot.approved_quote_snapshot_id),
        ("Quote Version", snapshot.quote_version),
        ("Case ID", snapshot.case_id),
        ("Official Quote Number", bundle.official_quote_number),
        ("Quote Number Candidate", bundle.quote_number_candidate),
        ("Approved At", _datetime_text(snapshot.approved_at)),
        ("Approved By", snapshot.approved_by),
        ("Exchange Rate", snapshot.exchange_rate),
        ("Tax Rate", snapshot.tax_rate),
        ("Gross Profit", snapshot.gross_profit_jpy),
        ("Gross Margin", snapshot.gross_margin_rate),
        ("Manufacturer Price Book Version", _price_book_versions(snapshot)),
        ("Customer Subtotal", snapshot.subtotal_ex_tax_jpy),
        ("Customer Tax", snapshot.tax_jpy),
        ("Customer Sales Total", snapshot.total_jpy),
        ("Generated At", _datetime_text(bundle.generated_at)),
        ("Generated By", bundle.generated_by),
    ]
    header_font = Font(bold=True)
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(index, 1, label).font = header_font
        cell = sheet.cell(index, 2, value)
        if label in {"Gross Profit", "Customer Subtotal", "Customer Tax", "Customer Sales Total"}:
            cell.number_format = "#,##0"
        elif label == "Gross Margin":
            cell.number_format = "0.00%"
        elif label == "Exchange Rate":
            cell.number_format = "0.00"
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 48


def _write_internal_calc_sheet(sheet, bundle: ApprovedQuoteExportBundle) -> None:
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    for column, title in enumerate(INTERNAL_COLUMNS, start=1):
        cell = sheet.cell(1, column, title)
        cell.font = header_font
        cell.fill = header_fill
    usd_columns = {4, 5}
    jpy_columns = {6, 7, 8, 9, 10, 11, 12, 13, 14}
    for row_index, row in enumerate(bundle.internal_transfer_payload.rows, start=2):
        values = [
            row.part_number,
            row.item_name,
            row.quantity,
            row.dealer_unit_price_usd,
            row.dealer_amount_usd,
            row.dealer_cost_jpy,
            row.import_tax_jpy,
            row.insurance_jpy,
            row.domestic_shipping_jpy,
            row.landed_subtotal_jpy,
            row.spaceone_standard_sales_jpy,
            row.adjusted_unit_price_jpy,
            row.sales_amount_jpy,
            row.gross_profit_jpy,
            row.gross_margin_rate,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            if column == 2:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if column in usd_columns:
                cell.number_format = "#,##0.00"
            elif column in jpy_columns:
                cell.number_format = "#,##0"
            elif column == 15:
                cell.number_format = "0.00%"
            if row.presentation_mode and row.presentation_mode.value in {"BUNDLED_WITH_PARENT", "INTERNAL_ONLY"}:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for column in range(1, len(INTERNAL_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 16
    sheet.column_dimensions["B"].width = 42
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(INTERNAL_COLUMNS))}{max(sheet.max_row, 1)}"
    sheet.freeze_panes = "A2"


def _write_internal_shipping_sheet(sheet, bundle: ApprovedQuoteExportBundle) -> None:
    headers = (
        "Shipping Type",
        "Quantity",
        "USD Rate",
        "USD Amount",
        "Exchange Rate",
        "Cost JPY",
        "Standard Sales Candidate",
        "Final Sales Price",
        "Gross Profit",
        "Gross Margin",
        "Source Snapshot",
    )
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    header_font = Font(bold=True)
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(1, column, title)
        cell.font = header_font
        cell.fill = header_fill
    for row_index, row in enumerate(bundle.internal_transfer_payload.shipping_rows, start=2):
        values = [
            row.shipping_type,
            row.quantity,
            row.usd_rate,
            row.usd_amount,
            row.exchange_rate,
            row.cost_jpy,
            row.standard_sales_candidate_jpy,
            row.final_sales_price_jpy,
            row.gross_profit_jpy,
            row.gross_margin_rate,
            row.source_snapshot_id,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            if column in {3, 4, 5}:
                cell.number_format = "#,##0.00"
            elif column in {6, 7, 8, 9}:
                cell.number_format = "#,##0.00"
            elif column == 10:
                cell.number_format = "0.00%"
            if row.line_role == "AGGREGATE":
                cell.font = Font(bold=True)
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22


def _write_internal_summary_sheet(sheet, snapshot, bundle: ApprovedQuoteExportBundle) -> None:
    payload = bundle.internal_transfer_payload
    header_font = Font(bold=True)
    rows = [
        ("Product Sales", payload.product_sales_ex_tax_jpy),
        ("Shipping Sales", payload.shipping_sales_ex_tax_jpy),
        ("Total Sales ex Tax", payload.customer_subtotal_ex_tax_jpy),
        ("Product Landed Cost", payload.product_landed_cost_jpy),
        ("Shipping Cost", payload.shipping_cost_jpy),
        ("Total Landed Cost", payload.total_landed_cost_jpy),
        ("Gross Profit", payload.gross_profit_jpy),
        ("Gross Margin", payload.gross_margin_rate),
        ("Approved Snapshot ID", snapshot.approved_quote_snapshot_id),
    ]
    for index, (label, value) in enumerate(rows, start=1):
        sheet.cell(index, 1, label).font = header_font
        cell = sheet.cell(index, 2, value)
        if label == "Gross Margin":
            cell.number_format = "0.00%"
        elif isinstance(value, float):
            cell.number_format = "#,##0.00"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 24


def _write_spaceone_quote_sheet(
    sheet,
    snapshot: ApprovedQuoteSnapshot,
    bundle: ApprovedQuoteExportBundle,
    purpose: ExportPurpose,
) -> None:
    payload = bundle.spaceone_quote_payload
    quote_number = bundle.official_quote_number or bundle.quote_number_candidate or ""
    if purpose == ExportPurpose.DEVELOPMENT:
        quote_number = f"{quote_number}（開発用）"
    title_font = Font(bold=True, size=18)
    label_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    sheet["A1"] = "見積書"
    sheet["A1"].font = title_font
    issuer = snapshot.issuer_snapshot
    if issuer is not None:
        sheet["C1"] = issuer.company_name
        sheet["C2"] = issuer.address
        sheet["C3"] = issuer.office_address
        sheet["C4"] = issuer.telephone
        for row in range(1, 5):
            sheet.cell(row, 3).alignment = Alignment(wrap_text=True)
    sheet["A6"] = "宛先"
    sheet["B6"] = f"{payload.customer} 御中" if payload.customer else ""
    sheet["A7"] = "件名"
    sheet["B7"] = payload.title
    sheet["A8"] = "見積番号"
    sheet["B8"] = quote_number
    sheet["A9"] = "発行日"
    _write_excel_date(sheet["B9"], payload.issue_date or snapshot.issue_date)
    sheet["A10"] = "有効期限"
    _write_excel_date(sheet["B10"], payload.valid_until or snapshot.valid_until)
    sheet["A12"] = "御見積金額"
    sheet["B12"] = snapshot.total_jpy
    sheet["B12"].number_format = '"¥"#,##0'
    sheet["B12"].font = Font(bold=True, size=14)
    for label_cell in ("A6", "A7", "A8", "A9", "A10", "A12"):
        sheet[label_cell].font = label_font
    headers = ("品目", "単価", "数量", "価格")
    header_row = 14
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, title)
        cell.font = label_font
        cell.fill = header_fill
        cell.border = thin
    current = header_row + 1
    for line in payload.lines:
        item_text = line.item_name or ""
        if line.item_detail:
            item_text = f"{item_text}\n{line.item_detail}" if item_text else line.item_detail
        name_cell = sheet.cell(current, 1, item_text)
        name_cell.alignment = Alignment(wrap_text=True, vertical="top")
        line_count = str(item_text).count("\n") + 1
        sheet.row_dimensions[current].height = max(18, 14 * line_count)
        unit = sheet.cell(current, 2, line.unit_price_jpy)
        unit.number_format = "#,##0"
        sheet.cell(current, 3, line.quantity)
        amount = sheet.cell(current, 4, line.amount_jpy)
        amount.number_format = "#,##0"
        for column in range(1, 5):
            sheet.cell(current, column).border = thin
        current += 1
    current += 1
    sheet.cell(current, 3, "小計").font = label_font
    subtotal = sheet.cell(current, 4, snapshot.subtotal_ex_tax_jpy)
    subtotal.number_format = "#,##0"
    current += 1
    sheet.cell(current, 3, "消費税").font = label_font
    tax = sheet.cell(current, 4, snapshot.tax_jpy)
    tax.number_format = "#,##0"
    current += 1
    sheet.cell(current, 3, "合計").font = label_font
    total = sheet.cell(current, 4, snapshot.total_jpy)
    total.number_format = "#,##0"
    total.font = Font(bold=True)
    current += 2
    sheet.cell(current, 1, "備考").font = label_font
    current += 1
    for remark in approved_remark_texts(snapshot):
        sheet.merge_cells(start_row=current, start_column=1, end_row=current, end_column=4)
        cell = sheet.cell(current, 1, remark)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(size=10)
        sheet.row_dimensions[current].height = _remark_row_height(remark)
        current += 1
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 36
    sheet.column_dimensions["D"].width = 16
    sheet.row_dimensions[1].height = 22


def _validate_internal_excel(path: Path, snapshot: ApprovedQuoteSnapshot, bundle: ApprovedQuoteExportBundle) -> None:
    workbook = load_workbook(path, data_only=True)
    metadata = {row[0]: row[1] for row in workbook["Metadata"].iter_rows(values_only=True)}
    if metadata.get("Approved Snapshot ID") != snapshot.approved_quote_snapshot_id:
        raise QuoteExportError("Internal Excel metadata snapshot id does not match.")
    if metadata.get("Customer Sales Total") != snapshot.total_jpy:
        raise QuoteExportError("Internal Excel customer sales total does not match Approved Snapshot.")
    calc = workbook["Calculation"]
    headers = [cell.value for cell in calc[1]]
    if headers != list(INTERNAL_COLUMNS):
        raise QuoteExportError("Internal Excel columns do not match the required order.")
    data_rows = [row for row in calc.iter_rows(min_row=2, values_only=True) if any(row)]
    if len(data_rows) != len(snapshot.configuration_snapshot):
        raise QuoteExportError("Internal Excel BOM count does not match Approved configuration.")
    parts = [row[0] for row in data_rows]
    for line in snapshot.configuration_snapshot:
        if line.manufacturer_sku not in parts:
            raise QuoteExportError(f"Internal Excel is missing BOM part {line.manufacturer_sku}.")
    if "Shipping" not in workbook.sheetnames or "Summary" not in workbook.sheetnames:
        raise QuoteExportError("Internal Excel must include Shipping and Summary sheets.")
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True)}
    if summary.get("Total Sales ex Tax") != snapshot.subtotal_ex_tax_jpy:
        raise QuoteExportError("Internal Summary sales do not match Approved Snapshot.")
    if summary.get("Total Landed Cost") != snapshot.total_landed_cost_jpy:
        raise QuoteExportError("Internal Summary landed cost does not match Approved Snapshot.")
    if summary.get("Gross Profit") != snapshot.gross_profit_jpy:
        raise QuoteExportError("Internal Summary gross profit does not match Approved Snapshot.")
    if summary.get("Gross Margin") != snapshot.gross_margin_rate:
        raise QuoteExportError("Internal Summary gross margin does not match Approved Snapshot.")


def _validate_spaceone_excel(
    path: Path,
    snapshot: ApprovedQuoteSnapshot,
    bundle: ApprovedQuoteExportBundle,
    purpose: ExportPurpose,
) -> None:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["見積書"]
    flat = []
    for row in sheet.iter_rows(values_only=True):
        flat.extend(row)
    if "見積書" not in flat or "御見積金額" not in flat:
        raise QuoteExportError("SpaceOne quote Excel is missing the required structure.")
    if snapshot.total_jpy not in flat or snapshot.subtotal_ex_tax_jpy not in flat or snapshot.tax_jpy not in flat:
        raise QuoteExportError("SpaceOne quote Excel totals do not match Approved Snapshot.")
    if purpose == ExportPurpose.FORMAL and bundle.official_quote_number not in flat:
        raise QuoteExportError("Formal SpaceOne quote Excel must contain the official quote number.")
    if purpose == ExportPurpose.FORMAL:
        if not _excel_date_cell_matches(sheet["B9"], snapshot.issue_date):
            raise QuoteExportError("Formal SpaceOne quote Excel must contain issue_date as yyyy/mm/dd.")
        if not _excel_date_cell_matches(sheet["B10"], snapshot.valid_until):
            raise QuoteExportError("Formal SpaceOne quote Excel must contain valid_until as yyyy/mm/dd.")
        if snapshot.issuer_snapshot and snapshot.issuer_snapshot.company_name not in flat:
            raise QuoteExportError("Formal SpaceOne quote Excel must use the issuer snapshot.")
    for remark in approved_remark_texts(snapshot):
        if remark not in flat:
            raise QuoteExportError("SpaceOne quote Excel is missing an Approved remark.")
    print_last = _print_area_last_row(sheet.print_area)
    used_last = _used_row(sheet)
    if print_last < used_last:
        raise QuoteExportError("SpaceOne quote print area does not include the final remarks row.")
    shipping = next((item for item in snapshot.customer_lines_snapshot if item.line_kind == "SHIPPING"), None)
    if shipping and shipping.display_name not in " ".join(str(item or "") for item in flat):
        raise QuoteExportError("SpaceOne quote Excel is missing the customer shipping display.")


def _append_manifest(
    bundle: ApprovedQuoteExportBundle,
    *,
    file_type: ExportFileType,
    path: Path,
    purpose: ExportPurpose,
    row_count: int,
    subtotal: Optional[float],
    tax: Optional[float],
    total: Optional[float],
    pdf_layout_prepared: bool = False,
) -> None:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    bundle.file_manifest.append(
        ExportFileManifest(
            file_type=file_type,
            filename=path.name,
            purpose=purpose,
            approved_snapshot_id=bundle.approved_quote_snapshot_id,
            quote_version=bundle.quote_version,
            generated_at=datetime.now(timezone.utc),
            subtotal_ex_tax=subtotal,
            tax=tax,
            total=total,
            row_count=row_count,
            sha256=digest,
            pdf_layout_prepared=pdf_layout_prepared,
        )
    )
    bundle.status = ExportBundleStatus.GENERATED


def _filename(bundle: ApprovedQuoteExportBundle, kind: str, suffix: str) -> str:
    if bundle.official_quote_number:
        stem = _safe_filename(bundle.official_quote_number)
    else:
        case = bundle.case_id or bundle.approved_quote_snapshot_id
        stem = f"{_safe_filename(case)}_v{bundle.quote_version}"
    return f"{stem}_{kind}{suffix}"


def _price_book_versions(snapshot: ApprovedQuoteSnapshot) -> str:
    versions = []
    for item in snapshot.manufacturer_price_snapshots:
        label = " ".join(part for part in [item.price_book, item.price_book_version] if part)
        if label and label not in versions:
            versions.append(label)
    return ", ".join(versions)


def _forbidden_terms_in(value: object) -> set[str]:
    if value is None or isinstance(value, (int, float)):
        return set()
    text = str(value)
    return {term for term in CUSTOMER_FORBIDDEN_TERMS if term.lower() in text.lower()}


def _assert_equal(actual: Optional[float], expected: Optional[float], label: str) -> None:
    if actual is None or expected is None or abs(actual - expected) >= 1:
        raise QuoteExportError(f"{label} does not match ApprovedQuoteSnapshot.")


def _safe_filename(value: str) -> str:
    cleaned = _UNSAFE_FILENAME.sub("_", value.strip())
    return cleaned.strip("._") or "quote"


def _clean_optional(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    cleaned = value.strip()
    return cleaned or None


def _write_excel_date(cell, value) -> None:
    parsed = parse_quote_date(value)
    cell.value = parsed
    if parsed is not None:
        cell.number_format = EXCEL_DATE_NUMBER_FORMAT


def _excel_date_cell_matches(cell, expected) -> bool:
    actual = parse_quote_date(cell.value)
    wanted = parse_quote_date(expected)
    if actual is None or wanted is None:
        return False
    if actual != wanted:
        return False
    if not isinstance(cell.value, (date, datetime)):
        return False
    return (cell.number_format or "").lower() == EXCEL_DATE_NUMBER_FORMAT


def _cell_text(value: Optional[str]) -> str:
    return "" if value is None else str(value).replace("\t", " ").replace("\n", " ")


def _number_text(value: Optional[float]) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return str(value)


def _datetime_text(value: Optional[datetime]) -> Optional[str]:
    if value is None:
        return None
    return value.isoformat()


def _used_row(sheet) -> int:
    last = 1
    for row in sheet.iter_rows(min_row=1, max_col=4):
        if any(cell.value not in (None, "") for cell in row):
            last = row[0].row
    return last


def _print_area_last_row(print_area: Optional[str]) -> int:
    if not print_area:
        return 0
    match = re.search(r"\$?D\$?(\d+)$", print_area.replace("'", ""))
    return int(match.group(1)) if match else 0


def _remark_row_height(text: str) -> float:
    wrapped = 0
    for line in str(text).splitlines() or [""]:
        wrapped += max(1, math.ceil(len(line) / 36))
    return max(16, 13 * wrapped)

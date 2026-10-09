from __future__ import annotations

from io import BytesIO
from typing import Optional
from urllib.parse import urlparse

from openpyxl import load_workbook

from agents.online_price_master import OnlinePriceReview, OnlinePriceReviewStatus, review_manufacturer_online_price
from agents.online_price_source import (
    OnlinePriceFetchError,
    OnlinePriceSource,
    fetch_workbook,
    resolve_registered_price_source,
)
from agents.price_master import XLSX_SIGNATURE
from models import ManufacturerPriceSource, PriceMasterType
from repositories.manufacturer_price_source_repository import (
    ManufacturerPriceSourceRepository,
    SOURCE_DEFAULTS,
    SOURCE_DT40,
    SOURCE_PT30,
    SOURCE_SPECTRA_GOLD,
    default_source,
)
from repositories.sqlite_price_master_repository import SqlitePriceMasterRepository

CHECK_OK = {"NO_CHANGE", "CHANGES_PENDING", "CONNECTED_FUTURE"}


class ManufacturerPriceSourceValidationError(ValueError):
    pass


def list_effective_sources(
    repository: Optional[ManufacturerPriceSourceRepository] = None,
) -> list[ManufacturerPriceSource]:
    repo = repository or ManufacturerPriceSourceRepository()
    result = []
    for source_key in SOURCE_DEFAULTS:
        try:
            stored = repo.get(source_key)
        except Exception:
            stored = None
        if stored is not None:
            result.append(stored)
            continue
        fallback = resolve_registered_price_source(
            source_key,
            master_type=_master_type(source_key),
            source_repository=repo,
        )
        result.append(default_source(source_key, source_url=fallback.url, enabled=fallback.configured))
    return result


def save_source_setting(
    source_key: str,
    source_url: str,
    enabled: bool,
    *,
    expected_row_version: int,
    repository: Optional[ManufacturerPriceSourceRepository] = None,
) -> ManufacturerPriceSource:
    clean_url = source_url.strip()
    if clean_url and not _valid_https_url(clean_url):
        raise ManufacturerPriceSourceValidationError("INVALID_URL")
    if enabled and not clean_url:
        raise ManufacturerPriceSourceValidationError("URL_REQUIRED")
    source = default_source(source_key, source_url=clean_url, enabled=enabled)
    return (repository or ManufacturerPriceSourceRepository()).save(
        source,
        expected_row_version=expected_row_version,
    )


def check_source_connection(
    source: ManufacturerPriceSource,
    *,
    source_repository: Optional[ManufacturerPriceSourceRepository] = None,
    price_master_repository: Optional[SqlitePriceMasterRepository] = None,
    fetcher=None,
) -> tuple[ManufacturerPriceSource, Optional[OnlinePriceReview]]:
    """Check and record status. This never imports or activates a price master."""
    repo = source_repository or ManufacturerPriceSourceRepository()
    if source.row_version < 1:
        raise ManufacturerPriceSourceValidationError("SAVE_BEFORE_CHECK")
    resolved = resolve_registered_price_source(
        source.source_key,
        master_type=_master_type(source.source_key),
        source_repository=repo,
    )
    online = OnlinePriceSource(
        master_type=_master_type(source.source_key),
        source_id=source.source_key,
        url=source.source_url,
        enabled=source.enabled,
        source_key=source.source_key,
        registry_row_version=source.row_version,
        authorization=resolved.authorization,
    )
    if source.source_key in {SOURCE_DT40, SOURCE_PT30}:
        review = review_manufacturer_online_price(
            PriceMasterType(source.source_key),
            repository=price_master_repository,
            source=online,
            fetcher=fetcher,
        )
        status = review.status.value
        error = None if review.status in {
            OnlinePriceReviewStatus.NO_CHANGE,
            OnlinePriceReviewStatus.CHANGES_PENDING,
        } else review.reason_code or review.status.value
    elif source.source_key == SOURCE_SPECTRA_GOLD:
        review = None
        status, error = _check_future_workbook(online, fetcher=fetcher)
    else:
        raise ManufacturerPriceSourceValidationError("UNKNOWN_SOURCE")
    updated = repo.record_check(
        source.source_key,
        status=status,
        error=error,
        expected_row_version=source.row_version,
    )
    if review is not None:
        # record_check is itself a Registry update; the approved review must bind to its new version.
        review.source_key = updated.source_key
        review.source_url = online.fetch_url
        review.source_enabled = updated.enabled
        review.source_row_version = updated.row_version
    return updated, review


def _check_future_workbook(source: OnlinePriceSource, *, fetcher=None) -> tuple[str, Optional[str]]:
    if not source.configured:
        return "NOT_CONFIGURED", "NOT_CONFIGURED"
    try:
        fetched = (fetcher or fetch_workbook)(source)
    except OnlinePriceFetchError as error:
        return error.reason_code, error.reason_code
    except Exception:
        return "FETCH_FAILED", "FETCH_FAILED"
    if not fetched.data.startswith(XLSX_SIGNATURE):
        return "NOT_XLSX", "NOT_XLSX"
    try:
        workbook = load_workbook(BytesIO(fetched.data), read_only=True, data_only=True)
        has_sheet = bool(workbook.sheetnames)
        workbook.close()
    except Exception:
        return "PARSER_FAILED", "PARSER_FAILED"
    if not has_sheet:
        return "NO_DATA", "NO_DATA"
    # FUTURE only: connectivity/xlsx structure is confirmed; no pricing identity is inferred.
    return "CONNECTED_FUTURE", None


def _master_type(source_key: str) -> Optional[PriceMasterType]:
    if source_key in {SOURCE_DT40, SOURCE_PT30}:
        return PriceMasterType(source_key)
    return None


def _valid_https_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme.lower() == "https" and bool(parsed.netloc) and not parsed.username and not parsed.password

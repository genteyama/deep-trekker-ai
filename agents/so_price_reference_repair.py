from __future__ import annotations

import csv
import hashlib
import json
import re
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from agents.so_price_reference_audit import (
    FIELD_DEALER,
    FIELD_MSRP,
    ReferenceAuditStatus,
    _locate_so_columns,
    audit_price_references,
)
from parsers.spaceone_master_parser import _iferror_fallback, parse_cell_reference

Source = Union[str, Path, BinaryIO]


class RepairPreviewError(RuntimeError):
    pass


@dataclass(frozen=True)
class SourceSnapshot:
    import_id: str
    sha256: str


@dataclass
class RepairPlanItem:
    so_sheet: str
    so_row: int
    so_sku: Optional[str]
    field: str
    target_so_cell: str
    current_formula: str
    current_workbook: str
    current_sheet: str
    current_cell: str
    current_referenced_sku: Optional[str]
    expected_workbook: str
    expected_sheet: str
    expected_cell: str
    expected_sku: str
    official_value: Optional[float]
    current_fallback_value: Optional[float]
    fallback_matches_official: Optional[bool]
    replacement_formula: str
    source_so_import_id: str
    source_so_sha256: str
    manufacturer_import_id: str
    manufacturer_sha256: str
    status: str = "READY"
    reason: str = "Existing exact manufacturer reference will be moved to the expected row."


@dataclass
class RepairPreviewResult:
    plan: list[RepairPlanItem] = field(default_factory=list)
    audit_summary_before: dict = field(default_factory=dict)
    audit_summary_after: dict = field(default_factory=dict)
    logical_diff: list[dict] = field(default_factory=list)
    output_paths: dict = field(default_factory=dict)

    def summary(self) -> dict:
        repair_rows = {(item.so_sheet, item.so_row) for item in self.plan}
        fallback_match = sum(item.fallback_matches_official is True for item in self.plan)
        fallback_mismatch = sum(item.fallback_matches_official is False for item in self.plan)
        fallback_unavailable = sum(item.fallback_matches_official is None for item in self.plan)
        return {
            "repair_plan_fields": len(self.plan),
            "unique_so_rows": len(repair_rows),
            "preview_applied": len(self.plan),
            "formula_validation_passed": len(self.plan),
            "logical_diff_cells": len(self.logical_diff),
            "fallback_match": fallback_match,
            "fallback_mismatch": fallback_mismatch,
            "fallback_unavailable": fallback_unavailable,
            "audit_before": self.audit_summary_before,
            "audit_after": self.audit_summary_after,
        }


def create_repair_preview(
    so_source: Source,
    dt40_source: Source,
    pt30_source: Source,
    *,
    output_dir: Union[str, Path],
    snapshots: dict[str, SourceSnapshot],
    expected_fields: int = 64,
    expected_rows: int = 32,
) -> RepairPreviewResult:
    """Build and apply an all-or-nothing repair plan to an xlsx copy only."""
    so_bytes = _source_bytes(so_source)
    dt40_bytes = _source_bytes(dt40_source)
    pt30_bytes = _source_bytes(pt30_source)
    _verify_source_sha("SO_MASTER", so_bytes, snapshots)
    _verify_source_sha("DT40", dt40_bytes, snapshots)
    _verify_source_sha("PT30", pt30_bytes, snapshots)

    audit = audit_price_references(
        BytesIO(so_bytes),
        BytesIO(dt40_bytes),
        BytesIO(pt30_bytes),
        source_names={name.lower(): snapshots[name].import_id for name in snapshots},
    )
    repairable = [
        row
        for row in audit.rows
        if row.status == ReferenceAuditStatus.EXACT_REPAIRABLE and row.safe_to_repair
    ]
    repair_rows = {(row.so_sheet, row.so_row) for row in repairable}
    if len(repairable) != expected_fields or len(repair_rows) != expected_rows:
        raise RepairPreviewError(
            f"STALE_AUDIT: expected {expected_fields} fields/{expected_rows} rows, "
            f"found {len(repairable)} fields/{len(repair_rows)} rows"
        )

    plan = _build_plan(so_bytes, repairable, snapshots)
    preview_bytes, logical_diff = _apply_plan(so_bytes, plan)
    after = audit_price_references(BytesIO(preview_bytes), BytesIO(dt40_bytes), BytesIO(pt30_bytes))
    _verify_after_audit(audit, after, plan)

    result = RepairPreviewResult(
        plan=plan,
        audit_summary_before=audit.summary(),
        audit_summary_after=after.summary(),
        logical_diff=logical_diff,
    )
    result.output_paths = _write_outputs(result, preview_bytes, output_dir)
    return result


def replace_reference_range(formula: str, *, current_sheet: str, current_cell: str, expected_sheet: str, expected_cell: str) -> str:
    column, row = _split_cell(current_cell)
    pattern = re.compile(
        rf"{re.escape(current_sheet)}!\$?{re.escape(column)}\$?{row}",
        re.IGNORECASE,
    )
    matches = list(pattern.finditer(formula))
    if len(matches) != 1:
        raise RepairPreviewError("UNSUPPORTED_FORMULA: expected one current manufacturer range")
    return pattern.sub(f"{expected_sheet}!{expected_cell}", formula, count=1)


def _build_plan(so_bytes: bytes, repairable, snapshots: dict[str, SourceSnapshot]) -> list[RepairPlanItem]:
    value_wb = load_workbook(BytesIO(so_bytes), data_only=True)
    formula_wb = load_workbook(BytesIO(so_bytes), data_only=False)
    plan = []
    try:
        columns_by_sheet = {}
        for sheet in value_wb.worksheets:
            columns_by_sheet[sheet.title] = _locate_so_columns(sheet)
        for row in repairable:
            if not all(
                (
                    row.current_formula,
                    row.current_workbook,
                    row.current_sheet,
                    row.current_cell,
                    row.expected_workbook,
                    row.expected_sheet,
                    row.expected_cell,
                    row.expected_sku,
                )
            ):
                raise RepairPreviewError("STALE_AUDIT: repairable row lacks required reference metadata")
            if row.current_workbook != row.expected_workbook:
                raise RepairPreviewError("WORKBOOK_SWITCH_REQUIRES_REVIEW")
            columns = columns_by_sheet.get(row.so_sheet)
            if not columns:
                raise RepairPreviewError(f"STALE_AUDIT: SO header missing for {row.so_sheet}")
            column_key = "msrp" if row.field == FIELD_MSRP else "dealer"
            column_index = columns.get(column_key)
            if column_index is None:
                raise RepairPreviewError(f"STALE_AUDIT: SO {row.field} column missing")
            target = f"{get_column_letter(column_index + 1)}{row.so_row}"
            current_formula = formula_wb[row.so_sheet][target].value
            if current_formula != row.current_formula:
                raise RepairPreviewError(f"STALE_AUDIT: current formula changed at {row.so_sheet}!{target}")
            replacement = replace_reference_range(
                current_formula,
                current_sheet=row.current_sheet,
                current_cell=row.current_cell,
                expected_sheet=row.expected_sheet,
                expected_cell=row.expected_cell,
            )
            parsed = parse_cell_reference(replacement)
            if (
                parsed is None
                or parsed.workbook != row.expected_workbook
                or parsed.sheet != row.expected_sheet
                or _canonical_cell(parsed.cell) != _canonical_cell(row.expected_cell)
            ):
                raise RepairPreviewError(f"FORMULA_VERIFICATION_FAILED: {row.so_sheet}!{target}")
            fallback = _iferror_fallback(current_formula)
            matches = None
            if fallback is not None and row.official_value is not None:
                matches = round(float(fallback), 2) == round(float(row.official_value), 2)
            manufacturer = snapshots[row.expected_workbook]
            plan.append(
                RepairPlanItem(
                    so_sheet=row.so_sheet,
                    so_row=row.so_row,
                    so_sku=row.normalized_sku or row.so_sku,
                    field=row.field,
                    target_so_cell=target,
                    current_formula=current_formula,
                    current_workbook=row.current_workbook,
                    current_sheet=row.current_sheet,
                    current_cell=row.current_cell,
                    current_referenced_sku=row.current_referenced_sku,
                    expected_workbook=row.expected_workbook,
                    expected_sheet=row.expected_sheet,
                    expected_cell=row.expected_cell,
                    expected_sku=row.expected_sku,
                    official_value=row.official_value,
                    current_fallback_value=fallback,
                    fallback_matches_official=matches,
                    replacement_formula=replacement,
                    source_so_import_id=snapshots["SO_MASTER"].import_id,
                    source_so_sha256=snapshots["SO_MASTER"].sha256,
                    manufacturer_import_id=manufacturer.import_id,
                    manufacturer_sha256=manufacturer.sha256,
                )
            )
    finally:
        formula_wb.close()
        value_wb.close()
    return plan


def _apply_plan(so_bytes: bytes, plan: list[RepairPlanItem]) -> tuple[bytes, list[dict]]:
    workbook = load_workbook(BytesIO(so_bytes), data_only=False)
    try:
        # Validate the complete batch before the first mutation.
        for item in plan:
            if workbook[item.so_sheet][item.target_so_cell].value != item.current_formula:
                raise RepairPreviewError(f"STALE_AUDIT: current formula changed at {item.so_sheet}!{item.target_so_cell}")
            parsed = parse_cell_reference(item.replacement_formula)
            if (
                parsed is None
                or parsed.workbook != item.expected_workbook
                or parsed.sheet != item.expected_sheet
                or _canonical_cell(parsed.cell) != _canonical_cell(item.expected_cell)
            ):
                raise RepairPreviewError("FORMULA_VERIFICATION_FAILED")
        for item in plan:
            workbook[item.so_sheet][item.target_so_cell] = item.replacement_formula
        output = BytesIO()
        workbook.save(output)
        preview_bytes = output.getvalue()
    finally:
        workbook.close()
    logical_diff = _logical_diff(so_bytes, preview_bytes)
    expected = {(item.so_sheet, item.target_so_cell) for item in plan}
    actual = {(item["sheet"], item["cell"]) for item in logical_diff}
    if actual != expected:
        raise RepairPreviewError("LOGICAL_DIFF_OUTSIDE_PLAN")
    return preview_bytes, logical_diff


def _logical_diff(original: bytes, preview: bytes) -> list[dict]:
    original_wb = load_workbook(BytesIO(original), data_only=False)
    preview_wb = load_workbook(BytesIO(preview), data_only=False)
    changes = []
    try:
        if original_wb.sheetnames != preview_wb.sheetnames:
            raise RepairPreviewError("LOGICAL_DIFF_SHEET_STRUCTURE_CHANGED")
        for sheet_name in original_wb.sheetnames:
            left = original_wb[sheet_name]
            right = preview_wb[sheet_name]
            max_row = max(left.max_row or 0, right.max_row or 0)
            max_column = max(left.max_column or 0, right.max_column or 0)
            for row in range(1, max_row + 1):
                for column in range(1, max_column + 1):
                    before = left.cell(row, column).value
                    after = right.cell(row, column).value
                    if before != after:
                        changes.append(
                            {
                                "sheet": sheet_name,
                                "cell": right.cell(row, column).coordinate,
                                "before": before,
                                "after": after,
                            }
                        )
    finally:
        original_wb.close()
        preview_wb.close()
    return changes


def _verify_after_audit(before, after, plan: list[RepairPlanItem]) -> None:
    targets = {(item.so_sheet, item.so_row, item.field) for item in plan}
    after_by_key = {(row.so_sheet, row.so_row, row.field): row for row in after.rows}
    for key in targets:
        row = after_by_key.get(key)
        if row is None or row.status != ReferenceAuditStatus.CORRECT or row.safe_to_repair:
            raise RepairPreviewError(f"PREVIEW_AUDIT_FAILED: {key}")
    before_untouched = {
        (row.so_sheet, row.so_row, row.field): row.status
        for row in before.rows
        if (row.so_sheet, row.so_row, row.field) not in targets
    }
    after_untouched = {
        key: after_by_key[key].status
        for key in before_untouched
        if key in after_by_key
    }
    if before_untouched != after_untouched:
        raise RepairPreviewError("PREVIEW_AUDIT_CHANGED_NON_TARGET_STATUS")


def _write_outputs(result: RepairPreviewResult, preview_bytes: bytes, output_dir: Union[str, Path]) -> dict:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    preview_path = directory / "SO_MASTER_DQ5B_PREVIEW.xlsx"
    csv_path = directory / "dq5_repair_plan.csv"
    json_path = directory / "dq5_repair_plan.json"
    diff_path = directory / "dq5_repair_logical_diff.json"
    review_path = directory / "dq5_human_review_summary.csv"
    preview_path.write_bytes(preview_bytes)
    rows = [asdict(item) for item in result.plan]
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(RepairPlanItem.__dataclass_fields__))
        writer.writeheader()
        writer.writerows(rows)
    human_rows = _human_review_rows(result.plan)
    with review_path.open("w", encoding="utf-8", newline="") as handle:
        fieldnames = list(human_rows[0]) if human_rows else []
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        if fieldnames:
            writer.writeheader()
            writer.writerows(human_rows)
    json_path.write_text(
        json.dumps(
            {"summary": result.summary(), "plan": rows, "human_review_rows": human_rows},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    diff_path.write_text(
        json.dumps({"changed_cells": result.logical_diff}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {
        "preview": str(preview_path),
        "plan_csv": str(csv_path),
        "plan_json": str(json_path),
        "logical_diff": str(diff_path),
        "human_review": str(review_path),
    }


def _human_review_rows(plan: list[RepairPlanItem]) -> list[dict]:
    grouped = {}
    for item in plan:
        key = (item.so_sheet, item.so_row, item.so_sku)
        row = grouped.setdefault(
            key,
            {
                "so_sheet": item.so_sheet,
                "so_row": item.so_row,
                "so_sku": item.so_sku,
                "old_msrp_ref": None,
                "new_msrp_ref": None,
                "old_dealer_ref": None,
                "new_dealer_ref": None,
                "old_referenced_sku": item.current_referenced_sku,
                "expected_sku": item.expected_sku,
                "msrp_official": None,
                "dealer_official": None,
                "fallback_status": None,
            },
        )
        prefix = "msrp" if item.field == FIELD_MSRP else "dealer"
        row[f"old_{prefix}_ref"] = f"{item.current_sheet}!{item.current_cell}"
        row[f"new_{prefix}_ref"] = f"{item.expected_sheet}!{item.expected_cell}"
        row[f"{prefix}_official"] = item.official_value
        fallback = (
            "MATCH"
            if item.fallback_matches_official is True
            else "MISMATCH"
            if item.fallback_matches_official is False
            else "UNAVAILABLE"
        )
        existing = row["fallback_status"]
        row["fallback_status"] = fallback if existing in (None, fallback) else f"{existing}/{fallback}"
    return list(grouped.values())


def _verify_source_sha(name: str, data: bytes, snapshots: dict[str, SourceSnapshot]) -> None:
    snapshot = snapshots.get(name)
    if snapshot is None:
        raise RepairPreviewError(f"SOURCE_METADATA_MISSING: {name}")
    actual = hashlib.sha256(data).hexdigest()
    if actual != snapshot.sha256:
        raise RepairPreviewError(f"SOURCE_CHANGED: {name}")


def _split_cell(cell: str) -> tuple[str, int]:
    match = re.fullmatch(r"\$?([A-Za-z]+)\$?([0-9]+)", cell or "")
    if match is None:
        raise RepairPreviewError(f"INVALID_CELL_REFERENCE: {cell}")
    return match.group(1), int(match.group(2))


def _canonical_cell(cell: Optional[str]) -> Optional[str]:
    if not cell:
        return None
    column, row = _split_cell(cell)
    return f"{column.upper()}{row}"


def _source_bytes(source: Source) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    data = source.read()
    if hasattr(source, "seek"):
        source.seek(0)
    return data

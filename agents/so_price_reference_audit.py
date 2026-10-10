from __future__ import annotations

import csv
import json
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Optional, Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from models import SkuSourceStatus
from parsers.price_book_parser import (
    SKIP_SHEET_KEYS,
    _find_header,
    _merge_aliases,
    _normalize_header,
    normalize_sku,
    parse_money,
    parse_price_book,
    sheet_key,
)
from parsers.spaceone_master_parser import (
    HEADER_SCAN_ROWS,
    HELPER_SHEET_PREFIX,
    _iferror_fallback,
    _is_headerish,
    _is_legacy_shipping,
    classify_spaceone_sku,
    parse_cell_reference,
)

Source = Union[str, Path, BinaryIO]
FIELD_MSRP = "MSRP"
FIELD_DEALER = "DEALER"
AUDIT_FIELDS = (FIELD_MSRP, FIELD_DEALER)


class ReferenceAuditStatus:
    CORRECT = "CORRECT"
    EXACT_REPAIRABLE = "EXACT_REPAIRABLE"
    NO_REFERENCE = "NO_REFERENCE"
    UNSUPPORTED_FORMULA = "UNSUPPORTED_FORMULA"
    SKU_NOT_FOUND = "SKU_NOT_FOUND"
    AMBIGUOUS_OCCURRENCE = "AMBIGUOUS_OCCURRENCE"
    SOURCE_PRICE_CONFLICT = "SOURCE_PRICE_CONFLICT"
    OBSOLETE_ONLY = "OBSOLETE_ONLY"
    INVALID_SKU = "INVALID_SKU"
    MANUAL_REVIEW = "MANUAL_REVIEW"


@dataclass
class ReferenceAuditRow:
    so_sheet: str
    so_row: int
    so_sku: Optional[str]
    normalized_sku: Optional[str]
    field: str
    current_formula: Optional[str] = None
    current_workbook: Optional[str] = None
    current_sheet: Optional[str] = None
    current_cell: Optional[str] = None
    current_referenced_sku: Optional[str] = None
    expected_workbook: Optional[str] = None
    expected_sheet: Optional[str] = None
    expected_row: Optional[int] = None
    expected_cell: Optional[str] = None
    expected_sku: Optional[str] = None
    so_display_value: Optional[float] = None
    official_value: Optional[float] = None
    status: str = ReferenceAuditStatus.MANUAL_REVIEW
    safe_to_repair: bool = False
    reason: str = ""


@dataclass
class ReferenceAuditReport:
    rows: list[ReferenceAuditRow] = field(default_factory=list)
    sources: dict = field(default_factory=dict)

    def summary(self) -> dict:
        repair_rows = {(row.so_sheet, row.so_row) for row in self.rows if row.safe_to_repair}
        item_rows = {(row.so_sheet, row.so_row) for row in self.rows}

        def count(status: str) -> int:
            return sum(1 for row in self.rows if row.status == status)

        def mismatches(field_name: str) -> int:
            return sum(
                1
                for row in self.rows
                if row.field == field_name and _values_differ(row.so_display_value, row.official_value)
            )

        ambiguous = count(ReferenceAuditStatus.AMBIGUOUS_OCCURRENCE) + count(
            ReferenceAuditStatus.SOURCE_PRICE_CONFLICT
        )
        return {
            "so_item_count": len(item_rows),
            "reference_field_count": len(self.rows),
            "CORRECT": count(ReferenceAuditStatus.CORRECT),
            "EXACT_REPAIRABLE": count(ReferenceAuditStatus.EXACT_REPAIRABLE),
            "MANUAL_REVIEW": count(ReferenceAuditStatus.MANUAL_REVIEW),
            "SKU_NOT_FOUND": count(ReferenceAuditStatus.SKU_NOT_FOUND),
            "AMBIGUOUS": ambiguous,
            "AMBIGUOUS_OCCURRENCE": count(ReferenceAuditStatus.AMBIGUOUS_OCCURRENCE),
            "SOURCE_PRICE_CONFLICT": count(ReferenceAuditStatus.SOURCE_PRICE_CONFLICT),
            "NO_REFERENCE": count(ReferenceAuditStatus.NO_REFERENCE),
            "UNSUPPORTED_FORMULA": count(ReferenceAuditStatus.UNSUPPORTED_FORMULA),
            "OBSOLETE_ONLY": count(ReferenceAuditStatus.OBSOLETE_ONLY),
            "INVALID_SKU": count(ReferenceAuditStatus.INVALID_SKU),
            "msrp_mismatch_count": mismatches(FIELD_MSRP),
            "dealer_mismatch_count": mismatches(FIELD_DEALER),
            "unique_so_rows_needing_repair": len(repair_rows),
            "sources": self.sources,
        }


@dataclass
class _ManufacturerBook:
    name: str
    occurrences_by_sku: dict
    columns_by_sheet: dict
    sku_by_row: dict


def audit_price_references(
    so_source: Source,
    dt40_source: Source,
    pt30_source: Source,
    *,
    source_names: Optional[dict] = None,
) -> ReferenceAuditReport:
    """Read SO, DT40, and PT30 workbooks and classify manufacturer price references.

    The function never writes workbooks, the price-master registry, or quote data.
    """
    so_bytes = _source_bytes(so_source)
    dt40_bytes = _source_bytes(dt40_source)
    pt30_bytes = _source_bytes(pt30_source)
    books = {
        "DT40": _index_manufacturer(dt40_bytes, "DT40"),
        "PT30": _index_manufacturer(pt30_bytes, "PT30"),
    }
    report = ReferenceAuditReport(sources=source_names or {})
    formula_wb = load_workbook(BytesIO(so_bytes), data_only=False)
    value_wb = load_workbook(BytesIO(so_bytes), data_only=True)
    try:
        for formula_sheet, value_sheet in zip(formula_wb.worksheets, value_wb.worksheets):
            if formula_sheet.title.startswith(HELPER_SHEET_PREFIX):
                continue
            columns = _locate_so_columns(value_sheet)
            if columns is None:
                continue
            max_row = max(formula_sheet.max_row or 0, value_sheet.max_row or 0)
            for row_number in range(columns["start_row"], max_row + 1):
                item = _so_item(formula_sheet, value_sheet, row_number, columns)
                if item is None:
                    continue
                report.rows.extend(_audit_item(item, books))
    finally:
        formula_wb.close()
        value_wb.close()
    return report


def write_audit_report(report: ReferenceAuditReport, output_dir: Union[str, Path]) -> tuple[Path, Path]:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    csv_path = directory / "dq5_reference_audit.csv"
    json_path = directory / "dq5_reference_audit.json"
    rows = [asdict(row) for row in report.rows]
    fieldnames = list(ReferenceAuditRow.__dataclass_fields__)
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    json_path.write_text(
        json.dumps({"summary": report.summary(), "rows": rows}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return csv_path, json_path


def _audit_item(item: dict, books: dict) -> list[ReferenceAuditRow]:
    _attach_referenced_skus(item, books)
    if item["invalid"]:
        return [
            _closed_row(item, field_name, ReferenceAuditStatus.INVALID_SKU, "SO SKU is not a usable exact SKU.")
            for field_name in AUDIT_FIELDS
        ]
    if item["shipping"]:
        return [
            _closed_row(
                item,
                field_name,
                ReferenceAuditStatus.MANUAL_REVIEW,
                "Legacy shipping row is not a manufacturer SKU reference.",
            )
            for field_name in AUDIT_FIELDS
        ]
    rows = []
    for field_name in AUDIT_FIELDS:
        current = item["fields"][field_name]
        named_workbook = current["workbook"] if current["workbook"] in books else None
        resolution = _resolve_occurrence(item["normalized_sku"], named_workbook, books)
        rows.append(_audit_field(item, field_name, resolution))
    return rows


def _audit_field(item: dict, field_name: str, resolution: dict) -> ReferenceAuditRow:
    current = item["fields"][field_name]
    row = ReferenceAuditRow(
        so_sheet=item["sheet"],
        so_row=item["row"],
        so_sku=item["so_sku"],
        normalized_sku=item["normalized_sku"],
        field=field_name,
        current_formula=current["formula"],
        current_workbook=current["workbook"],
        current_sheet=current["sheet"],
        current_cell=current["cell"],
        current_referenced_sku=current["referenced_sku"],
        so_display_value=current["display_value"],
    )
    if current.get("column_missing"):
        row.status = ReferenceAuditStatus.MANUAL_REVIEW
        row.reason = "SO manufacturer price column could not be identified from the header."
        return row
    occurrence = resolution.get("occurrence")
    if occurrence is not None:
        book = resolution["books"][occurrence.source_price_book]
        expected_cell = _expected_cell(book, occurrence, field_name)
        row.expected_workbook = occurrence.source_price_book
        row.expected_sheet = occurrence.source_sheet
        row.expected_row = occurrence.source_row
        row.expected_cell = expected_cell
        row.expected_sku = occurrence.sku
        row.official_value = occurrence.msrp_usd if field_name == FIELD_MSRP else occurrence.dealer_price_usd
    if resolution["status"] != ReferenceAuditStatus.CORRECT:
        row.status = resolution["status"]
        row.reason = resolution["reason"]
        return row
    if not row.expected_cell or not _both_expected_cells(resolution):
        row.status = ReferenceAuditStatus.MANUAL_REVIEW
        row.reason = "Manufacturer price column could not be identified from the workbook header."
        return row
    if current["kind"] == "unsupported":
        row.status = ReferenceAuditStatus.UNSUPPORTED_FORMULA
        row.reason = "Formula is not a single manufacturer cell reference."
        return row
    if _reference_matches(current, row):
        row.status = ReferenceAuditStatus.CORRECT
        row.reason = "Current reference matches the unique active manufacturer cell."
        if _values_differ(row.so_display_value, row.official_value):
            row.reason = "Reference cell matches, but the displayed value differs from the official manufacturer price."
        return row
    if current["kind"] == "none":
        row.status = ReferenceAuditStatus.EXACT_REPAIRABLE
        row.safe_to_repair = True
        row.reason = "No manufacturer reference is present. The unique active cell is known."
        return row
    row.status = ReferenceAuditStatus.EXACT_REPAIRABLE
    row.safe_to_repair = True
    row.reason = "Current reference differs from the unique active manufacturer cell."
    return row


def _resolve_occurrence(sku: Optional[str], named_workbook: Optional[str], books: dict) -> dict:
    if not sku:
        return {"status": ReferenceAuditStatus.MANUAL_REVIEW, "reason": "SO SKU is empty.", "occurrence": None}
    selected = [books[named_workbook]] if named_workbook in books else list(books.values())
    occurrences = [item for book in selected for item in book.occurrences_by_sku.get(sku, [])]
    active = [item for item in occurrences if item.source_status != SkuSourceStatus.OBSOLETE]
    if not occurrences:
        scope = named_workbook or "DT40/PT30"
        return {
            "status": ReferenceAuditStatus.SKU_NOT_FOUND,
            "reason": f"No exact SKU occurrence exists in {scope}.",
            "occurrence": None,
        }
    if not active:
        return {
            "status": ReferenceAuditStatus.OBSOLETE_ONLY,
            "reason": "Exact SKU exists only on an obsolete manufacturer sheet.",
            "occurrence": None,
        }
    prices = {(_money(item.msrp_usd), _money(item.dealer_price_usd)) for item in active}
    if len(prices) > 1:
        return {
            "status": ReferenceAuditStatus.SOURCE_PRICE_CONFLICT,
            "reason": "Exact SKU has conflicting active manufacturer prices.",
            "occurrence": None,
        }
    if len(active) != 1:
        return {
            "status": ReferenceAuditStatus.AMBIGUOUS_OCCURRENCE,
            "reason": "Exact SKU has more than one active manufacturer row.",
            "occurrence": None,
        }
    return {
        "status": ReferenceAuditStatus.CORRECT,
        "reason": "",
        "occurrence": active[0],
        "books": books,
    }


def _both_expected_cells(resolution: dict) -> bool:
    occurrence = resolution.get("occurrence")
    if occurrence is None:
        return False
    book = resolution["books"][occurrence.source_price_book]
    return all(_expected_cell(book, occurrence, field_name) for field_name in AUDIT_FIELDS)


def _expected_cell(book: _ManufacturerBook, occurrence, field_name: str) -> Optional[str]:
    columns = book.columns_by_sheet.get(occurrence.source_sheet) or {}
    key = "msrp" if field_name == FIELD_MSRP else "dealer_price"
    index = columns.get(key)
    if index is None or occurrence.source_row is None:
        return None
    return f"{get_column_letter(index + 1)}{occurrence.source_row}"


def _reference_matches(current: dict, row: ReferenceAuditRow) -> bool:
    return (
        current["kind"] == "reference"
        and current["workbook"] == row.expected_workbook
        and current["sheet"] == row.expected_sheet
        and _canonical_cell(current["cell"]) == _canonical_cell(row.expected_cell)
    )


def _so_item(formula_sheet, value_sheet, row_number: int, columns: dict) -> Optional[dict]:
    sku_value = value_sheet.cell(row_number, columns["sku"] + 1).value
    normalized, invalid, _cell_type, raw = classify_spaceone_sku(sku_value)
    row_values = [cell.value for cell in value_sheet[row_number]]
    shipping = _is_legacy_shipping(normalized, " ".join(str(value) for value in row_values if value), None)
    if _is_headerish(raw or normalized):
        return None
    if normalized is None and not invalid and not shipping:
        return None
    fields = {}
    for field_name, column_key in ((FIELD_MSRP, "msrp"), (FIELD_DEALER, "dealer")):
        column = columns.get(column_key)
        formula = formula_sheet.cell(row_number, column + 1).value if column is not None else None
        display = value_sheet.cell(row_number, column + 1).value if column is not None else None
        parsed = _field_reference(formula, display)
        parsed["column_missing"] = column is None
        fields[field_name] = parsed
    return {
        "sheet": value_sheet.title,
        "row": row_number,
        "so_sku": raw if raw is not None else normalized,
        "normalized_sku": None if invalid else normalized,
        "invalid": bool(invalid),
        "shipping": shipping,
        "fields": fields,
    }


def _field_reference(formula, display) -> dict:
    formula_text = formula if isinstance(formula, str) else None
    display_value, invalid = parse_money(display)
    if display_value is None:
        display_value = _iferror_fallback(formula_text)
    referenced = {
        "formula": formula_text,
        "kind": "none",
        "workbook": None,
        "sheet": None,
        "cell": None,
        "referenced_sku": None,
        "display_value": None if invalid else display_value,
    }
    if formula_text is None:
        return referenced
    reference = parse_cell_reference(formula_text)
    if reference and reference.sheet and reference.cell and ":" not in reference.cell:
        referenced.update(
            {
                "kind": "reference",
                "workbook": reference.workbook,
                "sheet": reference.sheet,
                "cell": _canonical_cell(reference.cell),
            }
        )
        return referenced
    if formula_text.startswith("=") or "IMPORTRANGE" in formula_text.upper():
        referenced["kind"] = "unsupported"
    return referenced


def _locate_so_columns(sheet) -> Optional[dict]:
    aliases = _merge_aliases(None)
    sku_names = set(aliases["sku"])
    msrp_names = set(aliases["msrp"])
    dealer_names = set(aliases["dealer_price"])
    rows = list(sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, max_col=40, values_only=True))
    header_index = None
    sku_column = None
    for index, values in enumerate(rows):
        for column, value in enumerate(values):
            if _normalize_header(value) in sku_names:
                header_index = index
                sku_column = column
                break
        if header_index is not None:
            break
    if header_index is None:
        return None
    label_rows = [rows[header_index]]
    if header_index + 1 < len(rows) and _row_has_price_label(rows[header_index + 1], msrp_names, dealer_names):
        label_rows.append(rows[header_index + 1])
    msrp_column = _price_column(label_rows, msrp_names, "定価")
    dealer_column = _price_column(label_rows, dealer_names, "卸値")
    return {
        "sku": sku_column,
        "msrp": msrp_column,
        "dealer": dealer_column,
        "start_row": header_index + len(label_rows) + 1,
    }


def _price_column(rows, names, marker: str) -> Optional[int]:
    # 定価 / 卸値 identify the manufacturer price columns. 「設定価格」 contains 定価
    # as part of the sales-price header, so that column is not a manufacturer price.
    for values in rows:
        for column, value in enumerate(values):
            text = _normalize_header(value)
            if marker in text and "設定" not in text:
                return column
    for values in rows:
        for column, value in enumerate(values):
            if _normalize_header(value) in names:
                return column
    return None


def _row_has_price_label(values, msrp_names, dealer_names) -> bool:
    for value in values:
        text = _normalize_header(value)
        if text in msrp_names or text in dealer_names or "定価" in text or "卸値" in text:
            return True
    return False


def _index_manufacturer(data: bytes, book_name: str) -> _ManufacturerBook:
    parsed = parse_price_book(BytesIO(data), source_price_book=book_name)
    # Cached header text is the column source. The formula workbook stores those labels
    # behind IMPORTRANGE fallbacks, which is not a reliable header.
    workbook = load_workbook(BytesIO(data), data_only=True)
    columns_by_sheet = {}
    sku_by_row = {}
    try:
        aliases = _merge_aliases(None)
        for sheet in workbook.worksheets:
            if sheet_key(sheet.title) in SKIP_SHEET_KEYS:
                continue
            header = _find_header(sheet, aliases)
            if header is None:
                continue
            columns_by_sheet[sheet.title] = header["columns"]
            sku_column = header["columns"]["sku"] + 1
            for row_number in range(1, (sheet.max_row or 0) + 1):
                sku = _normalize_manufacturer_sku(sheet.cell(row_number, sku_column).value)
                if sku:
                    sku_by_row[(sheet.title, row_number)] = sku
    finally:
        workbook.close()
    grouped = defaultdict(list)
    for occurrence in parsed.occurrences:
        grouped[occurrence.sku].append(occurrence)
    return _ManufacturerBook(book_name, grouped, columns_by_sheet, sku_by_row)


def _normalize_manufacturer_sku(value) -> Optional[str]:
    return normalize_sku(value)


def _attach_referenced_skus(item: dict, books: dict) -> None:
    for parsed in item["fields"].values():
        book = books.get(parsed["workbook"])
        if book is None or not parsed["sheet"] or not parsed["cell"]:
            continue
        try:
            _letters, row_number = coordinate_from_string(parsed["cell"])
        except ValueError:
            continue
        parsed["referenced_sku"] = book.sku_by_row.get((parsed["sheet"], row_number))


def _closed_row(item: dict, field_name: str, status: str, reason: str) -> ReferenceAuditRow:
    current = item["fields"][field_name]
    return ReferenceAuditRow(
        so_sheet=item["sheet"],
        so_row=item["row"],
        so_sku=item["so_sku"],
        normalized_sku=item["normalized_sku"],
        field=field_name,
        current_formula=current["formula"],
        current_workbook=current["workbook"],
        current_sheet=current["sheet"],
        current_cell=current["cell"],
        current_referenced_sku=current["referenced_sku"],
        so_display_value=current["display_value"],
        status=status,
        safe_to_repair=False,
        reason=reason,
    )


def _canonical_cell(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    letters, row_number = coordinate_from_string(str(value).replace("$", ""))
    return f"{letters.upper()}{row_number}"


def _values_differ(left: Optional[float], right: Optional[float]) -> bool:
    if left is None or right is None:
        return False
    return _money(left) != _money(right)


def _money(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    return round(float(value), 2)


def _source_bytes(source: Source) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    data = source.read()
    if hasattr(source, "seek"):
        source.seek(0)
    return data

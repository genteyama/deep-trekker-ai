from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Optional
import csv
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from agents.activity_catalog import CATEGORY_LABELS, ActivityFilter, ActivityRow, display_day
from agents.quote_pdf import resolve_pdf_fonts
from models import CustomerRecord
from repositories.sqlite import now_iso, strip_secrets
from repositories.sqlite_customer_repository import customer_key


REPORT_TITLE = "Deep Trekker / PipeTrekker 対応履歴レポート"
CSV_HEADERS = ("会社名", "部署", "担当者名", "メール", "電話", "住所", "End User", "備考")
TABLE_HEADERS = ("No.", "日付", "顧客", "End User", "区分", "機種", "案件名", "対応ステータス")


def report_stamp(now: Optional[str] = None) -> str:
    if now:
        return now
    return datetime.now().strftime("%Y%m%d")


def activity_report_filename(ext: str, *, now: Optional[str] = None) -> str:
    return f"DeepTrekker_Activity_Report_{report_stamp(now)}{ext}"


def customer_csv_filename(*, now: Optional[str] = None) -> str:
    return f"DeepTrekker_Customers_{report_stamp(now)}.csv"


def describe_filters(filt: Optional[ActivityFilter]) -> str:
    current = filt or ActivityFilter()
    parts = []
    if current.date_from or current.date_to:
        parts.append(f"期間: {current.date_from or '-'} 〜 {current.date_to or '-'}")
    if current.customers:
        parts.append("顧客: " + " / ".join(current.customers))
    if current.end_users:
        parts.append("End User: " + " / ".join(current.end_users))
    if current.products:
        parts.append("機種: " + " / ".join(current.products))
    if current.categories:
        parts.append("区分: " + " / ".join(CATEGORY_LABELS.get(item, item) for item in current.categories))
    if current.statuses:
        parts.append("対応ステータス: " + " / ".join(current.statuses))
    if current.keyword:
        parts.append(f"キーワード: {current.keyword}")
    if current.include_archived:
        parts.append("アーカイブを含む")
    return " / ".join(parts) if parts else "なし"


def _safe_text(value) -> str:
    cleaned = strip_secrets(value if value is not None else "")
    if isinstance(cleaned, str):
        return cleaned
    return str(cleaned or "")


def report_rows(rows: list[ActivityRow]) -> list[tuple[str, ...]]:
    data = []
    for index, row in enumerate(rows, start=1):
        data.append(
            (
                str(index),
                display_day(row.date) or _safe_text(row.date)[:10],
                _safe_text(row.customer),
                _safe_text(row.end_user),
                CATEGORY_LABELS.get(row.category, row.category),
                _safe_text(row.product),
                _safe_text(row.title),
                _safe_text(row.status),
            )
        )
    return data


def _set_east_asia(run, name: str = "Yu Gothic") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def render_activity_report_docx(
    rows: list[ActivityRow],
    path: Path,
    *,
    filt: Optional[ActivityFilter] = None,
    exported_at: Optional[str] = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    heading = document.add_heading(REPORT_TITLE, level=1)
    for run in heading.runs:
        _set_east_asia(run)
    meta = document.add_paragraph()
    stamp = exported_at or now_iso()
    meta.add_run(f"出力日時: {stamp}")
    _set_east_asia(meta.runs[0])
    filters = document.add_paragraph()
    filters.add_run(f"Filter条件: {describe_filters(filt)}")
    _set_east_asia(filters.runs[0])
    count = document.add_paragraph()
    count.add_run(f"選択件数: {len(rows)}")
    _set_east_asia(count.runs[0])
    table = document.add_table(rows=1 + len(rows), cols=len(TABLE_HEADERS))
    table.style = "Table Grid"
    for index, header in enumerate(TABLE_HEADERS):
        cell = table.rows[0].cells[index]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                _set_east_asia(run)
    for row_index, values in enumerate(report_rows(rows), start=1):
        for col, value in enumerate(values):
            cell = table.rows[row_index].cells[col]
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_east_asia(run)
    document.save(path)
    return path


def render_activity_report_pdf(
    rows: list[ActivityRow],
    path: Path,
    *,
    filt: Optional[ActivityFilter] = None,
    exported_at: Optional[str] = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    regular, bold, _source = resolve_pdf_fonts()
    styles = {
        "title": ParagraphStyle("ActivityTitle", fontName=bold, fontSize=14, leading=18),
        "meta": ParagraphStyle("ActivityMeta", fontName=regular, fontSize=9, leading=13),
        "cell": ParagraphStyle("ActivityCell", fontName=regular, fontSize=8, leading=11),
        "head": ParagraphStyle("ActivityHead", fontName=bold, fontSize=8, leading=11),
    }
    story = [
        Paragraph(REPORT_TITLE, styles["title"]),
        Spacer(1, 6),
        Paragraph(f"出力日時: {exported_at or now_iso()}", styles["meta"]),
        Paragraph(f"Filter条件: {describe_filters(filt)}", styles["meta"]),
        Paragraph(f"選択件数: {len(rows)}", styles["meta"]),
        Spacer(1, 8),
    ]
    data = [[Paragraph(_safe_text(header), styles["head"]) for header in TABLE_HEADERS]]
    for values in report_rows(rows):
        data.append([Paragraph(_safe_text(value), styles["cell"]) for value in values])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), regular),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12324A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9AA4B2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(table)
    document = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=REPORT_TITLE,
    )
    document.build(story)
    return path


def customer_records_from_rows(
    rows: list[ActivityRow],
    customer_repo=None,
    *,
    include_master: bool = False,
) -> list[CustomerRecord]:
    seen: dict[str, CustomerRecord] = {}
    now = now_iso()
    for row in rows:
        name = (row.customer or "").strip()
        if not name:
            continue
        key = customer_key(name, None)
        master = customer_repo.find(name, None) if customer_repo is not None else None
        if master is not None:
            if key not in seen:
                seen[key] = master
            continue
        if key not in seen:
            seen[key] = CustomerRecord(
                customer_id=f"OBS-{key.replace('|', '-')}",
                customer_name=name,
                end_user=row.end_user or None,
                created_at=now,
                updated_at=now,
            )
    if include_master and customer_repo is not None:
        for master in customer_repo.list_customers():
            key = customer_key(master.customer_name, master.contact_name)
            seen[key] = master
    return sorted(seen.values(), key=lambda item: customer_key(item.customer_name, item.contact_name))


def render_customer_csv(records: list[CustomerRecord], path: Optional[Path] = None) -> tuple[str, Optional[Path]]:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(CSV_HEADERS)
    for record in records:
        writer.writerow(
            [
                _safe_text(record.customer_name),
                _safe_text(record.department),
                _safe_text(record.contact_name),
                _safe_text(record.email),
                _safe_text(record.phone),
                _safe_text(record.address),
                _safe_text(record.end_user),
                _safe_text(record.notes),
            ]
        )
    text = buffer.getvalue()
    encoded = ("\ufeff" + text).encode("utf-8")
    if path is not None:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(encoded)
    return "\ufeff" + text, path


def find_existing_formal_pdf(snapshot, search_dir: Optional[Path] = None) -> Optional[Path]:
    number = getattr(snapshot, "official_quote_number", None) or getattr(snapshot, "quote_number_candidate", None)
    if not number:
        return None
    directory = Path(search_dir) if search_dir else Path("runtime") / "exports"
    if not directory.exists():
        return None
    matches = sorted(directory.glob(f"*{number}*.pdf"))
    return matches[-1] if matches else None
